"""Sweep the corpus decks for where replace_text puts new words.

Usage, from the repo root:
    uv run --project deckcheck python deckcheck/scripts/splice_sweep.py

Three checks, on every text paragraph of the cached corpus decks:
- boundaries: where one run ends a word and the next, differently formatted run starts with a space, insert " gross"
  at the run edge. The new word must land in the next run, beside its space.
- same run: an edit inside one run must change that run's text and nothing else in the paragraph.
- rewrites: random rewrites across runs must leave exactly the new text, or be refused with a SpliceError.

Each check expects 0 wrong. The cached corpus holds 261 such run edges.

Exit 0 on SWEEP PASS or when the corpus cache is missing, 1 on SWEEP FAIL.
"""

from __future__ import annotations

import copy
import random
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from deckcheck.changeset import text as T

CACHE = Path(__file__).resolve().parents[2] / "artifacts/verify-pptx/corpus-cache"
WORDS = ["new", "the", "c.", "12%", "growth", ",", "and", "$4.1B", "(2023)", "—", "x", "  "]


def shapes(coll):
    for sh in coll:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shapes(sh.shapes)
        else:
            yield sh


def paragraphs() -> list[tuple[str, etree._Element]]:
    out = []
    for path in sorted(CACHE.glob("*.pptx")):
        for slide in Presentation(str(path)).slides:
            for sh in shapes(slide.shapes):
                if sh.has_text_frame:
                    out += [(f"{path.name[:8]} slide {slide.slide_id} shape {sh.shape_id}", p) for p in sh.text_frame._txBody.findall(qn("a:p"))]
    return out


def run_text(x: T.Atom) -> str:
    return x.el.findtext(T.A_T) or ""


def look(el: etree._Element) -> str:
    rpr = el.find(qn("a:rPr"))
    if rpr is None:
        return ""
    bits = [f"{k}={v}" for k, v in sorted(rpr.attrib.items()) if k in ("b", "i", "u", "baseline")]
    if rpr.find(qn("a:hlinkClick")) is not None:
        bits.append("link")
    fill = rpr.find(qn("a:solidFill"))
    if fill is not None and len(fill):
        bits.append("fill=" + (fill[0].get("val") or ""))
    return ",".join(bits)


def splice(p: etree._Element, at: int, old: str, new: str) -> etree._Element:
    q = copy.deepcopy(p)
    T.write_splice(q, T.plan_splice(q, at, old, new))
    return q


def boundaries(paras) -> tuple[int, list[str]]:
    total, wrong = 0, []
    for where, p in paras:
        xs = T.atoms(p)
        if any(not x.is_run for x in xs):
            continue
        full = T.paragraph_text(p)
        for a, b in zip(xs, xs[1:]):
            if not re.search(r"\S$", run_text(a)) or not re.match(r" \S", run_text(b)) or look(a.el) == look(b.el):
                continue
            new = full[: a.hi] + " gross" + full[a.hi :]
            q = splice(p, 0, full, new)
            home = next(r for r in q.findall(T.A_R) if "gross" in (r.findtext(T.A_T) or ""))
            total += 1
            if T.paragraph_text(q) != new or look(home) != look(b.el):
                wrong.append(f"{where}: {run_text(a)[-20:]!r} [{look(a.el)}] | {run_text(b)[:20]!r} [{look(b.el)}] put 'gross' in [{look(home)}]")
    return total, wrong


def same_run(paras, rng: random.Random) -> tuple[int, list[str]]:
    total, wrong = 0, []
    for where, p in paras:
        for x in T.atoms(p):
            if not x.is_run:
                continue
            for m in re.finditer(r"\w{3,}", run_text(x)):
                if m.start() < 2 or m.end() > len(run_text(x)) - 2 or rng.random() > 0.05:
                    continue
                lo, hi = max(0, m.start() - 8), min(len(run_text(x)), m.end() + 8)
                old, a, b = run_text(x)[lo:hi], m.start() - lo, m.end() - lo
                new = rng.choice([old[:a] + "alpha" + old[b:], old[:b] + " beta gamma" + old[b:], old[:a] + old[b:], old[:b] + "," + old[b:]])
                expected = copy.deepcopy(p)
                t = list(expected)[list(p).index(x.el)].find(T.A_T)
                t.text = run_text(x)[:lo] + new + run_text(x)[hi:]
                total += 1
                if etree.tostring(splice(p, x.lo + lo, old, new)) != etree.tostring(expected):
                    wrong.append(f"{where}: {old!r} -> {new!r} changed more than its run")
    return total, wrong


def rewrites(paras, rng: random.Random, n: int = 3000) -> tuple[int, list[str]]:
    pool = [p for _, p in paras if len(T.atoms(p)) >= 2 and len(T.paragraph_text(p)) > 15]
    total, wrong = 0, []
    for _ in range(n):
        p = rng.choice(pool)
        full = T.paragraph_text(p)
        a = rng.randrange(len(full))
        b = rng.randrange(a + 1, len(full) + 1)
        words = re.findall(r"\s+|\S+", full[a:b])
        for _ in range(rng.randint(1, 4)):
            i = min(rng.randrange(len(words) + 1), len(words) - 1)
            match rng.choice(["insert", "delete", "replace"]):
                case "insert":
                    words.insert(i + 1, rng.choice(WORDS) + rng.choice([" ", ""]))
                case "delete" if len(words) > 1:
                    words.pop(i)
                case _:
                    words[i] = rng.choice(WORDS)
        new = "".join(words)
        try:
            q = splice(p, a, full[a:b], new)
        except T.SpliceError:
            continue
        total += 1
        if T.paragraph_text(q) != full[:a] + new + full[b:]:
            wrong.append(f"{full[a:b]!r} -> {new!r} wrote {T.paragraph_text(q)!r}")
    return total, wrong


def main() -> int:
    if not any(CACHE.glob("*.pptx")):
        print(f"SWEEP SKIPPED: no corpus decks in {CACHE}; run .claude/skills/verify-pptx/scripts/corpus.py to fetch them")
        return 0
    paras = paragraphs()
    rng = random.Random(13)
    results = {"boundaries": boundaries(paras), "same run": same_run(paras, rng), "rewrites": rewrites(paras, rng)}
    for name, (total, wrong) in results.items():
        print(f"{name}: {len(wrong)} wrong of {total}")
        for line in wrong[:5]:
            print(f"  {line}")
    failed = any(wrong for _, wrong in results.values())
    print("SWEEP FAIL" if failed else "SWEEP PASS")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())

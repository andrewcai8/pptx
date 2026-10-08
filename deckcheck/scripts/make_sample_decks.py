"""Build sample decks for deckcheck.

Usage: uv run --project deckcheck python deckcheck/scripts/make_sample_decks.py OUT_DIR

Writes clean.pptx (passes the house style), dirty.pptx (one violation of each
rule), and clean-v2.pptx (clean with slide 2 edited and slide 5 appended), then
prints the expected (slide, rule) violations for dirty.pptx.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Inches, Pt

TITLE_SLIDE, TITLE_AND_CONTENT, TITLE_ONLY = 0, 1, 5

EXPECTED_DIRTY: list[tuple[int, str]] = [
    (1, "no-placeholder-text"),
    (2, "max-fonts-per-slide"),
    (2, "no-bullet-end-punctuation"),
    (3, "source-on-data-slides"),
    (3, "title-max-chars"),
    (4, "min-font-size"),
    (4, "slide-has-title"),
    (4, "within-slide-bounds"),
]


def build_clean(path: Path) -> None:
    _build(path, dirty=False, v2=False)


def build_dirty(path: Path) -> None:
    _build(path, dirty=True, v2=False)


def build_clean_v2(path: Path) -> None:
    _build(path, dirty=False, v2=True)


def _build(path: Path, *, dirty: bool, v2: bool) -> None:
    prs = Presentation()

    s1 = prs.slides.add_slide(prs.slide_layouts[TITLE_SLIDE])
    s1.shapes.title.text = "Project Atlas market entry review"
    s1.placeholders[1].text = "Prepared by XX" if dirty else "Prepared for the steering committee, October 2026"

    s2 = prs.slides.add_slide(prs.slide_layouts[TITLE_AND_CONTENT])
    s2.shapes.title.text = "Demand is shifting to mid-market buyers"
    bullets = [
        "Mid-market revenue grew 18% year over year",
        "Enterprise renewals flat at 95%" if v2 else "Enterprise renewals flat at 94%",
        "Three competitors exited the segment in 2025" + ("." if dirty else ""),
    ]
    _bullets(s2.placeholders[1].text_frame, bullets, fonts=["Arial", "Georgia", "Verdana"] if dirty else None)

    s3 = prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
    s3.shapes.title.text = (
        "Mid-market share doubled since 2022 as enterprise buyers consolidated vendors and three "
        "regional competitors exited the segment entirely"
        if dirty
        else "Mid-market share doubled since 2022"
    )
    data = CategoryChartData()
    data.categories = ["2022", "2023", "2024", "2025"]
    data.add_series("Mid-market share %", (12, 16, 20, 24))
    s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(9), Inches(4.5), data)
    if not dirty:
        _text_box(s3, "Source: Company filings, 2022-2025", Pt(12))

    s4 = prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
    if not dirty:
        s4.shapes.title.text = "Mid-market carries the best unit economics"
    left = Inches(2) if dirty else Inches(0.5)
    table = s4.shapes.add_table(3, 3, left, Inches(1.5), Inches(9), Inches(2)).table
    for r, row in enumerate([["Segment", "Revenue", "Margin"], ["Mid-market", "$42M", "31%"], ["Enterprise", "$88M", "22%"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    _text_box(s4, "Source: Internal finance data, FY2025", Pt(8) if dirty else Pt(12))

    if v2:
        s5 = prs.slides.add_slide(prs.slide_layouts[TITLE_AND_CONTENT])
        s5.shapes.title.text = "Next steps"
        _bullets(s5.placeholders[1].text_frame, ["Validate pricing with five pilot customers", "Pick a channel partner by Q1"])

    prs.save(str(path))


def _bullets(frame, texts: list[str], fonts: list[str] | None = None) -> None:
    for i, text in enumerate(texts):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = p.add_run()
        run.text = text
        if fonts:
            run.font.name = fonts[i]


def _text_box(slide, text: str, size: Emu) -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(6.6), Inches(9), Inches(0.4))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = size


def main(argv: list[str]) -> int:
    if len(argv) != 1:
        print(__doc__, file=sys.stderr)
        return 2
    out = Path(argv[0])
    out.mkdir(parents=True, exist_ok=True)
    build_clean(out / "clean.pptx")
    build_dirty(out / "dirty.pptx")
    build_clean_v2(out / "clean-v2.pptx")
    print(f"wrote {out / 'clean.pptx'}, {out / 'dirty.pptx'}, {out / 'clean-v2.pptx'}")
    print("expected dirty.pptx violations (slide, rule):")
    for slide, rule in EXPECTED_DIRTY:
        print(f"{slide} {rule}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))

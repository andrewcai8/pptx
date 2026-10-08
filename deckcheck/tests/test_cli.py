from __future__ import annotations

import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

from deckcheck.cli import main
from make_sample_decks import build_clean, build_clean_v2, build_dirty

HOUSE_STYLE = Path(__file__).resolve().parents[2] / "standards" / "house-style.yaml"


def run(argv: list[str], capsys: pytest.CaptureFixture[str]) -> tuple[int, str, str]:
    code = main(argv)
    out, err = capsys.readouterr()
    return code, out, err


def write_rules(tmp_path: Path, body: str) -> Path:
    path = tmp_path / "rules.yaml"
    path.write_text(body)
    return path


def violations(deck: Path, rules: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> tuple[int, list]:
    code, _, _ = run(["check", str(deck), "--rules", str(rules), "--out", str(tmp_path / "report")], capsys)
    report = json.loads((tmp_path / "report" / "report.json").read_text())
    return code, [(v["slide"], v["rule"]) for v in report["violations"]]


BOUNDS_RULES = "rules:\n  within-slide-bounds:\n    tolerance_pt: 18\n"
BLANK, TITLE_ONLY = 6, 5


def deck_of_boxes(path: Path, boxes: list[tuple[Emu, Emu, Emu, Emu, str]]) -> Path:
    prs = Presentation()
    for left, top, width, height, text in boxes:
        slide = prs.slides.add_slide(prs.slide_layouts[BLANK])
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.text_frame.text = text
    prs.save(str(path))
    return path


def test_within_slide_bounds_ignores_shapes_without_text(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck = deck_of_boxes(
        tmp_path / "bleed.pptx",
        [
            (Inches(-1), Inches(-1), Inches(12), Inches(9.5), ""),
            (Inches(-1), Inches(-1), Inches(12), Inches(9.5), "Full-bleed caption"),
        ],
    )

    assert violations(deck, write_rules(tmp_path, BOUNDS_RULES), tmp_path, capsys) == (
        1,
        [(2, "within-slide-bounds")],
    )


def test_within_slide_bounds_allows_overhang_up_to_tolerance(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    deck = deck_of_boxes(
        tmp_path / "overhang.pptx",
        [
            (Inches(6) + Pt(18), Inches(1), Inches(4), Inches(1), "Wide text box"),
            (Inches(6) + Pt(19), Inches(1), Inches(4), Inches(1), "Wide text box"),
        ],
    )

    assert violations(deck, write_rules(tmp_path, BOUNDS_RULES), tmp_path, capsys) == (
        1,
        [(2, "within-slide-bounds")],
    )


def test_within_slide_bounds_requires_tolerance(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck = tmp_path / "clean.pptx"
    build_clean(deck)
    rules = write_rules(tmp_path, "rules:\n  within-slide-bounds: {}\n")

    assert run(["check", str(deck), "--rules", str(rules)], capsys) == (
        2,
        "",
        f"error: {rules}: within-slide-bounds: missing params tolerance_pt\n",
    )


def test_slide_has_title_exempts_layouts_without_a_title_placeholder(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[BLANK])
    prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
    deck = tmp_path / "titles.pptx"
    prs.save(str(deck))
    rules = write_rules(tmp_path, "rules:\n  slide-has-title:\n    exempt_layouts: []\n")

    assert violations(deck, rules, tmp_path, capsys) == (1, [(2, "slide-has-title")])


def test_clean_deck_passes_house_style(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck = tmp_path / "clean.pptx"
    build_clean(deck)

    assert run(["check", str(deck), "--rules", str(HOUSE_STYLE)], capsys) == (
        0,
        f"PASS {deck} (4 slides, 8 rules)\n",
        "",
    )


def test_dirty_deck_reports_one_violation_per_rule(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck = tmp_path / "dirty.pptx"
    build_dirty(deck)

    code, out, _ = run(["check", str(deck), "--rules", str(HOUSE_STYLE), "--out", str(tmp_path / "report")], capsys)

    report = json.loads((tmp_path / "report" / "report.json").read_text())
    assert code == 1
    assert out.splitlines()[0] == f"FAIL {deck}: 8 violations"
    assert report["passed"] is False
    assert sorted((v["slide"], v["rule"]) for v in report["violations"]) == [
        (1, "no-placeholder-text"),
        (2, "max-fonts-per-slide"),
        (2, "no-bullet-end-punctuation"),
        (3, "source-on-data-slides"),
        (3, "title-max-chars"),
        (4, "min-font-size"),
        (4, "slide-has-title"),
        (4, "within-slide-bounds"),
    ]


def test_diff_reports_edited_and_appended_slides(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    old, new = tmp_path / "clean.pptx", tmp_path / "clean-v2.pptx"
    build_clean(old)
    build_clean_v2(new)

    code, _, _ = run(["diff", str(old), str(new), "--out", str(tmp_path / "diff")], capsys)

    result = json.loads((tmp_path / "diff" / "diff.json").read_text())
    assert code == 0
    assert [(s["slide"], s["status"]) for s in result["slides"]] == [
        (1, "unchanged"),
        (2, "changed"),
        (3, "unchanged"),
        (4, "unchanged"),
        (5, "added"),
    ]
    changed = result["slides"][1]["diff"].splitlines()
    assert "-Content Placeholder 2: Enterprise renewals flat at 94%" in changed
    assert "+Content Placeholder 2: Enterprise renewals flat at 95%" in changed


def test_unknown_rule_id_is_a_config_error(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck = tmp_path / "clean.pptx"
    build_clean(deck)
    rules = write_rules(tmp_path, "rules:\n  no-comic-sans: {}\n")

    assert run(["check", str(deck), "--rules", str(rules)], capsys) == (
        2,
        "",
        f"error: {rules}: unknown rule id 'no-comic-sans'\n",
    )


def test_theme_fallback_and_theme_refs_resolve_to_theme_fonts(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Fonts"
    para = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text_frame.paragraphs[0]
    for font in ["Arial", "Georgia", "Verdana", "+mn-lt"]:
        run_ = para.add_run()
        run_.text = f"{font} "
        run_.font.name = font
    deck = tmp_path / "fonts.pptx"
    prs.save(str(deck))
    rules = write_rules(tmp_path, "rules:\n  max-fonts-per-slide:\n    max: 3\n")

    assert run(["check", str(deck), "--rules", str(rules)], capsys) == (
        1,
        f"FAIL {deck}: 1 violations\n"
        "slide 1 [max-fonts-per-slide] 4 fonts on slide, max 3 | Arial, Calibri, Georgia, Verdana\n",
        "",
    )


def test_bu_none_paragraph_in_body_placeholder_is_not_a_bullet(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Bullets"
    frame = slide.placeholders[1].text_frame
    frame.paragraphs[0].text = "A prose sentence."
    frame.paragraphs[0]._p.get_or_add_pPr().insert(0, OxmlElement("a:buNone"))
    frame.add_paragraph().text = "An inherited bullet."
    deck = tmp_path / "bullets.pptx"
    prs.save(str(deck))
    rules = write_rules(tmp_path, 'rules:\n  no-bullet-end-punctuation:\n    chars: "."\n')

    assert run(["check", str(deck), "--rules", str(rules)], capsys) == (
        1,
        f"FAIL {deck}: 1 violations\n"
        "slide 1 [no-bullet-end-punctuation] Content Placeholder 2: bullet ends with '.' | An inherited bullet.\n",
        "",
    )


@pytest.mark.parametrize(("char", "expected"), [("\u200b", (0, [])), ("\u2022", (1, [(1, "no-bullet-end-punctuation")]))])
def test_body_paragraph_inherits_its_bullet_from_the_master(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], char: str, expected: tuple[int, list]
) -> None:
    prs = Presentation()
    body_style = prs.slide_master.element.find(qn("p:txStyles")).find(qn("p:bodyStyle"))
    body_style.find(qn("a:lvl1pPr")).find(qn("a:buChar")).set("char", char)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Bullets"
    slide.placeholders[1].text_frame.text = "A prose sentence."
    deck = tmp_path / "bullets.pptx"
    prs.save(str(deck))
    rules = write_rules(tmp_path, 'rules:\n  no-bullet-end-punctuation:\n    chars: "."\n')

    assert violations(deck, rules, tmp_path, capsys) == expected


SOURCE_RULES = "rules:\n  source-on-data-slides:\n    prefix: Source\n"


def add_chart(slide) -> None:
    data = CategoryChartData()
    data.categories = ["2024", "2025"]
    data.add_series("Share", (12, 16))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(9), Inches(4.5), data)


def add_table(slide, cells: list[str]) -> None:
    table = slide.shapes.add_table(1, len(cells), Inches(0.5), Inches(1.5), Inches(9), Inches(1)).table
    for c, value in enumerate(cells):
        table.cell(0, c).text = value


def add_footnote(slide, text: str) -> None:
    slide.shapes.add_textbox(Inches(0.5), Inches(6.6), Inches(9), Inches(0.4)).text_frame.text = text


def test_source_may_follow_a_note_on_a_later_line(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    prs = Presentation()
    for footnote in ["Note: Shares are rounded\vSource: Company filings", "Note: Shares are rounded"]:
        slide = prs.slides.add_slide(prs.slide_layouts[BLANK])
        add_chart(slide)
        add_footnote(slide, footnote)
    deck = tmp_path / "footnotes.pptx"
    prs.save(str(deck))

    assert violations(deck, write_rules(tmp_path, SOURCE_RULES), tmp_path, capsys) == (
        1,
        [(2, "source-on-data-slides")],
    )


def test_only_tables_with_numbers_need_a_source(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    prs = Presentation()
    for cells in [["Stores", "Offices", "Warehouses"], ["Stores", "Offices", "42%"]]:
        add_table(prs.slides.add_slide(prs.slide_layouts[BLANK]), cells)
    deck = tmp_path / "tables.pptx"
    prs.save(str(deck))

    assert violations(deck, write_rules(tmp_path, SOURCE_RULES), tmp_path, capsys) == (
        1,
        [(2, "source-on-data-slides")],
    )


def test_title_length_counts_only_the_headline(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    prs = Presentation()
    for title in ["Short headline\v" + "s" * 160, "h" * 151]:
        prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY]).shapes.title.text = title
    deck = tmp_path / "titles.pptx"
    prs.save(str(deck))
    rules = write_rules(tmp_path, "rules:\n  title-max-chars:\n    max: 150\n")

    assert violations(deck, rules, tmp_path, capsys) == (1, [(2, "title-max-chars")])


def deck_of_text_boxes(path: Path, boxes: list[tuple[Emu, Emu, Emu, Emu, float]]) -> Path:
    prs = Presentation()
    for left, top, width, height, rotation in boxes:
        box = prs.slides.add_slide(prs.slide_layouts[BLANK]).shapes.add_textbox(left, top, width, height)
        box.text_frame.text = "Share of stores open"
        box.rotation = rotation
    prs.save(str(path))
    return path


def test_within_slide_bounds_measures_a_rotated_box(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck = deck_of_text_boxes(
        tmp_path / "axis.pptx",
        [
            (Inches(-1.2), Inches(3), Inches(3), Inches(0.5), 270.0),
            (Inches(-1.2), Inches(3), Inches(3), Inches(0.5), 0.0),
        ],
    )

    assert violations(deck, write_rules(tmp_path, BOUNDS_RULES), tmp_path, capsys) == (
        1,
        [(2, "within-slide-bounds")],
    )


def test_within_slide_bounds_skips_a_box_parked_off_the_slide(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    deck = deck_of_text_boxes(
        tmp_path / "parked.pptx",
        [
            (Inches(0), Inches(-1), Inches(1.4), Inches(0.5), 0.0),
            (Inches(0), Inches(-0.3), Inches(1.4), Inches(0.5), 0.0),
        ],
    )

    assert violations(deck, write_rules(tmp_path, BOUNDS_RULES), tmp_path, capsys) == (
        1,
        [(2, "within-slide-bounds")],
    )


def test_slide_has_title_exempts_a_copied_layout(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    prs = Presentation()
    copied = prs.slide_layouts[0]
    copied.element.cSld.set("name", "1_Title Slide")
    prs.slides.add_slide(copied)
    prs.slides.add_slide(prs.slide_layouts[1])
    deck = tmp_path / "copied.pptx"
    prs.save(str(deck))
    rules = write_rules(tmp_path, 'rules:\n  slide-has-title:\n    exempt_layouts: ["Title Slide"]\n')

    assert violations(deck, rules, tmp_path, capsys) == (1, [(2, "slide-has-title")])

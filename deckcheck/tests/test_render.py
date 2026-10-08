from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from deckcheck.cli import main
from deckcheck.render import find_soffice

BLANK = 6


def run(argv: list[str], capsys: pytest.CaptureFixture[str]) -> tuple[int, str, str]:
    code = main(argv)
    out, err = capsys.readouterr()
    return code, out, err


def deck_in_font(path: Path, font: str) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "Revenue grew 12% in 2024"
    box.text_frame.paragraphs[0].runs[0].font.name = font
    prs.save(str(path))
    return path


@pytest.mark.skipif(
    not (find_soffice() and shutil.which("pdftoppm") and shutil.which("fc-match")),
    reason="needs soffice, pdftoppm, and fc-match",
)
def test_render_reports_a_substituted_font(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck = deck_in_font(tmp_path / "deck.pptx", "Nonexistent Sans QA")
    out = tmp_path / "render"

    code, stdout, stderr = run(["render", str(deck), "--out", str(out)], capsys)

    fonts = json.loads((out / "fonts.json").read_text())
    family = fonts["Nonexistent Sans QA"]["family"]
    assert (code, stderr) == (0, "")
    assert fonts == {"Nonexistent Sans QA": {"family": family, "substituted": True}}
    assert stdout == f"{out}/slide-1.png\nsubstituted: Nonexistent Sans QA -> {family}\n"


def test_render_without_fc_match_exits_3_before_writing(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    deck = deck_in_font(tmp_path / "deck.pptx", "Nonexistent Sans QA")
    out = tmp_path / "render"
    empty = tmp_path / "empty"
    empty.mkdir()
    monkeypatch.setenv("PATH", str(empty))

    assert run(["render", str(deck), "--out", str(out)], capsys) == (
        3,
        "",
        "error: fc-match not found; install fontconfig with `brew install fontconfig`\n",
    )
    assert not out.exists()

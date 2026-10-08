from __future__ import annotations

import copy
import hashlib
import json
import os
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

from deckcheck import model
from deckcheck.cli import main
from deckcheck.fix import FIXERS, Change, fix_bullet_end
from deckcheck.rules import RULES
from make_sample_decks import build_clean, build_dirty

HOUSE_STYLE = Path(__file__).resolve().parents[2] / "standards" / "house-style.yaml"
BULLET_RULES = 'rules:\n  no-bullet-end-punctuation:\n    chars: ".;,"\n'
BOUNDS_RULES = "rules:\n  within-slide-bounds:\n    tolerance_pt: 18\n"
TITLE_AND_CONTENT, TITLE_ONLY, BLANK = 1, 5, 6


def run(argv: list[str], capsys: pytest.CaptureFixture[str]) -> tuple[int, str, str]:
    code = main(argv)
    out, err = capsys.readouterr()
    return code, out, err


def write_rules(tmp_path: Path, body: str) -> Path:
    path = tmp_path / "rules.yaml"
    path.write_text(body)
    return path


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def check_pairs(deck: Path, rules: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> list:
    run(["check", str(deck), "--rules", str(rules), "--out", str(tmp_path / "check")], capsys)
    report = json.loads((tmp_path / "check" / "report.json").read_text())
    return [(v["slide"], v["rule"]) for v in report["violations"]]


def box(shape) -> tuple[int, int, int, int]:
    return shape.left, shape.top, shape.width, shape.height


REPORT_ONLY_DIRTY = [
    (1, "no-placeholder-text"),
    (3, "source-on-data-slides"),
    (3, "title-max-chars"),
    (4, "slide-has-title"),
]


def test_fix_dirty_deck_fixes_four_and_reports_four(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck, fixed = tmp_path / "dirty.pptx", tmp_path / "out" / "fixed.pptx"
    build_dirty(deck)
    before = sha256(deck)

    code, out, err = run(
        ["fix", str(deck), "--out", str(fixed), "--rules", str(HOUSE_STYLE), "--report", str(tmp_path / "fix")], capsys
    )

    assert (code, err) == (1, "")
    assert out == (
        f"FAIL {deck} -> {fixed}: 4 fixed in 1 pass, 4 remain\n"
        "fixed slide 2 [max-fonts-per-slide] 4 fonts on slide, max 3 | font Georgia -> +mn-lt (Calibri), 1 run\n"
        "fixed slide 2 [no-bullet-end-punctuation] Content Placeholder 2: bullet ends with '.' | "
        "text 'Three competitors exited the segment in 2025.' -> 'Three competitors exited the segment in 2025'\n"
        "fixed slide 4 [min-font-size] TextBox 3: 6pt text, min 7pt | size 6pt -> 7pt, 1 run\n"
        "fixed slide 4 [within-slide-bounds] Table 2: text shape extends 72pt past the slide edge, tolerance 18pt | "
        "box (1828800, 1371600) 8229600x1828800 EMU -> (914400, 1371600) 8229600x1828800 EMU\n"
        "remains slide 1 [no-placeholder-text] Subtitle 2: placeholder text 'XX' | Prepared by XX "
        "(report-only: fix by hand)\n"
        "remains slide 3 [source-on-data-slides] Chart 2: chart slide has no line starting with 'Source' | Chart 2 "
        "(report-only: fix by hand)\n"
        "remains slide 3 [title-max-chars] Title 1: title is 167 chars, max 150 | Mid-market share doubled since "
        "2022 as enterprise buyers consolidated vendors and three regional competitors exited the segment entirely, "
        "while list prices held steady (report-only: fix by hand)\n"
        "remains slide 4 [slide-has-title] slide has no title | layout Title Only (report-only: fix by hand)\n"
    )
    assert sha256(deck) == before
    report = json.loads((tmp_path / "fix" / "fix.json").read_text())
    assert (report["input_sha256"], report["output_sha256"], report["passes"], report["passed"]) == (
        before,
        sha256(fixed),
        1,
        False,
    )
    assert [(f["slide"], f["rule"], f["pass"]) for f in report["fixed"]] == [
        (2, "max-fonts-per-slide", 1),
        (2, "no-bullet-end-punctuation", 1),
        (4, "min-font-size", 1),
        (4, "within-slide-bounds", 1),
    ]
    assert [(r["slide"], r["rule"], r["why"]) for r in report["remaining"]] == [
        (s, r, "report-only") for s, r in REPORT_ONLY_DIRTY
    ]
    assert check_pairs(fixed, HOUSE_STYLE, tmp_path, capsys) == REPORT_ONLY_DIRTY

    run(["diff", str(deck), str(fixed), "--out", str(tmp_path / "diff")], capsys)
    diff = json.loads((tmp_path / "diff" / "diff.json").read_text())
    assert [(s["slide"], s["status"]) for s in diff["slides"]] == [
        (1, "unchanged"),
        (2, "changed"),
        (3, "unchanged"),
        (4, "unchanged"),
    ]
    assert [line for line in diff["slides"][1]["diff"].splitlines() if line.startswith(("+", "-"))] == [
        "--- old/slide-2",
        "+++ new/slide-2",
        "-Content Placeholder 2: Three competitors exited the segment in 2025.",
        "+Content Placeholder 2: Three competitors exited the segment in 2025",
    ]


def test_fix_writes_the_fixed_values_into_the_dirty_deck(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck, fixed = tmp_path / "dirty.pptx", tmp_path / "fixed.pptx"
    build_dirty(deck)

    run(["fix", str(deck), "--out", str(fixed), "--rules", str(HOUSE_STYLE)], capsys)

    slides = Presentation(str(fixed)).slides
    bullets = slides[1].placeholders[1].text_frame.paragraphs
    assert [(p.runs[0].text, p.runs[0].font.name) for p in bullets] == [
        ("Mid-market revenue grew 18% year over year", "Arial"),
        ("Enterprise renewals flat at 94%", "+mn-lt"),
        ("Three competitors exited the segment in 2025", "Verdana"),
    ]
    table, source = slides[3].shapes[1], slides[3].shapes[2]
    assert box(table) == (Inches(1), Inches(1.5), Inches(9), Inches(2))
    assert source.text_frame.paragraphs[0].runs[0].font.size == Pt(7)


def test_fixing_a_fixed_deck_changes_nothing(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck, first, second = tmp_path / "dirty.pptx", tmp_path / "first.pptx", tmp_path / "second.pptx"
    build_dirty(deck)
    run(["fix", str(deck), "--out", str(first), "--rules", str(HOUSE_STYLE)], capsys)

    code, out, _ = run(["fix", str(first), "--out", str(second), "--rules", str(HOUSE_STYLE)], capsys)

    assert (code, out.splitlines()[0]) == (1, f"FAIL {first} -> {second}: 0 fixed, 4 remain")
    assert second.read_bytes() == first.read_bytes()


def test_fix_copies_a_clean_deck_byte_for_byte(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck, fixed = tmp_path / "clean.pptx", tmp_path / "fixed.pptx"
    build_clean(deck)

    assert run(["fix", str(deck), "--out", str(fixed), "--rules", str(HOUSE_STYLE)], capsys) == (
        0,
        f"PASS {deck} -> {fixed}: 0 fixed, 0 remain\n",
        "",
    )
    assert fixed.read_bytes() == deck.read_bytes()


def zip_entries(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as z:
        return {name: z.read(name) for name in z.namelist()}


def test_fix_copies_every_part_it_did_not_write_byte_for_byte(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    deck, fixed = tmp_path / "dirty.pptx", tmp_path / "fixed.pptx"
    build_dirty(deck)
    with zipfile.ZipFile(deck, "a") as z:
        z.writestr("customXml/unreferenced.xml", "<note>python-pptx drops parts no relationship reaches</note>")
    before = zip_entries(deck)

    run(["fix", str(deck), "--out", str(fixed), "--rules", str(HOUSE_STYLE)], capsys)

    after = zip_entries(fixed)
    assert list(after) == list(before)
    assert [name for name in before if after[name] != before[name]] == [
        "ppt/slides/slide2.xml",
        "ppt/slides/slide4.xml",
    ]


def test_fix_writes_nothing_when_reading_the_deck_changes_it(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    deck, fixed = tmp_path / "dirty.pptx", tmp_path / "fixed.pptx"
    build_dirty(deck)
    read_run = model._run

    def writing_read(r, theme, is_title):
        r.get_or_add_rPr()
        return read_run(r, theme, is_title)

    monkeypatch.setattr(model, "_run", writing_read)

    assert run(["fix", str(deck), "--out", str(fixed), "--rules", str(HOUSE_STYLE)], capsys) == (
        2,
        "",
        f"error: reading {deck} changed slide 1, so fix cannot tell its own edits apart; nothing written\n",
    )
    assert not fixed.exists()


@pytest.mark.parametrize("link", ["same", "dotted", "symlink", "hardlink"])
def test_fix_refuses_to_overwrite_its_input(tmp_path: Path, capsys: pytest.CaptureFixture[str], link: str) -> None:
    deck = tmp_path / "decks" / "dirty.pptx"
    deck.parent.mkdir()
    build_dirty(deck)
    before = sha256(deck)
    out = {
        "same": deck,
        "dotted": tmp_path / "decks" / ".." / "decks" / "dirty.pptx",
        "symlink": tmp_path / "link.pptx",
        "hardlink": tmp_path / "hard.pptx",
    }[link]
    if link == "symlink":
        out.symlink_to(deck)
    if link == "hardlink":
        os.link(deck, out)

    assert run(["fix", str(deck), "--out", str(out), "--rules", str(HOUSE_STYLE)], capsys) == (
        2,
        "",
        f"error: --out {out} is the input deck; fix never overwrites its input\n",
    )
    assert sha256(deck) == before


def test_fix_refuses_a_directory_as_out(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck = tmp_path / "dirty.pptx"
    build_dirty(deck)

    assert run(["fix", str(deck), "--out", str(tmp_path), "--rules", str(HOUSE_STYLE)], capsys) == (
        2,
        "",
        f"error: --out {tmp_path} is a directory; pass the path of the new deck\n",
    )


def test_fix_exits_2_when_it_cannot_write_its_output(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck, blocker = tmp_path / "dirty.pptx", tmp_path / "afile"
    build_dirty(deck)
    blocker.write_text("not a directory")

    assert run(["fix", str(deck), "--out", str(blocker / "x.pptx"), "--rules", str(HOUSE_STYLE)], capsys) == (
        2,
        "",
        f"error: cannot write the fix output: [Errno 17] File exists: '{blocker}'\n",
    )


def test_fix_checks_the_report_directory_before_writing_the_deck(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    deck, fixed, blocker = tmp_path / "dirty.pptx", tmp_path / "fixed.pptx", tmp_path / "afile"
    build_dirty(deck)
    blocker.write_text("not a directory")

    assert run(
        ["fix", str(deck), "--out", str(fixed), "--rules", str(HOUSE_STYLE), "--report", str(blocker)], capsys
    ) == (2, "", f"error: cannot write the fix output: [Errno 17] File exists: '{blocker}'\n")
    assert not fixed.exists()


def test_fix_output_mode_follows_the_umask(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck, fixed = tmp_path / "dirty.pptx", tmp_path / "fixed.pptx"
    build_dirty(deck)
    deck.chmod(0o400)
    old = os.umask(0o027)
    try:
        run(["fix", str(deck), "--out", str(fixed), "--rules", str(HOUSE_STYLE)], capsys)
    finally:
        os.umask(old)

    assert oct(fixed.stat().st_mode & 0o777) == "0o640"


def bullet_deck(path: Path, texts: list[str]) -> Path:
    prs = Presentation()
    for text in texts:
        slide = prs.slides.add_slide(prs.slide_layouts[TITLE_AND_CONTENT])
        slide.shapes.title.text = "Bullets"
        slide.placeholders[1].text_frame.text = text
    prs.save(str(path))
    return path


def bullet_texts(path: Path) -> list[str]:
    return [s.placeholders[1].text_frame.text for s in Presentation(str(path)).slides]


@pytest.mark.parametrize(
    ("text", "stripped"),
    [
        ("Grew 18% year over year.", "Grew 18% year over year"),
        ("Costs fell;", "Costs fell"),
        ("Margins held.;", "Margins held"),
        ("Rates of 3.5.", "Rates of 3.5"),
        ("One group per location/lab/classroom.", "One group per location/lab/classroom"),
        ("Pursue M&A.", "Pursue M&A"),
        ("Strengthen R&D.", "Strengthen R&D"),
        ("Close with a Q&A.", "Close with a Q&A"),
        ("Margins won't.", "Margins won't"),
        ("Report to the CEO's.", "Report to the CEO's"),
        ("Report to the CEO’s.", "Report to the CEO’s"),
    ],
)
def test_bullet_fix_strips_sentence_endings(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], text: str, stripped: str
) -> None:
    deck, fixed = bullet_deck(tmp_path / "bullets.pptx", [text]), tmp_path / "fixed.pptx"

    code, _, _ = run(
        ["fix", str(deck), "--out", str(fixed), "--rules", str(write_rules(tmp_path, BULLET_RULES))], capsys
    )

    assert (code, bullet_texts(fixed)) == (0, [stripped])


@pytest.mark.parametrize(
    ("text", "detail"),
    [
        ("Sanitize carts, baskets, etc.", "abbreviation etc."),
        ("Train store mgmt.", "abbreviation mgmt."),
        ("Brief the sales reps.", "abbreviation reps."),
        ("Add a cashier asst.", "abbreviation asst."),
        ("Etc.", "abbreviation Etc."),
        ("Cut fixed costs, e.g.", "abbreviation e.g."),
        ("Stores in the U.S.", "abbreviation U.S."),
        ("Pick Plan A.", "abbreviation A."),
        ("More to come..", "ellipsis"),
        ("More to come...", "ellipsis"),
        ("More to come....", "ellipsis"),
        ("More to come….", "ellipsis"),
        (".", "the bullet is only punctuation"),
        ("Revenue grew. .", "punctuation separated by spaces"),
    ],
)
def test_bullet_fix_reports_endings_it_must_not_strip(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], text: str, detail: str
) -> None:
    deck, fixed = bullet_deck(tmp_path / "bullets.pptx", [text]), tmp_path / "fixed.pptx"

    code, out, _ = run(
        ["fix", str(deck), "--out", str(fixed), "--rules", str(write_rules(tmp_path, BULLET_RULES))], capsys
    )

    assert (code, out.splitlines()[1:]) == (
        1,
        [
            (
                f"remains slide 1 [no-bullet-end-punctuation] Content Placeholder 2: bullet ends with {text[-1]!r} | "
                f"{text} (declined: {detail})"
            )
        ],
    )
    assert bullet_texts(fixed) == [text]


def test_bullet_fix_strips_across_runs_and_a_line_break(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[TITLE_AND_CONTENT])
    slide.shapes.title.text = "Bullets"
    p = slide.placeholders[1].text_frame.paragraphs[0]
    for text in ["Revenue grew 18%", ".", ";"]:
        p.add_run().text = text
    p.add_line_break()
    p.add_run().text = "  "
    deck, fixed = tmp_path / "runs.pptx", tmp_path / "fixed.pptx"
    prs.save(str(deck))

    code, _, _ = run(
        ["fix", str(deck), "--out", str(fixed), "--rules", str(write_rules(tmp_path, BULLET_RULES))], capsys
    )

    p = Presentation(str(fixed)).slides[0].placeholders[1].text_frame.paragraphs[0]
    assert (code, [r.text for r in p.runs]) == (0, ["Revenue grew 18%", "", "", "  "])


def test_bullet_fix_declines_a_bullet_ending_in_a_field(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[TITLE_AND_CONTENT])
    slide.shapes.title.text = "Bullets"
    p = slide.placeholders[1].text_frame.paragraphs[0]
    p.add_run().text = "See slide "
    fld = OxmlElement("a:fld")
    fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
    fld.set("type", "slidenum")
    t = OxmlElement("a:t")
    t.text = "4."
    fld.append(t)
    p._p.append(fld)
    deck, fixed = tmp_path / "field.pptx", tmp_path / "fixed.pptx"
    prs.save(str(deck))

    code, out, _ = run(
        ["fix", str(deck), "--out", str(fixed), "--rules", str(write_rules(tmp_path, BULLET_RULES))], capsys
    )

    assert (code, out.splitlines()[1:]) == (
        1,
        [
            (
                "remains slide 1 [no-bullet-end-punctuation] Content Placeholder 2: bullet ends with '.' | "
                "See slide 4. (declined: ends in a field)"
            )
        ],
    )
    assert fixed.read_bytes() == deck.read_bytes()


def test_font_fix_keeps_theme_and_symbol_fonts(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
    slide.shapes.title.text = "Fonts"
    para = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text_frame.paragraphs[0]
    for font, text in [("Arial", "Most of the words "), ("Georgia", "a few "), ("Wingdings", "\uf0a7")]:
        r = para.add_run()
        r.text, r.font.name = text, font
    deck, fixed = tmp_path / "fonts.pptx", tmp_path / "fixed.pptx"
    prs.save(str(deck))
    rules = write_rules(tmp_path, "rules:\n  max-fonts-per-slide:\n    max: 3\n")

    code, out, _ = run(["fix", str(deck), "--out", str(fixed), "--rules", str(rules)], capsys)

    assert (code, out.splitlines()[1:]) == (
        0,
        [
            "fixed slide 1 [max-fonts-per-slide] 4 fonts on slide, max 3 | font Georgia -> +mn-lt (Calibri), 1 run",
        ],
    )
    runs = Presentation(str(fixed)).slides[0].shapes[1].text_frame.paragraphs[0].runs
    assert [r.font.name for r in runs] == ["Arial", "+mn-lt", "Wingdings"]


def test_font_fix_gives_titles_the_major_theme_font(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
    title = slide.shapes.title.text_frame.paragraphs[0].add_run()
    title.text, title.font.name = "Fonts", "Impact"
    para = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text_frame.paragraphs[0]
    for font, text in [("Arial", "a "), ("Georgia", "bb "), ("Verdana", "Most of the words on the slide")]:
        r = para.add_run()
        r.text, r.font.name = text, font
    deck, fixed = tmp_path / "fonts.pptx", tmp_path / "fixed.pptx"
    prs.save(str(deck))
    rules = write_rules(tmp_path, "rules:\n  max-fonts-per-slide:\n    max: 2\n")

    code, out, _ = run(["fix", str(deck), "--out", str(fixed), "--rules", str(rules)], capsys)

    assert (code, out.splitlines()[1:]) == (
        0,
        [
            (
                "fixed slide 1 [max-fonts-per-slide] 4 fonts on slide, max 2 | "
                "font Arial -> +mn-lt (Calibri), 1 run; font Georgia -> +mn-lt (Calibri), 1 run; "
                "font Impact -> +mj-lt (Calibri), 1 run"
            ),
        ],
    )
    shapes = Presentation(str(fixed)).slides[0].shapes
    assert [r.font.name for r in shapes.title.text_frame.paragraphs[0].runs] == ["+mj-lt"]
    assert [r.font.name for r in shapes[1].text_frame.paragraphs[0].runs] == ["+mn-lt", "+mn-lt", "Verdana"]


def text_box(shapes, name: str, left: Emu, top: Emu, width: Emu, height: Emu, rotation: float = 0.0):
    shape = shapes.add_textbox(left, top, width, height)
    shape.name, shape.rotation = name, rotation
    shape.text_frame.text = "Share of stores open"
    return shape


def test_bounds_fix_moves_only_the_overhanging_shape_of_two_with_one_name(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK])
    text_box(slide.shapes, "ColumnHeader", Inches(-1), Inches(1), Inches(3), Inches(0.5))
    text_box(slide.shapes, "ColumnHeader", Inches(4), Inches(1), Inches(3), Inches(0.5))
    deck, fixed = tmp_path / "headers.pptx", tmp_path / "fixed.pptx"
    prs.save(str(deck))

    code, out, _ = run(
        ["fix", str(deck), "--out", str(fixed), "--rules", str(write_rules(tmp_path, BOUNDS_RULES))], capsys
    )

    assert (code, out.splitlines()[1:]) == (
        0,
        [
            (
                "fixed slide 1 [within-slide-bounds] ColumnHeader: text shape extends 72pt past the slide edge, "
                "tolerance 18pt | box (-914400, 914400) 2743200x457200 EMU -> (0, 914400) 2743200x457200 EMU"
            )
        ],
    )
    assert [box(s) for s in Presentation(str(fixed)).slides[0].shapes] == [
        (0, Inches(1), Inches(3), Inches(0.5)),
        (Inches(4), Inches(1), Inches(3), Inches(0.5)),
    ]


def scaled_group(slide, rot: int = 0):
    group = slide.shapes.add_group_shape()
    text_box(group.shapes, "Grouped label", Inches(2), Inches(1), Inches(2), Inches(1))
    xfrm = group._element.grpSpPr.find(qn("a:xfrm"))
    xfrm.off.x, xfrm.off.y, xfrm.ext.cx, xfrm.ext.cy = Inches(6), Inches(1), Inches(4), Inches(2)
    xfrm.chOff.x, xfrm.chOff.y, xfrm.chExt.cx, xfrm.chExt.cy = Inches(1), Inches(1), Inches(2), Inches(1)
    if rot:
        xfrm.set("rot", str(rot))
    return group


def test_bounds_fix_writes_a_group_child_in_group_coordinates(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    prs = Presentation()
    scaled_group(prs.slides.add_slide(prs.slide_layouts[BLANK]))
    deck, fixed = tmp_path / "group.pptx", tmp_path / "fixed.pptx"
    prs.save(str(deck))

    code, out, _ = run(
        ["fix", str(deck), "--out", str(fixed), "--rules", str(write_rules(tmp_path, BOUNDS_RULES))], capsys
    )

    assert (code, out.splitlines()[1:]) == (
        0,
        [
            (
                "fixed slide 1 [within-slide-bounds] Grouped label: text shape extends 144pt past the slide edge, "
                "tolerance 18pt | box (7315200, 914400) 3657600x1828800 EMU -> (5486400, 914400) 3657600x1828800 EMU"
            )
        ],
    )
    group = Presentation(str(fixed)).slides[0].shapes[0]
    assert box(group.shapes[0]) == (Inches(1), Inches(1), Inches(2), Inches(1))
    assert box(group) == (Inches(6), Inches(1), Inches(4), Inches(2))


def offset_group(slide, ext_cx: Emu, child_left: Emu):
    group = slide.shapes.add_group_shape()
    text_box(group.shapes, "Grouped label", child_left, Inches(1), Inches(1), Inches(1))
    xfrm = group._element.grpSpPr.find(qn("a:xfrm"))
    xfrm.off.x, xfrm.off.y, xfrm.ext.cx, xfrm.ext.cy = 2, Inches(1), ext_cx, Inches(1)
    xfrm.chOff.x, xfrm.chOff.y, xfrm.chExt.cx, xfrm.chExt.cy = 2, Inches(1), Inches(1), Inches(1)
    return group


def test_bounds_fix_rounds_a_group_child_onto_the_slide(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    prs = Presentation()
    offset_group(prs.slides.add_slide(prs.slide_layouts[BLANK]), Inches(3), Inches(-0.5))
    deck, fixed = tmp_path / "group.pptx", tmp_path / "fixed.pptx"
    prs.save(str(deck))

    code, _, _ = run(
        ["fix", str(deck), "--out", str(fixed), "--rules", str(write_rules(tmp_path, BOUNDS_RULES))], capsys
    )

    child = Presentation(str(fixed)).slides[0].shapes[0].shapes[0]
    assert (code, box(child)) == (0, (2, Inches(1), Inches(1), Inches(1)))


def test_bounds_fix_declines_a_child_of_a_group_with_zero_width(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    prs = Presentation()
    group = offset_group(prs.slides.add_slide(prs.slide_layouts[BLANK]), 0, Inches(1))
    group.shapes[0].top = Inches(7)
    deck, fixed = tmp_path / "group.pptx", tmp_path / "fixed.pptx"
    prs.save(str(deck))

    code, out, _ = run(
        ["fix", str(deck), "--out", str(fixed), "--rules", str(write_rules(tmp_path, BOUNDS_RULES))], capsys
    )

    assert (code, out.splitlines()[1:]) == (
        1,
        [
            (
                "remains slide 1 [within-slide-bounds] Grouped label: text shape extends 36pt past the slide edge, "
                "tolerance 18pt | box (2, 6400800, 2, 7315200) vs slide (9144000, 6858000) EMU "
                "(declined: inside a group with zero width or height)"
            )
        ],
    )


def test_bounds_fix_shrinks_only_a_box_wider_than_the_slide(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    prs = Presentation()
    for left, width, rotation in [(Inches(-1), Inches(12), 0.0), (Inches(3), Inches(9), 90.0)]:
        text_box(
            prs.slides.add_slide(prs.slide_layouts[BLANK]).shapes, "Wide", left, Inches(3), width, Inches(1), rotation
        )
    deck, fixed = tmp_path / "wide.pptx", tmp_path / "fixed.pptx"
    prs.save(str(deck))

    code, _, _ = run(
        ["fix", str(deck), "--out", str(fixed), "--rules", str(write_rules(tmp_path, BOUNDS_RULES))], capsys
    )

    assert code == 0
    assert [box(s.shapes[0]) for s in Presentation(str(fixed)).slides] == [
        (0, Inches(3), Inches(10), Inches(1)),
        (Inches(3.75), Inches(3.25), Inches(7.5), Inches(1)),
    ]


def test_bounds_fix_declines_oversized_tables_and_rotated_groups(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    prs = Presentation()
    table = prs.slides.add_slide(prs.slide_layouts[BLANK]).shapes.add_table(
        1, 2, Inches(-1), Inches(1), Inches(12), Inches(1)
    )
    table.table.cell(0, 0).text = "Segment"
    scaled_group(prs.slides.add_slide(prs.slide_layouts[BLANK]), rot=5400000)
    deck, fixed = tmp_path / "declined.pptx", tmp_path / "fixed.pptx"
    prs.save(str(deck))

    code, out, _ = run(
        ["fix", str(deck), "--out", str(fixed), "--rules", str(write_rules(tmp_path, BOUNDS_RULES))], capsys
    )

    assert (code, out.splitlines()[1:]) == (
        1,
        [
            (
                "remains slide 1 [within-slide-bounds] Table 1: text shape extends 72pt past the slide edge, "
                "tolerance 18pt | box (-914400, 914400, 10058400, 1828800) vs slide (9144000, 6858000) EMU "
                "(declined: table is larger than the slide)"
            ),
            (
                "remains slide 2 [within-slide-bounds] Grouped label: text shape extends 144pt past the slide edge, "
                "tolerance 18pt | box (7315200, 914400, 10972800, 2743200) vs slide (9144000, 6858000) EMU "
                "(declined: inside a rotated or flipped group)"
            ),
        ],
    )
    assert fixed.read_bytes() == deck.read_bytes()


def small_text_deck(path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK])
    run_ = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.paragraphs[0].add_run()
    run_.text, run_.font.size = "Footnote", Pt(6)
    prs.save(str(path))
    return path


def test_a_fix_that_does_not_stick_is_reported_with_its_change(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    deck, fixed = small_text_deck(tmp_path / "small.pptx"), tmp_path / "fixed.pptx"
    rules = write_rules(tmp_path, "rules:\n  min-font-size:\n    min_pt: 7\n")
    monkeypatch.setitem(FIXERS, "min-font-size", lambda s, deck, params: (Change("size", "6pt", "7pt, 1 run"),))

    code, out, _ = run(
        ["fix", str(deck), "--out", str(fixed), "--rules", str(rules), "--report", str(tmp_path)], capsys
    )

    assert (code, out) == (
        1,
        f"FAIL {deck} -> {fixed}: 1 fixed in 1 pass, 1 remains\n"
        "fixed slide 1 [min-font-size] TextBox 1: 6pt text, min 7pt | size 6pt -> 7pt, 1 run\n"
        "remains slide 1 [min-font-size] TextBox 1: 6pt text, min 7pt | Footnote "
        "(did-not-stick: fixed in pass 1 and still fires)\n",
    )
    report = json.loads((tmp_path / "fix.json").read_text())
    assert [(f["rule"], f["pass"]) for f in report["fixed"]] == [("min-font-size", 1)]
    assert [(r["rule"], r["why"]) for r in report["remaining"]] == [("min-font-size", "did-not-stick")]
    assert fixed.read_bytes() == deck.read_bytes()


def test_fix_stops_after_three_passes_that_each_cause_a_new_violation(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    deck, fixed = bullet_deck(tmp_path / "bullets.pptx", ["Grew."]), tmp_path / "fixed.pptx"

    def strip_and_add_another(p, deck, params):
        p.xml.addnext(copy.deepcopy(p.xml))
        return fix_bullet_end(p, deck, params)

    monkeypatch.setitem(FIXERS, "no-bullet-end-punctuation", strip_and_add_another)

    code, out, _ = run(
        ["fix", str(deck), "--out", str(fixed), "--rules", str(write_rules(tmp_path, BULLET_RULES))], capsys
    )

    fixed_line = (
        "fixed slide 1 [no-bullet-end-punctuation] Content Placeholder 2: bullet ends with '.' | text 'Grew.' -> 'Grew'"
    )
    assert (code, out.splitlines()) == (
        1,
        [
            f"FAIL {deck} -> {fixed}: 3 fixed in 3 passes, 1 remains",
            fixed_line,
            fixed_line,
            fixed_line,
            "remains slide 1 [no-bullet-end-punctuation] Content Placeholder 2: bullet ends with '.' | Grew. "
            "(pass-limit: first fired after pass 3)",
        ],
    )
    assert bullet_texts(fixed) == ["Grew\nGrew\nGrew\nGrew."]


def test_doctor_lists_the_fixable_rules(capsys: pytest.CaptureFixture[str]) -> None:
    main(["doctor"])

    assert (
        "fixable rule ids: max-fonts-per-slide, no-bullet-end-punctuation, min-font-size, within-slide-bounds"
        in capsys.readouterr().out.splitlines()
    )


def test_every_fixer_fixes_a_registered_rule() -> None:
    assert set(FIXERS) <= set(RULES)

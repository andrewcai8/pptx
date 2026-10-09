from __future__ import annotations

import datetime
import hashlib
import io
import json
import zipfile
from pathlib import Path

import openpyxl
import pytest
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from deckcheck.changeset.cli import main
from deckcheck.changeset.model import SCHEMA_DIR, schemas
from deckcheck.package import Package

TITLE_AND_CONTENT, TITLE_ONLY = 1, 5
REF = {"t": "00:01:00", "speaker": "Ana Ruiz", "quote": "Use the new figures."}
ASK = {"id": "a1", "text": "Refresh the figures.", "refs": [REF]}
FLAG = {"id": "f1", "question": "Which year do the prices start?", "slides": [256], "refs": [REF]}
HELD = {"id": "h1", "text": "Keep the outlook wording.", "slides": [256], "refs": [REF]}
LABEL = (
    '<a:fld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" id="{00000000-0000-0000-0000-000000000001}"'
    " type=\"datetime'''+''''9''''%'''\"><a:rPr lang=\"en-US\" b=\"1\"/><a:t>+9%</a:t></a:fld>"
)
CUSTOM_SHOW = (
    '<p:custShowLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<p:custShow name="Short" id="0"><p:sldLst><p:sld r:id="{one}"/><p:sld r:id="{two}"/></p:sldLst></p:custShow>'
    "</p:custShowLst>"
)


def build_deck(path: Path) -> Path:
    prs = Presentation()
    one = prs.slides.add_slide(prs.slide_layouts[TITLE_AND_CONTENT])
    one.shapes.title.text = "Market outlook"
    body = one.placeholders[1].text_frame
    body.text = "Demand grows 4% a year"
    body.add_paragraph().text = "Prices hold"
    note = one.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(1)).text_frame.paragraphs[0]
    for text, size, bold in (("Revenue grew ", 14, False), ("12%", 18, True), (" in 2025", 12, False)):
        run = note.add_run()
        run.text, run.font.size, run.font.bold = text, Pt(size), bold
    table = one.shapes.add_table(2, 2, Inches(1), Inches(6), Inches(4), Inches(1)).table
    for (r, c), text in {(0, 0): "Year", (0, 1): "Sales", (1, 0): "2025", (1, 1): "$120m"}.items():
        table.cell(r, c).text = text

    two = prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
    two.shapes.title.text = "Sales by year"
    data = CategoryChartData()
    data.categories = ["2024", "2025"]
    data.add_series("Sales", (100, 120))
    two.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), data)

    three = prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
    three.shapes.title.text = "Next steps"
    label = three.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1)).text_frame.paragraphs[0]
    label._p.append(parse_xml(LABEL))

    rids = [s.get(qn("r:id")) for s in prs.slides._sldIdLst]
    prs.part._element.find(qn("p:notesSz")).addnext(parse_xml(CUSTOM_SHOW.format(one=rids[0], two=rids[1])))
    prs.save(str(path))
    return path


REVENUE = {"kind": "replace_text", "slide": 256, "shape": 4, "old": "12%", "new": "15%"}
CELL = {"kind": "set_cell", "slide": 256, "shape": 5, "row": 1, "col": 1, "old": "$120m", "new": "$130m"}
POINT = {"kind": "set_chart_value", "slide": 257, "shape": 3, "series": 0, "point": 1, "old": 120, "new": 130}
ADD = {"kind": "add_slide", "layout": "Title and Content", "after": 257}
FILL = {"kind": "fill_placeholder", "slide": "add", "shape": 3, "paragraphs": [{"text": "Raise list prices"}, {"text": "Hold discounts", "level": 1}]}
MOVE = {"kind": "move_slide", "slide": 256, "after": 257}
DELETE = {"kind": "delete_slide", "slide": 257}
INSERT = {"kind": "insert_paragraph", "slide": 256, "shape": 4, "after": 0, "text": "Up from 9%"}
LABEL_OP = {"kind": "replace_text", "slide": 258, "shape": 3, "old": "+9%", "new": "+10%"}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def change(cid: str, op: dict, decision: str | dict | None = None, **extra) -> dict:
    c = {"id": cid, "ask_id": "a1", "rationale": "The figures moved.", "refs": [REF], "op": op, **extra}
    if decision is not None:
        c["decision"] = decision
    return c


def changeset(deck: Path, changes: list[dict], name: str = "changeset.json", lists: dict | None = None, **source) -> Path:
    doc = {
        "meeting": {"title": "Pricing review", "date": "2026-10-01"},
        "source": {"path": str(deck), "sha256": sha(deck), **source},
        "asks": [ASK],
        "changes": changes,
        **(lists or {}),
    }
    path = deck.parent / name
    path.write_text(json.dumps(doc, indent=2))
    return path


def run(argv: list[str], capsys: pytest.CaptureFixture[str]) -> tuple[int, str, str]:
    code = main([str(a) for a in argv])
    out, err = capsys.readouterr()
    return code, out, err


@pytest.fixture
def deck(tmp_path: Path) -> Path:
    return build_deck(tmp_path / "deck.pptx")


def paragraphs(path: Path, slide_id: int, shape_id: int) -> list[str]:
    slide = next(s for s in Presentation(str(path)).slides if s.slide_id == slide_id)
    shape = next(s for s in slide.shapes if s.shape_id == shape_id)
    return [p.text for p in shape.text_frame.paragraphs]


def cell(path: Path) -> str:
    slide = next(s for s in Presentation(str(path)).slides if s.slide_id == 256)
    return next(s for s in slide.shapes if s.shape_id == 5).table.cell(1, 1).text


def chart_values(path: Path) -> tuple[tuple[float, ...], object]:
    slide = next(s for s in Presentation(str(path)).slides if s.slide_id == 257)
    chart = next(s for s in slide.shapes if s.has_chart).chart
    book = openpyxl.load_workbook(io.BytesIO(chart.part.chart_workbook.xlsx_part.blob))
    return tuple(chart.plots[0].series[0].values), book.active["B3"].value


def slide_ids(path: Path) -> list[int]:
    return [s.slide_id for s in Presentation(str(path)).slides]


def test_committed_schemas_match_the_models() -> None:
    for name, text in schemas().items():
        assert (SCHEMA_DIR / name).read_text() == text, (
            f"{name} is stale; run uv run --project deckcheck python -m deckcheck.changeset.model"
        )


def test_reading_the_rels_of_a_part_that_has_none_adds_no_part(deck: Path) -> None:
    data = deck.read_bytes()
    pkg = Package(data)

    assert (len(pkg.rels("ppt/slides/slide9.xml")), pkg.has("ppt/slides/_rels/slide9.xml.rels"), pkg.to_bytes() is data) == (0, False, True)


def test_validate_lists_each_change_with_its_source_text(deck: Path, capsys: pytest.CaptureFixture[str]) -> None:
    cs = changeset(deck, [change("c1", REVENUE), change("c2", CELL)])

    assert run(["validate", cs], capsys) == (
        0,
        f"VALID {cs}: 2 changes (1 text-only, 1 structural) on slide 1; 1 ask, 0 flags, 0 held\n"
        "  c1 replace_text slide 1 shape 4 'TextBox 3': '12%' -> '15%' (text-only)\n"
        "  c2 set_cell slide 1 shape 5 'Table 4': '$120m' -> '$130m' (structural)\n",
        "",
    )


def test_execute_writes_every_change_as_written(deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    cs = changeset(deck, [change("c1", REVENUE), change("c2", CELL), change("c3", POINT), change("add", ADD), change("fill", FILL)])
    out = tmp_path / "executed.pptx"

    code, printed, _ = run(["execute", cs, "--out", out], capsys)

    assert (code, printed.splitlines()[0]) == (
        0,
        f"EXECUTED {cs} -> {out}: 5 changes (1 text-only, 4 structural) written, sha256 {sha(out)}",
    )
    assert paragraphs(out, 256, 4) == ["Revenue grew 15% in 2025"]
    assert cell(out) == "$130m"
    assert chart_values(out) == ((100.0, 130.0), 130)
    assert slide_ids(out) == [256, 257, 259, 258]
    assert paragraphs(out, 259, 3) == ["Raise list prices", "Hold discounts"]


def test_edited_decisions_write_the_reviewers_text(deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    cs = changeset(
        deck,
        [
            change("c1", REVENUE, {"edited": "18%"}),
            change("c2", CELL, {"edited": "$125m"}),
            change("c3", POINT, {"edited": "1,234.5"}),
            change("add", ADD, "keep_new"),
            change("fill", FILL, {"edited": "Raise prices\nHold discounts\nReview in May"}),
        ],
    )
    out = tmp_path / "final.pptx"

    code, printed, _ = run(["apply", cs, "--out", out], capsys)

    assert (code, printed) == (0, f"APPLIED {cs} -> {out}: 1 kept new, 4 edited, 0 kept old, 0 dropped, sha256 {sha(out)}\n")
    assert paragraphs(out, 256, 4) == ["Revenue grew 18% in 2025"]
    assert cell(out) == "$125m"
    assert chart_values(out) == ((100.0, 1234.5), 1234.5)
    slide = next(s for s in Presentation(str(out)).slides if s.slide_id == 259)
    body = next(s for s in slide.shapes if s.shape_id == 3)
    assert [(p.text, p.level) for p in body.text_frame.paragraphs] == [("Raise prices", 0), ("Hold discounts", 1), ("Review in May", 1)]


def test_keep_old_leaves_a_change_out_and_keep_new_writes_it(deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    cs = changeset(deck, [change("c1", REVENUE, "keep_old"), change("c2", CELL, "keep_new"), change("c3", POINT, "keep_old")])
    out = tmp_path / "final.pptx"

    code, _, _ = run(["apply", cs, "--out", out], capsys)

    assert (code, paragraphs(out, 256, 4), cell(out), chart_values(out)) == (0, ["Revenue grew 12% in 2025"], "$130m", ((100.0, 120.0), 120))


def test_apply_refuses_pending_decisions_and_writes_nothing(deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    cs = changeset(deck, [change("c1", REVENUE), change("c2", CELL, "keep_new"), change("c3", POINT, "pending")])
    out = tmp_path / "final.pptx"

    assert run(["apply", cs, "--out", out], capsys) == (1, f"PENDING {cs}: 2 decisions pending (c1, c3); nothing written\n", "")
    assert not out.exists()


def test_a_fill_is_dropped_when_its_added_slide_is_kept_old(deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    cs = changeset(deck, [change("add", ADD, "keep_old"), change("fill", FILL, "keep_new"), change("c1", REVENUE, "keep_old")])
    out = tmp_path / "final.pptx"

    assert run(["apply", cs, "--out", out], capsys) == (
        0,
        f"APPLIED {cs} -> {out}: 0 kept new, 0 edited, 2 kept old, 1 dropped, sha256 {sha(deck)}\n"
        "  dropped fill: it fills the slide add adds, which was kept old\n",
        "",
    )
    assert out.read_bytes() == deck.read_bytes()


def test_a_tab_is_written_as_a_tab(deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    cs = changeset(deck, [change("c1", {**REVENUE, "new": "15%\t"}, {"edited": "\t18%"}), change("c2", CELL, {"edited": "$125m\tnet"})])
    out = tmp_path / "final.pptx"

    assert run(["apply", cs, "--out", out], capsys)[0] == 0
    assert (paragraphs(out, 256, 4), cell(out)) == (["Revenue grew \t18% in 2025"], "$125m\tnet")


ALL_OPS = [change("c1", REVENUE), change("c2", CELL), change("c3", POINT), change("add", ADD), change("fill", FILL), change("c4", MOVE), change("c5", LABEL_OP)]


def decided(decision: str) -> list[dict]:
    return [{**c, "decision": decision} for c in ALL_OPS]


def test_decisions_replay_onto_the_source_byte_for_byte(deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    source_sha = sha(deck)
    executed = tmp_path / "executed.pptx"
    run(["execute", changeset(deck, ALL_OPS), "--out", executed], capsys)
    executed_sha = sha(executed)
    old, new, again = tmp_path / "old.pptx", tmp_path / "new.pptx", tmp_path / "again.pptx"

    run(["apply", changeset(deck, decided("keep_old"), "old.json"), "--out", old], capsys)
    run(["apply", changeset(deck, decided("keep_new"), "new.json"), "--out", new], capsys)
    run(["apply", changeset(deck, decided("keep_new"), "new.json"), "--out", again], capsys)

    # Equal bytes for an edited deck hold for one zlib build: untouched members are recompressed.
    assert old.read_bytes() == deck.read_bytes()
    assert new.read_bytes() == executed.read_bytes() == again.read_bytes()
    assert (sha(deck), sha(executed)) == (source_sha, executed_sha)
    assert executed.read_bytes() != deck.read_bytes()


def test_untouched_runs_in_an_edited_paragraph_keep_their_formatting(deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = tmp_path / "executed.pptx"
    run(["execute", changeset(deck, [change("c1", {**REVENUE, "old": "grew 12%", "new": "grew 15%"})]), "--out", out], capsys)

    def runs(path: Path) -> list[tuple[str, bytes]]:
        slide = next(s for s in Presentation(str(path)).slides if s.slide_id == 256)
        p = next(s for s in slide.shapes if s.shape_id == 4).text_frame.paragraphs[0]._p
        return [(r.text, etree.tostring(r.find(qn("a:rPr")), method="c14n")) for r in p.r_lst]

    before, after = runs(deck), runs(out)
    assert [t for t, _ in after] == ["Revenue grew ", "15%", " in 2025"]
    assert [rpr for _, rpr in after] == [rpr for _, rpr in before]


@pytest.mark.parametrize("order", [("grew", "pct"), ("pct", "grew")], ids=["text-order", "reverse-order"])
def test_abutting_quotes_that_both_insert_at_their_shared_edge_keep_text_order(
    deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str], order: tuple[str, str]
) -> None:
    ops = {"grew": {**REVENUE, "old": "grew ", "new": "grew by "}, "pct": {**REVENUE, "old": "12%", "new": "c.12%"}}
    out = tmp_path / "executed.pptx"
    run(["execute", changeset(deck, [change(cid, ops[cid]) for cid in order]), "--out", out], capsys)

    assert paragraphs(out, 256, 4) == ["Revenue grew by c.12% in 2025"]


def test_inserted_paragraphs_follow_their_anchor_in_changeset_order_styled_like_it(
    deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    first = {"kind": "insert_paragraph", "slide": 256, "shape": 4, "after": 0, "text": "Up from 9%"}
    out = tmp_path / "executed.pptx"
    run(["execute", changeset(deck, [change("c1", first), change("c2", {**first, "text": "Margins hold"})]), "--out", out], capsys)

    slide = next(s for s in Presentation(str(out)).slides if s.slide_id == 256)
    box = next(s for s in slide.shapes if s.shape_id == 4).text_frame
    assert [[(r.text, r.font.size.pt) for r in p.runs] for p in box.paragraphs] == [
        [("Revenue grew ", 14.0), ("12%", 18.0), (" in 2025", 12.0)],
        [("Up from 9%", 14.0)],
        [("Margins hold", 14.0)],
    ]


def test_a_think_cell_label_rewrites_its_field_format_with_its_text(deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = tmp_path / "executed.pptx"
    run(["execute", changeset(deck, [change("c1", LABEL_OP)]), "--out", out], capsys)

    slide = next(s for s in Presentation(str(out)).slides if s.slide_id == 258)
    field = next(s for s in slide.shapes if s.shape_id == 3).text_frame.paragraphs[0]._p.find(qn("a:fld"))
    assert (field.findtext(qn("a:t")), field.get("type")[8:].replace("'", "")) == ("+10%", "+10%")


def test_text_typed_into_an_empty_think_cell_label_is_quoted_in_its_field_format(
    deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    prs = Presentation(str(deck))
    p = prs.slides[2].shapes.add_textbox(Inches(4), Inches(2), Inches(2), Inches(1)).text_frame.paragraphs[0]
    p._p.append(parse_xml(LABEL.replace("'''+''''9''''%'''", "''''").replace("+9%", "")))
    p.add_run().text = " p.a."
    prs.save(str(deck))
    out = tmp_path / "executed.pptx"
    run(["execute", changeset(deck, [change("c1", {**LABEL_OP, "shape": 4, "old": " p.a.", "new": "5% p.a."})]), "--out", out], capsys)

    slide = next(s for s in Presentation(str(out)).slides if s.slide_id == 258)
    label = next(s for s in slide.shapes if s.shape_id == 4).text_frame.paragraphs[0]
    assert (label.text, label._p.find(qn("a:fld")).get("type")) == ("5% p.a.", "datetime'''''5%'")


def test_a_chart_edit_writes_the_cache_and_the_embedded_xlsx(deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = tmp_path / "executed.pptx"
    run(["execute", changeset(deck, [change("c1", POINT)]), "--out", out], capsys)

    with zipfile.ZipFile(out) as z:
        chart = z.read("ppt/charts/chart1.xml").decode()
        book = openpyxl.load_workbook(io.BytesIO(z.read("ppt/embeddings/Microsoft_Excel_Sheet1.xlsx")))
    assert ('<c:pt idx="1"><c:v>130</c:v></c:pt>' in chart, book.active["B3"].value) == (True, 130)


@pytest.mark.parametrize(
    ("value", "held"),
    [
        ("n/a", "holds text"),
        (True, "holds a true/false value"),
        ("#N/A", "holds an error"),
        (datetime.datetime(2025, 1, 1), "holds a date"),
        (None, "is blank"),
    ],
    ids=["text", "boolean", "error", "date", "blank"],
)
def test_a_chart_point_whose_workbook_cell_is_not_a_number_is_refused(
    deck: Path, capsys: pytest.CaptureFixture[str], value: object, held: str
) -> None:
    prs = Presentation(str(deck))
    workbook = next(s for s in prs.slides[1].shapes if s.has_chart).chart.part.chart_workbook
    book = openpyxl.load_workbook(io.BytesIO(workbook.xlsx_part.blob))
    book.iso_dates = True
    book.active["B3"] = value
    buf = io.BytesIO()
    book.save(buf)
    workbook.update_from_xlsx_blob(buf.getvalue())
    prs.save(str(deck))
    cs = changeset(deck, [change("c1", POINT)])

    assert run(["validate", cs], capsys) == (
        1,
        f"INVALID {cs}: 1 problem\n  c1 op: shape 3 'Chart 2': workbook cell Sheet1!B3 {held}, not a number\n",
        "",
    )


def test_delete_drops_the_slide_its_parts_and_its_custom_show_entry(deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = tmp_path / "executed.pptx"
    run(["execute", changeset(deck, [change("c1", DELETE)]), "--out", out], capsys)

    with zipfile.ZipFile(out) as z:
        names = set(z.namelist())
        pres = etree.fromstring(z.read("ppt/presentation.xml"))
    shown = pres.findall(f".//{qn('p:custShow')}/{qn('p:sldLst')}/{qn('p:sld')}")
    assert slide_ids(out) == [256, 258]
    assert {"ppt/slides/slide2.xml", "ppt/charts/chart1.xml", "ppt/embeddings/Microsoft_Excel_Sheet1.xlsx"} & names == set()
    assert len(shown) == 1


def test_the_review_view_reads_each_change_and_slide_from_the_source_and_executed_decks(
    deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    prs = Presentation(str(deck))
    group = prs.slides[0].shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "Pilot in Q3"
    group.left = Inches(2)
    prs.save(str(deck))
    changes = [
        change("text", {"kind": "replace_text", "slide": 256, "shape": 7, "old": "Q3", "new": "Q4"}),
        change("move", MOVE),
        change("add", ADD),
        change("title", {**FILL, "shape": 2, "paragraphs": [{"text": "Pricing"}, {"text": "actions"}]}),
        change("body", FILL),
        change("drop", {**DELETE, "slide": 258}),
    ]
    cs = changeset(deck, changes, lists={"flags": [FLAG], "held": [HELD]})
    out, review = tmp_path / "executed.pptx", tmp_path / "review.json"

    assert run(["execute", cs, "--out", out, "--review", review], capsys)[0] == 0

    common = {"ask_id": "a1", "decision": "pending", "rationale": "The figures moved.", "refs": [REF], "notes": []}
    edit, slide_only = ["keep_new", "keep_old", "edited"], ["keep_new", "keep_old"]
    assert json.loads(review.read_text()) == {
        "source": {"path": str(deck), "sha256": sha(deck)},
        "executed": {"path": str(out), "sha256": sha(out)},
        "meeting": {"title": "Pricing review", "date": "2026-10-01"},
        "asks": [ASK],
        "flags": [FLAG],
        "held": [HELD],
        "slides": [
            {"key": 256, "source_index": 1, "executed_index": 2, "title": "Market outlook"},
            {"key": 257, "source_index": 2, "executed_index": 1, "title": "Sales by year"},
            {"key": 258, "source_index": 3, "executed_index": None, "title": "Next steps"},
            {"key": "add", "source_index": None, "executed_index": 3, "title": "Pricing actions"},
        ],
        "changes": [
            {
                **common,
                "id": "text",
                "kind": "replace_text",
                "structural": False,
                "slide": 256,
                "source_index": 1,
                "executed_index": 2,
                "shape": {"id": 7, "name": "TextBox 6", "box": {"x": 1828800, "y": 914400, "w": 2743200, "h": 914400}},
                "before": "Pilot in Q3",
                "after": "Pilot in Q4",
                "span": [9, 11],
                "depends_on": None,
                "admits": edit,
            },
            {
                **common,
                "id": "move",
                "kind": "move_slide",
                "structural": True,
                "slide": 256,
                "source_index": 1,
                "executed_index": 2,
                "shape": None,
                "before": "first",
                "after": "after slide 2 (id 257)",
                "span": None,
                "depends_on": None,
                "admits": slide_only,
            },
            {
                **common,
                "id": "add",
                "kind": "add_slide",
                "structural": True,
                "slide": "add",
                "source_index": None,
                "executed_index": 3,
                "shape": None,
                "before": None,
                "after": "new slide on layout 'Title and Content', after slide 2 (id 257)",
                "span": None,
                "depends_on": None,
                "admits": slide_only,
            },
            {
                **common,
                "id": "title",
                "kind": "fill_placeholder",
                "structural": True,
                "slide": "add",
                "source_index": None,
                "executed_index": 3,
                "shape": {"id": 2, "name": "Title 1", "box": None},
                "before": None,
                "after": "Pricing\nactions",
                "span": None,
                "depends_on": "add",
                "admits": edit,
            },
            {
                **common,
                "id": "body",
                "kind": "fill_placeholder",
                "structural": True,
                "slide": "add",
                "source_index": None,
                "executed_index": 3,
                "shape": {"id": 3, "name": "Content Placeholder 2", "box": None},
                "before": None,
                "after": "Raise list prices\nHold discounts",
                "span": None,
                "depends_on": "add",
                "admits": edit,
            },
            {
                **common,
                "id": "drop",
                "kind": "delete_slide",
                "structural": True,
                "slide": 258,
                "source_index": 3,
                "executed_index": None,
                "shape": None,
                "before": "Next steps",
                "after": None,
                "span": None,
                "depends_on": None,
                "admits": slide_only,
            },
        ],
    }


SECTIONS = (
    '<p:extLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    '<p:ext uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}">'
    '<p14:sectionLst xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">'
    '<p14:section name="Intro" id="{00000000-0000-0000-0000-00000000000A}">'
    '<p14:sldIdLst><p14:sldId id="256"/><p14:sldId id="257"/></p14:sldIdLst></p14:section>'
    '<p14:section name="Body" id="{00000000-0000-0000-0000-00000000000B}">'
    '<p14:sldIdLst><p14:sldId id="258"/><p14:sldId id="259"/></p14:sldIdLst></p14:section>'
    "</p14:sectionLst></p:ext></p:extLst>"
)


def test_two_added_slides_get_their_own_ids_and_parts(deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    changes = [
        change("add", ADD, "keep_old"),
        change("fill", FILL, "keep_new"),
        change("lead", {**ADD, "after": None}, "keep_new"),
        change("lead-fill", {**FILL, "slide": "lead", "paragraphs": [{"text": "Agenda"}]}, "keep_new"),
    ]
    executed, final = tmp_path / "executed.pptx", tmp_path / "final.pptx"

    run(["execute", changeset(deck, changes), "--out", executed], capsys)
    run(["apply", changeset(deck, changes), "--out", final], capsys)

    with zipfile.ZipFile(executed) as z:
        added = sorted(n for n in z.namelist() if n in ("ppt/slides/slide4.xml", "ppt/slides/slide5.xml"))
    assert (slide_ids(executed), added) == ([260, 256, 257, 259, 258], ["ppt/slides/slide4.xml", "ppt/slides/slide5.xml"])
    assert (paragraphs(executed, 259, 3), paragraphs(executed, 260, 3)) == (["Raise list prices", "Hold discounts"], ["Agenda"])
    assert (slide_ids(final), paragraphs(final, 260, 3)) == ([260, 256, 257, 258], ["Agenda"])

def test_added_moved_and_deleted_slides_keep_the_sections_in_step(deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    prs = Presentation(str(deck))
    prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY]).shapes.title.text = "Appendix"
    prs.part._element.append(parse_xml(SECTIONS))
    prs.save(str(deck))
    changes = [change("move", {**MOVE, "after": 258}), change("add", ADD), change("drop", {**DELETE, "slide": 259})]
    out = tmp_path / "executed.pptx"

    run(["execute", changeset(deck, changes), "--out", out], capsys)

    with zipfile.ZipFile(out) as z:
        pres = etree.fromstring(z.read("ppt/presentation.xml"))
    p14 = "{http://schemas.microsoft.com/office/powerpoint/2010/main}"
    sections = [(s.get("name"), [int(i.get("id")) for i in s.iter(f"{p14}sldId")]) for s in pres.iter(f"{p14}section")]
    assert (slide_ids(out), sections) == ([257, 260, 258, 256], [("Intro", [257, 260]), ("Body", [258, 256])])


@pytest.mark.parametrize(
    ("changes", "order"),
    [
        ([change("c1", MOVE)], [257, 256, 258]),
        ([change("c1", {**MOVE, "after": None})], [256, 257, 258]),
        ([change("c1", {**MOVE, "slide": 258, "after": None})], [258, 256, 257]),
        ([change("add", ADD), change("c1", MOVE)], [257, 259, 256, 258]),
        ([change("c1", {**MOVE, "after": 258}), change("add", {**ADD, "after": 256})], [257, 258, 256, 259]),
    ],
    ids=["after", "first", "last-to-first", "chained-after-one-anchor", "follows-a-moved-anchor"],
)
def test_moved_and_added_slides_land_after_their_anchor(
    deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str], changes: list[dict], order: list[int]
) -> None:
    out = tmp_path / "executed.pptx"
    run(["execute", changeset(deck, changes), "--out", out], capsys)

    assert slide_ids(out) == order


def bad_json(deck: Path) -> Path:
    path = deck.parent / "changeset.json"
    path.write_text("{")
    return path


def unquoted_label(deck: Path) -> bytes:
    with zipfile.ZipFile(deck) as z:
        return z.read("ppt/slides/slide3.xml").replace(b"datetime'''+''''9''''%'''", b"datetime'+'9'%'")


def date_field(deck: Path) -> Path:
    prs = Presentation(str(deck))
    p = prs.slides[2].shapes.add_textbox(Inches(4), Inches(2), Inches(2), Inches(1)).text_frame.paragraphs[0]
    p._p.append(parse_xml(LABEL.replace("'''+''''9''''%'''", "yyyy" + "''''" * 70).replace("+9%", "2025")))
    prs.save(str(deck))
    return deck


def first_text(deck: Path, shape_id: int, text: str) -> Path:
    prs = Presentation(str(deck))
    shape = next(s for s in prs.slides[0].shapes if s.shape_id == shape_id)
    shape._element.find(f".//{qn('a:t')}").text = text
    prs.save(str(deck))
    return deck


def rezip(deck: Path, part: str, body: bytes | None) -> Path:
    data = deck.read_bytes()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(deck, "w") as dst:
        for info in src.infolist():
            if info.filename != part:
                dst.writestr(info, src.read(info))
            elif body is not None:
                dst.writestr(info, body)
    return deck


@pytest.mark.parametrize(
    ("write", "problem"),
    [
        (lambda d: changeset(d, [change("c1", {**REVENUE, "slide": 999})]), "c1 op.slide: no slide 999; slide ids: 256 (slide 1), 257 (slide 2), 258 (slide 3)"),
        (
            lambda d: changeset(d, [change("c1", {**REVENUE, "shape": 9})]),
            "c1 op.shape: slide 1 (id 256) has no shape 9; shapes with text: 2 'Title 1' ('Market outlook'), "
            "3 'Content Placeholder 2' ('Demand grows 4% a year Prices…'), 4 'TextBox 3' ('Revenue grew 12% in 2025')",
        ),
        (lambda d: changeset(d, [change("c1", {**CELL, "shape": 9})]), "c1 op.shape: slide 1 (id 256) has no shape 9; tables: 5 'Table 4'"),
        (lambda d: changeset(d, [change("c1", {**POINT, "shape": 9})]), "c1 op.shape: slide 2 (id 257) has no shape 9; charts: 3 'Chart 2'"),
        (lambda d: changeset(d, [change("c1", {**REVENUE, "old": "13%"})]), "c1 op.old: '13%' is not in shape 4 'TextBox 3'; its text is 'Revenue grew 12% in 2025'"),
        (
            lambda d: changeset(d, [change("c1", {**REVENUE, "shape": 3, "old": "s "})]),
            "c1 op.old: 's ' occurs 2 times in shape 3 'Content Placeholder 2', in paragraphs 0, 1; quote more of the text or name the paragraph",
        ),
        (lambda d: changeset(d, [change("c1", REVENUE), change("c2", {**REVENUE, "old": "12% in", "new": "15% in"})]), "c2 op.old: '12% in' overlaps c1's quote '12%' in paragraph 0"),
        (lambda d: changeset(d, [change("c1", CELL), change("c2", CELL)]), "c2 op: shape 5 row 1 col 1 is already changed by c1"),
        (lambda d: changeset(d, [change("c1", {**DELETE, "slide": 256}), change("c2", REVENUE)]), "c2 op.slide: slide 1 (id 256) is deleted by c1"),
        (lambda d: changeset(d, [change("c1", {**MOVE, "after": 256})]), "c1 op.after: a slide cannot follow itself"),
        (lambda d: changeset(d, [change("c1", MOVE), change("c2", {**MOVE, "after": 258})]), "c2 op.slide: slide 1 (id 256) is already moved by c1"),
        (
            lambda d: changeset(d, [change("c1", REVENUE)], sha256="0" * 64),
            lambda d: f"source.sha256: {d} has sha256 {sha(d)}; the deck changed since the maker read it",
        ),
        (lambda d: changeset(d, [change("c1", REVENUE, before="12%")]), "c1 before: the engine reads this from the source deck; remove it"),
        (lambda d: changeset(d, [change("c1", {**REVENUE, "after": "15%"})]), "c1 op.after: replace_text has no field 'after'"),
        (lambda d: changeset(d, [change("c1", DELETE, {"edited": "Costs"})]), "c1 decision: delete_slide admits keep_new or keep_old only"),
        (lambda d: changeset(d, [change("c1", POINT, {"edited": "lots"})]), "c1 decision: edited value 'lots' is not a number like 1234.5 or 1,234.5"),
        (lambda d: changeset(d, [change("c1", POINT, {"edited": "1,5"})]), "c1 decision: edited value '1,5' is not a number like 1234.5 or 1,234.5"),
        (lambda d: changeset(d, [change("c1", POINT, {"edited": "1_000"})]), "c1 decision: edited value '1_000' is not a number like 1234.5 or 1,234.5"),
        (lambda d: changeset(d, [change("c1", INSERT, {"edited": ""})]), "c1 decision: edited text '': String should have at least 1 character"),
        (
            lambda d: changeset(d, [change("c1", {**LABEL_OP, "new": "+9'%"})]),
            "c1 op.new: a field's text cannot take an apostrophe, which its format uses for quoting",
        ),
        (bad_json, "changeset: not JSON: Expecting property name enclosed in double quotes: line 1 column 2 (char 1)"),
        (
            lambda d: changeset(d, [change("c1", {"kind": "recolor", "slide": 256})]),
            "c1 op: Input tag 'recolor' found using 'kind' does not match any of the expected tags: 'replace_text', "
            "'insert_paragraph', 'set_cell', 'set_chart_value', 'add_slide', 'fill_placeholder', 'delete_slide', 'move_slide'",
        ),
        (lambda d: changeset(d, [change("c1", REVENUE), change("c2", {**FILL, "slide": "c1"})]), "c2 op.slide: c1 is a replace_text change, not an add_slide"),
        (
            lambda d: changeset(d, [change("c1", REVENUE)], lists={"flags": [{**FLAG, "slides": [256, 999]}]}),
            "f1 slides: no slide 999; slide ids: 256 (slide 1), 257 (slide 2), 258 (slide 3)",
        ),
        (
            lambda d: changeset(d, [change("c1", REVENUE)], lists={"held": [{**HELD, "slides": [998]}]}),
            "h1 slides: no slide 998; slide ids: 256 (slide 1), 257 (slide 2), 258 (slide 3)",
        ),
        (
            lambda d: changeset(d, [change("c1", REVENUE)], lists={"asks": [ASK, {**ASK, "text": "Again."}]}),
            "a1 id: asks[0] and asks[1] share this id",
        ),
        (lambda d: changeset(d, [change("c1", REVENUE)], lists={"flags": [FLAG, FLAG]}), "f1 id: flags[0] and flags[1] share this id"),
        (lambda d: changeset(d, [change("c1", REVENUE)], lists={"held": [HELD, HELD]}), "h1 id: held[0] and held[1] share this id"),
        (
            lambda d: changeset(rezip(d, "ppt/slides/slide3.xml", unquoted_label(d)), [change("c1", {**LABEL_OP, "new": "+7%"})]),
            "c1 op.new: the change touches a \"datetime'+'9'%'\" field, which only PowerPoint fills in",
        ),
        (
            lambda d: changeset(date_field(d), [change("c1", {**LABEL_OP, "shape": 4, "old": "2025", "new": "2026"})]),
            "c1 op.new: the change touches a " + repr("datetimeyyyy" + "'" * 27 + "…") + " field, which only PowerPoint fills in",
        ),
        (
            lambda d: changeset(first_text(d, 4, "Revenue\ngrew "), [change("c1", {**REVENUE, "old": "Revenue\ngrew", "new": "Revenue\nrose"})]),
            "c1 op.old: shape 4 'TextBox 3' holds a line feed (\"\\n\") inside the text of paragraph 0, which no op can write; "
            "quote the text on one side of it",
        ),
        (
            lambda d: changeset(first_text(d, 5, "Year\nended"), [change("c1", {**CELL, "row": 0, "col": 0, "old": "Year\nended", "new": "Years\nended"})]),
            "c1 op: row 0 col 0 holds a line feed (\"\\n\") inside a paragraph, so \"\\n\" cannot mark where its paragraphs split",
        ),
        (lambda d: changeset(d, [change("c1", REVENUE)], path="deck\x00.pptx"), "source.path: 'deck\\x00.pptx' is not a path: embedded null byte"),
        (
            lambda d: changeset(rezip(d, "ppt/embeddings/Microsoft_Excel_Sheet1.xlsx", b"not a workbook"), [change("c1", POINT)]),
            "c1 op: shape 3 'Chart 2': the workbook ppt/embeddings/Microsoft_Excel_Sheet1.xlsx cannot be read: File is not a zip file",
        ),
        (lambda d: changeset(rezip(d, "ppt/charts/chart1.xml", None), [change("c1", POINT)]), "c1 op: cannot read the deck: ppt/charts/chart1.xml is missing"),
        (lambda d: changeset(d, [change("c1", {**REVENUE, "new": "15%\f"})]), "c1 op.new: holds the control character '\\x0c' at index 3, which a deck cannot hold; remove it"),
        (lambda d: changeset(d, [change("c1", {**INSERT, "text": "Up\x07"})]), "c1 op.text: holds the control character '\\x07' at index 2, which a deck cannot hold; remove it"),
        (lambda d: changeset(d, [change("c1", {**CELL, "new": "$1\r30m"})]), "c1 op.new: holds the control character '\\r' at index 2, which a deck cannot hold; remove it"),
        (lambda d: changeset(d, [change("c1", {**REVENUE, "new": "15\ufffe%"})]), "c1 op.new: holds the control character '\\ufffe' at index 2, which a deck cannot hold; remove it"),
        (
            lambda d: changeset(d, [change("add", ADD), change("fill", {**FILL, "paragraphs": [{"text": "a\x1fb"}]})]),
            "fill op.paragraphs.0.text: holds the control character '\\x1f' at index 1, which a deck cannot hold; remove it",
        ),
        (
            lambda d: changeset(d, [change("c1", REVENUE, {"edited": "15%\x1f"})]),
            "c1 decision: edited text '15%\\x1f': holds the control character '\\x1f' at index 3, which a deck cannot hold; remove it",
        ),
        (
            lambda d: changeset(d, [change("add", ADD), change("fill", FILL, {"edited": "Raise prices\r\nHold discounts"})]),
            "fill decision: edited text 'Raise prices\\r\\nHold discounts': holds the control character '\\r' at index 12, which a deck cannot hold; remove it",
        ),
    ],
    ids=[
        "unknown-slide",
        "unknown-shape",
        "unknown-table",
        "unknown-chart",
        "old-not-in-shape",
        "old-occurs-twice",
        "overlapping-quotes",
        "same-cell-twice",
        "edit-on-deleted-slide",
        "move-after-itself",
        "two-moves",
        "wrong-sha256",
        "maker-written-before",
        "unknown-op-field",
        "edited-delete",
        "non-numeric-chart-edit",
        "decimal-comma-chart-edit",
        "underscore-chart-edit",
        "empty-edited-paragraph",
        "apostrophe-inserted-into-a-field",
        "malformed-json",
        "unknown-op-kind",
        "fill-names-a-non-add",
        "flag-names-a-missing-slide",
        "held-names-a-missing-slide",
        "duplicate-ask-ids",
        "duplicate-flag-ids",
        "duplicate-held-ids",
        "format-character-outside-quotes",
        "long-field-format",
        "line-feed-in-the-source-quote",
        "line-feed-in-a-cell",
        "nul-in-source-path",
        "workbook-not-a-zip",
        "chart-part-missing",
        "form-feed-in-new",
        "bell-in-inserted-text",
        "carriage-return-in-a-cell",
        "noncharacter-in-new",
        "unit-separator-in-a-fill",
        "unit-separator-in-an-edit",
        "windows-line-ending-in-an-edited-fill",
    ],
)
def test_a_bad_changeset_names_each_problem_and_exits_1(deck: Path, capsys: pytest.CaptureFixture[str], write, problem) -> None:
    cs = write(deck)
    expected = problem(deck) if callable(problem) else problem

    assert run(["validate", cs], capsys) == (1, f"INVALID {cs}: 1 problem\n  {expected}\n", "")


@pytest.mark.parametrize(
    ("argv", "message"),
    [
        (["execute", "changeset.json", "--out", "deck.pptx"], "--out deck.pptx is the source deck; the engine never writes over it"),
        (["execute", "changeset.json", "--out", "new.pptx", "--review", "deck.pptx"], "--review deck.pptx is the source deck; the engine never writes over it"),
        (["execute", "changeset.json", "--out", "./changeset.json"], "--out changeset.json is the ChangeSet; the engine never writes over it"),
        (["execute", "changeset.json", "--out", "new.pptx", "--review", "changeset.json"], "--review changeset.json is the ChangeSet; the engine never writes over it"),
        (["execute", "changeset.json", "--out", "new.pptx", "--review", "sub/../new.pptx"], "--review sub/../new.pptx is also --out; give each its own path"),
        (["apply", "changeset.json", "--out", "changeset.json"], "--out changeset.json is the ChangeSet; the engine never writes over it"),
    ],
    ids=["out-is-source", "review-is-source", "out-is-changeset", "review-is-changeset", "review-is-out", "apply-out-is-changeset"],
)
def test_no_output_overwrites_an_input_or_the_other_output(
    deck: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], argv: list[str], message: str
) -> None:
    changeset(deck, [change("c1", REVENUE, "keep_new")])
    (tmp_path / "sub").mkdir()
    monkeypatch.chdir(tmp_path)
    files = {p.name: p.read_bytes() for p in tmp_path.iterdir() if p.is_file()}

    assert run(argv, capsys) == (2, "", f"error: {message}\n")
    assert {p.name: p.read_bytes() for p in tmp_path.iterdir() if p.is_file()} == files



def not_a_deck(path: Path, kind: str) -> Path:
    match kind:
        case "workbook":
            openpyxl.Workbook().save(path)
        case "zip":
            with zipfile.ZipFile(path, "w") as z:
                z.writestr("notes.txt", "hello")
        case _:
            path.write_text("hello")
    return path


@pytest.mark.parametrize(
    ("kind", "message"),
    [
        ("workbook", "xl/workbook.xml is not a presentation"),
        ("zip", "_rels/.rels is missing"),
        ("text", "File is not a zip file"),
    ],
    ids=["workbook", "zip-without-parts", "not-a-zip"],
)
def test_a_source_that_is_not_a_deck_exits_2(tmp_path: Path, capsys: pytest.CaptureFixture[str], kind: str, message: str) -> None:
    source = not_a_deck(tmp_path / "source.pptx", kind)
    cs = changeset(source, [change("c1", REVENUE)])

    assert run(["validate", cs], capsys) == (2, "", f"error: cannot read deck {source}: {message}\n")

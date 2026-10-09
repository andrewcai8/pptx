from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from test_changeset import build_deck, change, changeset, run

LAYOUTS = (
    'layouts\n'
    '  "Title Slide": 2 "Title 1" (ctrTitle), 3 "Subtitle 2" (subTitle)\n'
    '  "Title and Content": 2 "Title 1" (title), 3 "Content Placeholder 2" (obj)\n'
    '  "Section Header": 2 "Title 1" (title), 3 "Text Placeholder 2" (body)\n'
    '  "Two Content": 2 "Title 1" (title), 3 "Content Placeholder 2" (obj), 4 "Content Placeholder 3" (obj)\n'
    '  "Comparison": 2 "Title 1" (title), 3 "Text Placeholder 2" (body), 4 "Content Placeholder 3" (obj), '
    '5 "Text Placeholder 4" (body), 6 "Content Placeholder 5" (obj)\n'
    '  "Title Only": 2 "Title 1" (title)\n'
    '  "Blank": none\n'
    '  "Content with Caption": 2 "Title 1" (title), 3 "Content Placeholder 2" (obj), 4 "Text Placeholder 3" (body)\n'
    '  "Picture with Caption": 2 "Title 1" (title), 4 "Text Placeholder 3" (body)\n'
    '  "Title and Vertical Text": 2 "Title 1" (title), 3 "Vertical Text Placeholder 2" (body)\n'
    '  "Vertical Title and Text": 2 "Vertical Title 1" (title), 3 "Vertical Text Placeholder 2" (body)\n'
)


def ph(pid: int, name: str, kind: str) -> dict:
    return {"id": pid, "name": name, "type": kind}


def text(sid: int, name: str, *paragraphs: str) -> dict:
    return {"id": sid, "name": name, "kind": "text", "hidden": False, "paragraphs": list(paragraphs)}


TITLE, CONTENT = ph(2, "Title 1", "title"), ph(3, "Content Placeholder 2", "obj")
OUTLINE_LAYOUTS = [
    {"name": "Title Slide", "placeholders": [ph(2, "Title 1", "ctrTitle"), ph(3, "Subtitle 2", "subTitle")]},
    {"name": "Title and Content", "placeholders": [TITLE, CONTENT]},
    {"name": "Section Header", "placeholders": [TITLE, ph(3, "Text Placeholder 2", "body")]},
    {"name": "Two Content", "placeholders": [TITLE, CONTENT, ph(4, "Content Placeholder 3", "obj")]},
    {
        "name": "Comparison",
        "placeholders": [
            TITLE,
            ph(3, "Text Placeholder 2", "body"),
            ph(4, "Content Placeholder 3", "obj"),
            ph(5, "Text Placeholder 4", "body"),
            ph(6, "Content Placeholder 5", "obj"),
        ],
    },
    {"name": "Title Only", "placeholders": [TITLE]},
    {"name": "Blank", "placeholders": []},
    {"name": "Content with Caption", "placeholders": [TITLE, CONTENT, ph(4, "Text Placeholder 3", "body")]},
    {"name": "Picture with Caption", "placeholders": [TITLE, ph(4, "Text Placeholder 3", "body")]},
    {"name": "Title and Vertical Text", "placeholders": [TITLE, ph(3, "Vertical Text Placeholder 2", "body")]},
    {"name": "Vertical Title and Text", "placeholders": [ph(2, "Vertical Title 1", "title"), ph(3, "Vertical Text Placeholder 2", "body")]},
]


@pytest.fixture
def deck(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.chdir(tmp_path)
    build_deck(tmp_path / "deck.pptx")
    return Path("deck.pptx")


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_outline_lists_every_id_and_text_a_changeset_can_address(deck: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert run(["outline", deck, "--json", "outline.json"], capsys) == (
        0,
        f'deck "deck.pptx" sha256 {sha(deck)}\n'
        + LAYOUTS
        + 'slide 1 id 256 layout "Title and Content": "Market outlook"\n'
        '  shape 2 "Title 1" text\n'
        '    p0 "Market outlook"\n'
        '  shape 3 "Content Placeholder 2" text\n'
        '    p0 "Demand grows 4% a year"\n'
        '    p1 "Prices hold"\n'
        '  shape 4 "TextBox 3" text\n'
        '    p0 "Revenue grew 12% in 2025"\n'
        '  shape 5 "Table 4" table 2x2\n'
        '    r0c0 "Year"\n'
        '    r0c1 "Sales"\n'
        '    r1c0 "2025"\n'
        '    r1c1 "$120m"\n'
        'slide 2 id 257 layout "Title Only": "Sales by year"\n'
        '  shape 2 "Title 1" text\n'
        '    p0 "Sales by year"\n'
        '  shape 3 "Chart 2" chart\n'
        '    series 0 "Sales": p0 "2024" 100, p1 "2025" 120\n'
        'slide 3 id 258 layout "Title Only": "Next steps"\n'
        '  shape 2 "Title 1" text\n'
        '    p0 "Next steps"\n'
        '  shape 3 "TextBox 2" text\n'
        '    p0 "+9%"\n',
        "",
    )
    assert json.loads(Path("outline.json").read_text()) == {
        "deck": {"path": "deck.pptx", "sha256": sha(deck)},
        "layouts": OUTLINE_LAYOUTS,
        "slides": [
            {
                "index": 1,
                "id": 256,
                "layout": "Title and Content",
                "title": "Market outlook",
                "shapes": [
                    text(2, "Title 1", "Market outlook"),
                    text(3, "Content Placeholder 2", "Demand grows 4% a year", "Prices hold"),
                    text(4, "TextBox 3", "Revenue grew 12% in 2025"),
                    {"id": 5, "name": "Table 4", "kind": "table", "hidden": False, "rows": [["Year", "Sales"], ["2025", "$120m"]]},
                ],
            },
            {
                "index": 2,
                "id": 257,
                "layout": "Title Only",
                "title": "Sales by year",
                "shapes": [
                    text(2, "Title 1", "Sales by year"),
                    {
                        "id": 3,
                        "name": "Chart 2",
                        "kind": "chart",
                        "hidden": False,
                        "series": [
                            {
                                "index": 0,
                                "name": "Sales",
                                "points": [{"point": 0, "category": "2024", "value": 100}, {"point": 1, "category": "2025", "value": 120}],
                            }
                        ],
                    },
                ],
            },
            {
                "index": 3,
                "id": 258,
                "layout": "Title Only",
                "title": "Next steps",
                "shapes": [text(2, "Title 1", "Next steps"), text(3, "TextBox 2", "+9%")],
            },
        ],
    }


def test_every_paragraph_cell_and_point_the_outline_lists_validates_as_old(deck: Path, capsys: pytest.CaptureFixture[str]) -> None:
    run(["outline", deck, "--json", "outline.json"], capsys)
    outline = json.loads(Path("outline.json").read_text())
    ops = []
    for slide in outline["slides"]:
        for shape in slide["shapes"]:
            at = {"slide": slide["id"], "shape": shape["id"]}
            ops += [{"kind": "replace_text", **at, "old": p, "new": p + "!", "paragraph": i} for i, p in enumerate(shape.get("paragraphs", []))]
            ops += [
                {"kind": "set_cell", **at, "row": r, "col": c, "old": v, "new": v + "!"}
                for r, row in enumerate(shape.get("rows", []))
                for c, v in enumerate(row)
            ]
            ops += [
                {"kind": "set_chart_value", **at, "series": s["index"], "point": p["point"], "old": p["value"], "new": p["value"] + 1}
                for s in shape.get("series", [])
                for p in s["points"]
            ]
    cs = changeset(deck, [change(f"c{i}", op) for i, op in enumerate(ops)])

    assert run(["validate", cs], capsys) == (
        0,
        f"VALID {cs}: 13 changes (7 text-only, 6 structural) on slides 1, 2, 3; 1 ask, 0 flags, 0 held\n"
        "  c0 replace_text slide 1 shape 2 'Title 1': 'Market outlook' -> 'Market outlook!' (text-only)\n"
        "  c1 replace_text slide 1 shape 3 'Content Placeholder 2': 'Demand grows 4% a year' -> 'Demand grows 4% a year!' (text-only)\n"
        "  c2 replace_text slide 1 shape 3 'Content Placeholder 2': 'Prices hold' -> 'Prices hold!' (text-only)\n"
        "  c3 replace_text slide 1 shape 4 'TextBox 3': 'Revenue grew 12% in 2025' -> 'Revenue grew 12% in 2025!' (text-only)\n"
        "  c4 set_cell slide 1 shape 5 'Table 4': 'Year' -> 'Year!' (structural)\n"
        "  c5 set_cell slide 1 shape 5 'Table 4': 'Sales' -> 'Sales!' (structural)\n"
        "  c6 set_cell slide 1 shape 5 'Table 4': '2025' -> '2025!' (structural)\n"
        "  c7 set_cell slide 1 shape 5 'Table 4': '$120m' -> '$120m!' (structural)\n"
        "  c8 replace_text slide 2 shape 2 'Title 1': 'Sales by year' -> 'Sales by year!' (text-only)\n"
        "  c9 set_chart_value slide 2 shape 3 'Chart 2': '100' -> '101' (structural)\n"
        "  c10 set_chart_value slide 2 shape 3 'Chart 2': '120' -> '121' (structural)\n"
        "  c11 replace_text slide 3 shape 2 'Title 1': 'Next steps' -> 'Next steps!' (text-only)\n"
        "  c12 replace_text slide 3 shape 3 'TextBox 2': '+9%' -> '+9%!' (text-only)\n",
        "",
    )


def test_outline_of_a_file_that_is_not_a_deck_exits_2(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    notes = tmp_path / "notes.pptx"
    notes.write_text("not a deck")

    assert run(["outline", notes], capsys) == (2, "", f"error: cannot read deck {notes}: File is not a zip file\n")


@pytest.mark.parametrize("json_path", ["deck.pptx", "./deck.pptx", "link.pptx"])
def test_outline_refuses_to_write_its_json_over_the_deck(deck: Path, capsys: pytest.CaptureFixture[str], json_path: str) -> None:
    Path("link.pptx").symlink_to("deck.pptx")
    before = sha(deck)

    assert run(["outline", deck, "--json", json_path], capsys) == (
        2,
        "",
        f"error: --json {json_path.removeprefix('./')} is the deck; the engine never writes over it\n",
    )
    assert sha(deck) == before


def test_a_hidden_shape_and_a_merged_cell_are_marked(deck: Path, capsys: pytest.CaptureFixture[str]) -> None:
    prs = Presentation(str(deck))
    frame = next(s for s in prs.slides[0].shapes if s.has_table)
    frame._element.find(f"{qn('p:nvGraphicFramePr')}/{qn('p:cNvPr')}").set("hidden", "1")
    frame.table.cell(1, 1)._tc.set("hMerge", "1")
    prs.save(str(deck))

    code, out, _ = run(["outline", deck, "--json", "outline.json"], capsys)

    assert (code, out.split("slide 2 ")[0].split("  shape 4 ")[1]) == (
        0,
        '"TextBox 3" text\n'
        '    p0 "Revenue grew 12% in 2025"\n'
        '  shape 5 "Table 4" table 2x2 hidden\n'
        '    r0c0 "Year"\n'
        '    r0c1 "Sales"\n'
        '    r1c0 "2025"\n'
        "    r1c1 merged\n",
    )
    assert json.loads(Path("outline.json").read_text())["slides"][0]["shapes"][3] == {
        "id": 5,
        "name": "Table 4",
        "kind": "table",
        "hidden": True,
        "rows": [["Year", "Sales"], ["2025", None]],
    }

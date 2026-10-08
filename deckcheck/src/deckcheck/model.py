from __future__ import annotations

import hashlib
import io
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from lxml.etree import XMLSyntaxError
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.exc import PackageNotFoundError
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.shapes.picture import Picture

TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)

Kind = Literal["title", "body", "text", "table", "chart", "picture", "other"]


class DeckError(Exception):
    pass


@dataclass(frozen=True)
class Run:
    text: str
    font: str
    size_pt: float | None


@dataclass(frozen=True)
class Paragraph:
    text: str
    runs: tuple[Run, ...]
    is_bullet: bool
    level: int


@dataclass(frozen=True)
class Shape:
    name: str
    kind: Kind
    left: int
    top: int
    width: int
    height: int
    paragraphs: tuple[Paragraph, ...]


@dataclass(frozen=True)
class Slide:
    index: int
    layout_name: str
    layout_has_title: bool
    shapes: tuple[Shape, ...]

    @property
    def title_shape(self) -> Shape | None:
        return next((s for s in self.shapes if s.kind == "title"), None)

    @property
    def title(self) -> str:
        shape = self.title_shape
        return " ".join(p.text.strip() for p in shape.paragraphs) if shape else ""


@dataclass(frozen=True)
class Deck:
    path: str
    sha256: str
    slide_width: int
    slide_height: int
    slides: tuple[Slide, ...]


@dataclass(frozen=True)
class Violation:
    rule: str
    slide: int
    shape: str | None
    message: str
    evidence: str


@dataclass(frozen=True)
class ThemeFonts:
    major: str
    minor: str

    def resolve(self, typeface: str | None, is_title: bool) -> str:
        if typeface is None:
            return self.major if is_title else self.minor
        if typeface.startswith("+mj-"):
            return self.major
        if typeface.startswith("+mn-"):
            return self.minor
        return typeface


# x' = x * sx + dx, y' = y * sy + dy: maps group-child coordinates onto the slide.
@dataclass(frozen=True)
class Transform:
    sx: float = 1.0
    dx: float = 0.0
    sy: float = 1.0
    dy: float = 0.0


def load_deck(path: str | Path) -> Deck:
    try:
        data = Path(path).read_bytes()
        prs = Presentation(io.BytesIO(data))
    except (OSError, PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError, XMLSyntaxError) as e:
        raise DeckError(f"cannot read deck {path}: {e}") from e
    themes: dict[int, ThemeFonts] = {}
    slides = []
    for index, slide in enumerate(prs.slides, start=1):
        layout = slide.slide_layout
        theme = themes.setdefault(id(layout.slide_master.part), _theme_fonts(layout.slide_master))
        shapes = tuple(_shapes(slide.shapes, theme, Transform()))
        layout_has_title = any(ph.placeholder_format.type in TITLE_TYPES for ph in layout.placeholders)
        slides.append(Slide(index, layout.name, layout_has_title, shapes))
    return Deck(
        path=str(path),
        sha256=hashlib.sha256(data).hexdigest(),
        slide_width=int(prs.slide_width),
        slide_height=int(prs.slide_height),
        slides=tuple(slides),
    )


def _theme_fonts(master) -> ThemeFonts:
    theme = parse_xml(master.part.part_related_by(RT.THEME).blob)

    def latin(which: str) -> str:
        found = theme.xpath(f".//a:fontScheme/a:{which}/a:latin/@typeface")
        return str(found[0]) if found else ""

    return ThemeFonts(major=latin("majorFont"), minor=latin("minorFont"))


def _shapes(shapes, theme: ThemeFonts, t: Transform):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes, theme, _compose(t, shape))
            continue
        kind = _kind(shape)
        left, top = shape.left or 0, shape.top or 0
        width, height = shape.width or 0, shape.height or 0
        yield Shape(
            name=shape.name,
            kind=kind,
            left=round(left * t.sx + t.dx),
            top=round(top * t.sy + t.dy),
            width=round(width * t.sx),
            height=round(height * t.sy),
            paragraphs=tuple(_paragraphs(shape, kind, theme)),
        )


def _compose(t: Transform, group) -> Transform:
    xfrm = group._element.grpSpPr.find(qn("a:xfrm"))
    if xfrm is None:
        return t
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    ch_off, ch_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
    if off is None or ext is None or ch_off is None or ch_ext is None:
        return t
    sx = int(ext.get("cx")) / (int(ch_ext.get("cx")) or 1)
    sy = int(ext.get("cy")) / (int(ch_ext.get("cy")) or 1)
    dx = int(off.get("x")) - int(ch_off.get("x")) * sx
    dy = int(off.get("y")) - int(ch_off.get("y")) * sy
    return Transform(sx=sx * t.sx, dx=dx * t.sx + t.dx, sy=sy * t.sy, dy=dy * t.sy + t.dy)


def _kind(shape) -> Kind:
    if shape.has_chart:
        return "chart"
    if shape.has_table:
        return "table"
    if isinstance(shape, Picture):
        return "picture"
    if shape.is_placeholder:
        ph_type = shape.placeholder_format.type
        if ph_type in TITLE_TYPES:
            return "title"
        if ph_type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT):
            return "body"
    if shape.has_text_frame:
        return "text"
    return "other"


def _paragraphs(shape, kind: Kind, theme: ThemeFonts):
    if kind == "table":
        frames = [cell.text_frame for cell in shape.table.iter_cells()]
    elif shape.has_text_frame:
        frames = [shape.text_frame]
    else:
        frames = []
    for frame in frames:
        for p in frame.paragraphs:
            if not p.text.strip():
                continue
            runs = tuple(
                Run(
                    text=r.text,
                    font=theme.resolve(r.font.name, kind == "title"),
                    size_pt=r.font.size.pt if r.font.size is not None else None,
                )
                for r in p.runs
            )
            yield Paragraph(text=p.text, runs=runs, is_bullet=_is_bullet(p, kind), level=p.level)


def _is_bullet(p, kind: Kind) -> bool:
    pPr = p._p.pPr
    if pPr is not None:
        if pPr.find(qn("a:buChar")) is not None or pPr.find(qn("a:buAutoNum")) is not None:
            return True
        if pPr.find(qn("a:buNone")) is not None:
            return False
    return kind == "body"

from __future__ import annotations

import io
import posixpath
import re
import zipfile
from dataclasses import dataclass

import openpyxl
from lxml import etree
from pptx.oxml.ns import qn
from pyxlsb import open_workbook

from deckcheck.package import FIXED_TIME, Package

C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
S_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R_ID = qn("r:id")
RT_PACKAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package"
RT_DOCUMENT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
RANGE = re.compile(r"^(?:'((?:[^']|'')+)'|([^'!]+))!\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?$")
PINNED_TIME = b"1980-01-01T00:00:00Z"


class ChartError(ValueError):
    pass


@dataclass(frozen=True)
class Cell:
    sheet: str
    ref: str


@dataclass(frozen=True)
class Point:
    """One chart point: its cached value in the chart part and its cell in the embedded workbook."""

    chart: str
    cache: etree._Element
    value: float
    workbook_rid: str
    workbook: str
    cell: Cell


def number_text(value: float) -> str:
    return str(int(value)) if value.is_integer() and abs(value) < 1e15 else repr(value)


def locate_point(pkg: Package, slide_part: str, frame: etree._Element, series: int, point: int) -> Point:
    ref = frame.find(f".//{{{C_NS}}}chart")
    chart = pkg.related(slide_part, ref.get(R_ID)) if ref is not None else None
    if chart is None:
        raise ChartError("the shape is not a chart")
    space = pkg.xml(chart)
    sers = list(space.iter(f"{{{C_NS}}}ser"))
    if series >= len(sers):
        raise ChartError(f"the chart has {len(sers)} series, numbered 0 to {len(sers) - 1}")
    num_ref = sers[series].find(f"{{{C_NS}}}val/{{{C_NS}}}numRef")
    if num_ref is None:
        raise ChartError(f"series {series} has no numeric values (c:val/c:numRef)")
    pts = {int(pt.get("idx")): pt for pt in num_ref.iterfind(f"{{{C_NS}}}numCache/{{{C_NS}}}pt")}
    if point not in pts:
        raise ChartError(f"series {series} has no point {point}; its points are {', '.join(map(str, sorted(pts)))}")
    cache = pts[point].find(f"{{{C_NS}}}v")
    try:
        value = float(cache.text)
    except (TypeError, ValueError):
        raise ChartError(f"series {series} point {point} caches {cache.text!r}, not a number") from None
    cell = _cell(num_ref.findtext(f"{{{C_NS}}}f") or "", point)
    rid, workbook = _workbook(pkg, chart, space)
    _check_cell(pkg.blob(workbook), workbook, cell)
    return Point(chart, cache, value, rid, workbook, cell)


def _cell(formula: str, point: int) -> Cell:
    m = RANGE.match(formula.strip())
    if not m:
        raise ChartError(f"the series range {formula!r} is not one row or column of a sheet, so no workbook cell matches")
    sheet = (m.group(1) or "").replace("''", "'") or m.group(2)
    col, row = _col(m.group(3)), int(m.group(4))
    end_col, end_row = (_col(m.group(5)), int(m.group(6))) if m.group(5) else (col, row)
    if col != end_col and row != end_row:
        raise ChartError(f"the series range {formula!r} spans rows and columns, so no workbook cell matches")
    length = max(end_col - col, end_row - row) + 1
    if point >= length:
        raise ChartError(f"point {point} is outside the series range {formula!r}")
    if col == end_col and row != end_row:
        return Cell(sheet, f"{_letters(col)}{row + point}")
    return Cell(sheet, f"{_letters(col + point)}{row}")


def _col(letters: str) -> int:
    n = 0
    for c in letters:
        n = n * 26 + ord(c) - 64
    return n


def _letters(col: int) -> str:
    out = ""
    while col:
        col, rem = divmod(col - 1, 26)
        out = chr(65 + rem) + out
    return out


def _workbook(pkg: Package, chart: str, space: etree._Element) -> tuple[str, str]:
    ext = space.find(f"{{{C_NS}}}externalData")
    rel = pkg.rel(chart, ext.get(R_ID)) if ext is not None else None
    if rel is None:
        raise ChartError("the chart has no embedded workbook")
    if rel.get("TargetMode") == "External" or rel.get("Type") != RT_PACKAGE:
        raise ChartError("the chart's workbook is linked or an OLE object, not an embedded workbook")
    part = pkg.related(chart, ext.get(R_ID))
    if not pkg.has(part) or posixpath.splitext(part)[1] not in (".xlsx", ".xlsb"):
        raise ChartError(f"the chart's workbook {part} is not an embedded .xlsx or .xlsb")
    return ext.get(R_ID), part


XLSX_HELD = {"s": "holds text", "str": "holds text", "inlineStr": "holds text", "b": "holds a true/false value", "e": "holds an error", "d": "holds a date"}


def _check_cell(blob: bytes, part: str, cell: Cell) -> None:
    """One rule for both workbook formats: the cell exists and holds a number, so writing one keeps it true."""
    if part.endswith(".xlsb"):
        values = _xlsb_values(blob)
        if cell.sheet not in values:
            raise ChartError(f"the workbook has no sheet {cell.sheet!r}")
        held = _xlsb_held(values[cell.sheet].get(_coordinates(cell.ref)))
    else:
        held = _xlsx_held(_xlsx_cell(Package(blob), cell))
    if held:
        raise ChartError(f"workbook cell {cell.sheet}!{cell.ref} {held}, not a number")


def _xlsx_held(c: etree._Element | None) -> str | None:
    if c is None:
        return "is blank"
    if c.find(f"{{{S_NS}}}f") is not None:
        return "is a formula"
    if c.get("t", "n") != "n":
        return XLSX_HELD.get(c.get("t"), f"holds a {c.get('t')!r} value")
    if c.find(f"{{{S_NS}}}v") is None:
        return "is blank"
    try:
        float(c.findtext(f"{{{S_NS}}}v"))
    except ValueError:
        return "holds text"
    return None


def _xlsb_held(value: object) -> str | None:
    match value:
        case None:
            return "is blank"
        case bool():
            return "holds a true/false value"
        case int() | float():
            return None
        case str():
            return "holds text"
    return f"holds a {type(value).__name__} value"


def _coordinates(ref: str) -> tuple[int, int]:
    m = re.fullmatch(r"([A-Z]+)(\d+)", ref)
    return int(m.group(2)) - 1, _col(m.group(1)) - 1


def _xlsx_cell(book: Package, cell: Cell) -> etree._Element | None:
    rels = book.xml("_rels/.rels")
    main = next((r.get("Target").lstrip("/") for r in rels if r.get("Type") == RT_DOCUMENT), None)
    if main is None:
        raise ChartError("the embedded workbook has no workbook part")
    sheet = next((s for s in book.xml(main).iter(f"{{{S_NS}}}sheet") if s.get("name") == cell.sheet), None)
    if sheet is None:
        raise ChartError(f"the workbook has no sheet {cell.sheet!r}")
    return book.xml(book.related(main, sheet.get(R_ID))).find(f".//{{{S_NS}}}c[@r='{cell.ref}']")


def _xlsb_values(blob: bytes) -> dict[str, dict[tuple[int, int], object]]:
    with open_workbook(io.BytesIO(blob)) as wb:
        out = {}
        for name in wb.sheets:
            with wb.get_sheet(name) as sheet:
                out[name] = {(c.r, c.c): c.v for row in sheet.rows(sparse=True) for c in row if c.v is not None}
        return out


def set_point(pkg: Package, point: Point, value: float) -> None:
    """locate_point proved the cell holds a number, so it has a <v> and no type that contradicts one."""
    point.cache.text = number_text(value)
    part = pkg.related(point.chart, point.workbook_rid)
    if part.endswith(".xlsb"):
        part = _convert(pkg, point.chart, point.workbook_rid, part)
    book = Package(pkg.blob(part))
    _xlsx_cell(book, point.cell).find(f"{{{S_NS}}}v").text = number_text(value)
    pkg.put(part, book.to_bytes())


# openpyxl cannot read .xlsb and PowerPoint reloads a chart from its workbook on Edit Data, so an .xlsb
# workbook is rewritten once as an .xlsx holding the same values. Formulas and formatting do not survive.
def _convert(pkg: Package, chart: str, rid: str, part: str) -> str:
    out = openpyxl.Workbook()
    out.remove(out.active)
    for name, cells in _xlsb_values(pkg.blob(part)).items():
        sheet = out.create_sheet(name)
        for (r, c), v in sorted(cells.items()):
            sheet.cell(row=r + 1, column=c + 1, value=int(v) if isinstance(v, float) and v.is_integer() else v)
    buf = io.BytesIO()
    out.save(buf)
    stem = posixpath.splitext(part)[0]
    target = f"{stem}.xlsx"
    n = 1
    while pkg.has(target):
        n += 1
        target = f"{stem}_{n}.xlsx"
    pkg.put(target, _pinned(buf.getvalue()), XLSX_TYPE)
    pkg.retarget(chart, rid, target)
    return target


def _pinned(data: bytes) -> bytes:
    """openpyxl stamps the zip members and the document's created and modified times with the clock."""
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info)
            if info.filename == "docProps/core.xml":
                body = re.sub(rb"(<dcterms:(created|modified)[^>]*>)[^<]*", rb"\g<1>" + PINNED_TIME, body)
            dst.writestr(zipfile.ZipInfo(info.filename, FIXED_TIME), body, zipfile.ZIP_DEFLATED)
    return buf.getvalue()

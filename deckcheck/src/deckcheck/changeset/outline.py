from __future__ import annotations

import hashlib
import json
import math
from typing import Annotated, Literal

from lxml import etree
from pptx.oxml.ns import qn
from pydantic import Field

from deckcheck.changeset import chart, slides, text
from deckcheck.changeset.model import Source, Wire
from deckcheck.model import DeckError
from deckcheck.package import Package, PartError

KINDS = {
    qn("p:sp"): "shape",
    qn("p:pic"): "picture",
    qn("p:grpSp"): "group",
    qn("p:cxnSp"): "connector",
    qn("p:graphicFrame"): "object",
    qn("p:contentPart"): "content",
}
TRUE = ("1", "true")


class Placeholder(Wire):
    id: int
    name: str
    type: str


class Layout(Wire):
    name: str
    placeholders: tuple[Placeholder, ...]


class Shape(Wire):
    id: int
    name: str
    kind: Literal["shape", "picture", "group", "connector", "object", "content"]
    hidden: bool


class TextShape(Shape):
    kind: Literal["text"]
    paragraphs: tuple[str, ...] = Field(description="text.paragraph_text of every a:p, blank ones included")


class TableShape(Shape):
    kind: Literal["table"]
    rows: tuple[tuple[str | None, ...], ...] = Field(description="text.cell_text of every a:tc; null for a merged cell")


class Point(Wire):
    point: int
    category: str | None
    value: int | float


class Series(Wire):
    index: int = Field(description="position of the c:ser in the chart part, as set_chart_value.series counts it")
    name: str | None
    points: tuple[Point, ...]


class ChartShape(Shape):
    kind: Literal["chart"]
    series: tuple[Series, ...]


AnyShape = Annotated[TextShape | TableShape | ChartShape | Shape, Field(discriminator="kind")]


class Slide(Wire):
    index: int
    id: int
    layout: str
    title: str
    shapes: tuple[AnyShape, ...]


class Outline(Wire):
    deck: Source
    layouts: tuple[Layout, ...]
    slides: tuple[Slide, ...]


def read(data: bytes, path: str) -> Outline:
    try:
        pkg = Package(data)
        deck = slides.read_deck(pkg)
        layouts = tuple(
            Layout(name=layout.name, placeholders=tuple(Placeholder(id=p.id, name=p.name, type=p.type) for p in layout.placeholders if p.holds_text))
            for layout in deck.layouts
        )
        return Outline(
            deck=Source(path=path, sha256=hashlib.sha256(data).hexdigest()),
            layouts=layouts,
            slides=tuple(_slide(pkg, s) for s in deck.slides),
        )
    except (PartError, StopIteration) as e:
        raise DeckError(f"cannot read deck {path}: {e}") from e


def _slide(pkg: Package, slide: slides.SourceSlide) -> Slide:
    xml = pkg.xml(slide.part)
    layout = pkg.xml(slides.related_by_type(pkg, slide.part, slides.RT_LAYOUT)).find(qn("p:cSld")).get("name", "")
    shapes = [_shape(pkg, slide.part, el) for el in slides.shapes(slides.shape_tree(xml))]
    return Slide(index=slide.index, id=slide.id, layout=layout, title=slides.title(xml), shapes=tuple(shapes))


def _kind(el: etree._Element) -> str:
    match KINDS[el.tag]:
        case "shape" if el.find(qn("p:txBody")) is not None:
            return "text"
        case "object" if el.find(f".//{qn('a:tbl')}") is not None:
            return "table"
        case "object" if el.find(f".//{{{chart.C_NS}}}chart") is not None:
            return "chart"
        case kind:
            return kind


def _shape(pkg: Package, part: str, el: etree._Element) -> Shape:
    hidden = el[0].find(qn("p:cNvPr")).get("hidden") in TRUE
    common = {"id": slides.shape_id(el), "name": slides.shape_name(el), "hidden": hidden}
    match _kind(el):
        case "text":
            paragraphs = tuple(text.paragraph_text(p) for p in el.find(qn("p:txBody")).iterfind(qn("a:p")))
            return TextShape(kind="text", paragraphs=paragraphs, **common)
        case "table":
            tbl = el.find(f".//{qn('a:tbl')}")
            rows = tuple(
                tuple(None if text.merged(tc) else text.cell_text(tc) for tc in tr.iterfind(qn("a:tc"))) for tr in tbl.iterfind(qn("a:tr"))
            )
            return TableShape(kind="table", rows=rows, **common)
        case "chart":
            return ChartShape(kind="chart", series=_series(pkg, part, el), **common)
        case kind:
            return Shape(kind=kind, **common)


def _series(pkg: Package, part: str, frame: etree._Element) -> tuple[Series, ...]:
    chart_part = chart.chart_part(pkg, part, frame)
    if chart_part is None or not pkg.has(chart_part):
        return ()
    return tuple(_one_series(i, ser) for i, ser in enumerate(chart.series_of(pkg.xml(chart_part))))


def _one_series(index: int, ser: etree._Element) -> Series:
    names = ser.xpath("c:tx/c:strRef/c:strCache/c:pt/c:v | c:tx/c:v")
    cats = {int(pt.get("idx")): pt.findtext(f"{{{chart.C_NS}}}v") for pt in ser.xpath("c:cat/*/c:strCache/c:pt | c:cat/*/c:numCache/c:pt")}
    num_ref = ser.find(f"{{{chart.C_NS}}}val/{{{chart.C_NS}}}numRef")
    cached = chart.cached_points(num_ref) if num_ref is not None else {}
    points = []
    for idx, pt in sorted(cached.items()):
        value = _number(pt)
        if value is not None:
            points.append(Point(point=idx, category=cats.get(idx), value=value))
    return Series(index=index, name=(names[0].text or "") if names else None, points=tuple(points))


def _number(pt: etree._Element) -> int | float | None:
    value = chart.cached_value(pt)
    return json.loads(chart.number_text(value)) if value is not None and math.isfinite(value) else None


def render(outline: Outline) -> str:
    lines = [f"deck {_q(outline.deck.path)} sha256 {outline.deck.sha256}", "layouts"]
    for layout in outline.layouts:
        placeholders = ", ".join(f"{p.id} {_q(p.name)} ({p.type})" for p in layout.placeholders)
        lines.append(f"  {_q(layout.name)}: {placeholders or 'none'}")
    for slide in outline.slides:
        lines.append(f"slide {slide.index} id {slide.id} layout {_q(slide.layout)}: {_q(slide.title)}")
        for shape in slide.shapes:
            lines += _shape_lines(shape)
    return "\n".join(lines) + "\n"


def _shape_lines(shape: Shape) -> list[str]:
    head = f"  shape {shape.id} {_q(shape.name)} {shape.kind}"
    hidden = " hidden" if shape.hidden else ""
    match shape:
        case TextShape():
            return [head + hidden, *(f"    p{i} {_q(p)}" for i, p in enumerate(shape.paragraphs))]
        case TableShape():
            size = f" {len(shape.rows)}x{max(map(len, shape.rows), default=0)}"
            cells = [f"    r{r}c{c} {'merged' if v is None else _q(v)}" for r, row in enumerate(shape.rows) for c, v in enumerate(row)]
            return [head + size + hidden, *cells]
        case ChartShape():
            return [head + hidden, *(f"    series {s.index} {_q(s.name)}: {_points(s)}" for s in shape.series)]
    return [head + hidden]


def _points(series: Series) -> str:
    return ", ".join(f"p{p.point} {_q(p.category)} {json.dumps(p.value)}" for p in series.points) or "none"


def _q(value: str | None) -> str:
    return json.dumps(value, ensure_ascii=False)

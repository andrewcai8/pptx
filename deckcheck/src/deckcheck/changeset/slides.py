from __future__ import annotations

import re
from collections.abc import Collection, Iterator, Mapping, Sequence
from dataclasses import dataclass

from lxml import etree
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn

from deckcheck.package import Package, PartError

RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RT_DOCUMENT, RT_SLIDE = f"{RT}/officeDocument", f"{RT}/slide"
RT_LAYOUT, RT_MASTER = f"{RT}/slideLayout", f"{RT}/slideMaster"
SLIDE_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
SECTIONS = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
MAX_SLIDE_ID = 2147483647
R_ID = qn("r:id")
SHAPE_TAGS = frozenset(qn(t) for t in ("p:sp", "p:pic", "p:graphicFrame", "p:cxnSp", "p:grpSp", "p:contentPart"))
TITLE_TYPES = frozenset({"title", "ctrTitle"})
UNCLONED = frozenset({"dt", "ftr", "sldNum"})
TEXT_TYPES = frozenset({"title", "ctrTitle", "subTitle", "body", "obj"})
NEW_SLIDE = (
    '<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    '<p:cSld><p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
    "<p:grpSpPr/></p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sld>"
)


@dataclass(frozen=True)
class SourceSlide:
    id: int
    rid: str
    part: str
    index: int

    @property
    def name(self) -> str:
        return f"slide {self.index} (id {self.id})"


@dataclass(frozen=True)
class Placeholder:
    id: int
    name: str
    type: str
    xml: etree._Element

    @property
    def holds_text(self) -> bool:
        return self.type in TEXT_TYPES


@dataclass(frozen=True)
class Layout:
    name: str
    part: str
    master: str
    placeholders: tuple[Placeholder, ...]


@dataclass(frozen=True)
class Deck:
    presentation: str
    slides: tuple[SourceSlide, ...]
    layouts: tuple[Layout, ...]
    last_part: int

    def slide(self, slide_id: int) -> SourceSlide | None:
        return next((s for s in self.slides if s.id == slide_id), None)


def read_deck(pkg: Package) -> Deck:
    pres = next(r for r in pkg.xml("_rels/.rels") if r.get("Type") == RT_DOCUMENT).get("Target").lstrip("/")
    sld_ids = pkg.xml(pres).find(qn("p:sldIdLst"))
    slides = tuple(
        SourceSlide(int(s.get("id")), s.get(R_ID), pkg.related(pres, s.get(R_ID)), i)
        for i, s in enumerate(sld_ids if sld_ids is not None else (), start=1)
    )
    masters: list[str] = []
    for s in slides:
        master = related_by_type(pkg, related_by_type(pkg, s.part, RT_LAYOUT), RT_MASTER)
        if master not in masters:
            masters.append(master)
    layouts = tuple(_layout(pkg, part, m) for m in masters for part in _master_layouts(pkg, m))
    numbers = [int(m.group(1)) for n in pkg.names() if (m := re.fullmatch(r"ppt/slides/slide(\d+)\.xml", n))]
    return Deck(pres, slides, layouts, max(numbers, default=0))


def related_by_type(pkg: Package, part: str, reltype: str) -> str:
    rid = next((r.get("Id") for r in pkg.rels(part) if r.get("Type") == reltype), None)
    if rid is None:
        raise PartError(f"{part} has no {reltype.rsplit('/', 1)[-1]} relationship")
    return pkg.related(part, rid)


def _master_layouts(pkg: Package, master: str) -> list[str]:
    ids = pkg.xml(master).find(qn("p:sldLayoutIdLst"))
    return [pkg.related(master, i.get(R_ID)) for i in (ids if ids is not None else ())]


def _layout(pkg: Package, part: str, master: str) -> Layout:
    xml = pkg.xml(part)
    placeholders = []
    for sp in xml.iterfind(f"{qn('p:cSld')}/{qn('p:spTree')}/{qn('p:sp')}"):
        ph = sp.find(f"{qn('p:nvSpPr')}/{qn('p:nvPr')}/{qn('p:ph')}")
        if ph is None or ph.get("type", "obj") in UNCLONED:
            continue
        c_nv = sp.find(f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}")
        placeholders.append(Placeholder(int(c_nv.get("id")), c_nv.get("name", ""), ph.get("type", "obj"), sp))
    return Layout(xml.find(qn("p:cSld")).get("name", ""), part, master, tuple(placeholders))


def shapes(container: etree._Element) -> Iterator[etree._Element]:
    for child in container:
        if child.tag in SHAPE_TAGS:
            yield child
            if child.tag == qn("p:grpSp"):
                yield from shapes(child)


def shape_tree(slide_xml: etree._Element) -> etree._Element:
    return slide_xml.find(f"{qn('p:cSld')}/{qn('p:spTree')}")


def shape_id(shape: etree._Element) -> int:
    return int(shape[0].find(qn("p:cNvPr")).get("id"))


def shape_name(shape: etree._Element) -> str:
    return shape[0].find(qn("p:cNvPr")).get("name", "")


def title(slide_xml: etree._Element) -> str:
    for shape in shapes(shape_tree(slide_xml)):
        ph = shape.find(f"{qn('p:nvSpPr')}/{qn('p:nvPr')}/{qn('p:ph')}")
        body = shape.find(qn("p:txBody"))
        if ph is not None and ph.get("type") in TITLE_TYPES and body is not None:
            return " ".join(p.xpath("string(.)").strip() for p in body.iterfind(qn("a:p"))).strip()
    return ""


def new_slide(pkg: Package, deck: Deck, layout: Layout, part: str) -> str:
    sld = parse_xml(NEW_SLIDE)
    tree = shape_tree(sld)
    for ph in layout.placeholders:
        tree.append(_placeholder(ph))
    pkg.put_xml(part, sld, SLIDE_TYPE)
    pkg.relate(part, RT_LAYOUT, layout.part)
    return pkg.relate(deck.presentation, RT_SLIDE, part)


def _placeholder(ph: Placeholder) -> etree._Element:
    sp = parse_xml(
        '<p:sp xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        '<p:nvSpPr><p:cNvPr id="0" name=""/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph/></p:nvPr>'
        "</p:nvSpPr><p:spPr/></p:sp>"
    )
    c_nv = sp.find(f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}")
    c_nv.set("id", str(ph.id))
    c_nv.set("name", ph.name)
    source = ph.xml.find(f"{qn('p:nvSpPr')}/{qn('p:nvPr')}/{qn('p:ph')}")
    target = sp.find(f"{qn('p:nvSpPr')}/{qn('p:nvPr')}/{qn('p:ph')}")
    for attr in ("type", "orient", "sz", "idx"):
        if source.get(attr) is not None:
            target.set(attr, source.get(attr))
    if ph.holds_text:
        sp.append(
            parse_xml(
                '<p:txBody xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
                ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
                "<a:bodyPr/><a:lstStyle/><a:p/></p:txBody>"
            )
        )
    return sp


def final_order(source: Sequence[int], deleted: Collection[int], placements: Sequence[tuple[int, int | None]]) -> list[int]:
    placed = {item for item, _ in placements}
    order = [s for s in source if s not in deleted and s not in placed]
    tail: dict[int | None, int] = {}
    todo = list(placements)
    while todo:
        waiting = [(item, anchor) for item, anchor in todo if anchor is not None and anchor not in order]
        if len(waiting) == len(todo):
            raise ValueError(f"slides {sorted(item for item, _ in todo)} are placed after each other in a cycle")
        for item, anchor in todo:
            if (item, anchor) in waiting:
                continue
            after = tail.get(anchor, anchor)
            order.insert(0 if after is None else order.index(after) + 1, item)
            tail[anchor] = item
        todo = waiting
    return order


def write_order(pkg: Package, deck: Deck, order: Sequence[tuple[int, str]], anchors: Mapping[int, int | None]) -> None:
    root = pkg.xml(deck.presentation)
    lst = root.find(qn("p:sldIdLst"))
    existing = {int(s.get("id")): s for s in lst}
    for s in list(lst):
        lst.remove(s)
    for sid, rid in order:
        el = existing.get(sid)
        if el is None:
            el = etree.SubElement(lst, qn("p:sldId"))
            el.set("id", str(sid))
            el.set(R_ID, rid)
        else:
            lst.append(el)
    kept = {rid for _, rid in order}
    for s in deck.slides:
        if s.rid not in kept:
            pkg.unrelate(deck.presentation, s.rid)
    for sld in root.iterfind(f"{qn('p:custShowLst')}/{qn('p:custShow')}/{qn('p:sldLst')}/{qn('p:sld')}"):
        if sld.get(R_ID) not in kept:
            sld.getparent().remove(sld)
    _sections(root, [sid for sid, _ in order], anchors)


def _sections(root: etree._Element, order: Sequence[int], anchors: Mapping[int, int | None]) -> None:
    lst = next((e[0] for e in root.iterfind(f"{qn('p:extLst')}/{qn('p:ext')}") if e.get("uri") == SECTIONS and len(e)), None)
    if lst is None:
        return
    sections = lst.findall(f"{{{P14}}}section")
    if not sections:
        return
    home = {int(s.get("id")): sec for sec in sections for s in sec.iterfind(f"{{{P14}}}sldIdLst/{{{P14}}}sldId")}
    assigned: dict[int, etree._Element] = {}
    for sid in order:
        if sid in anchors:
            anchor = anchors[sid]
            assigned[sid] = assigned.get(anchor, sections[0]) if anchor is not None else sections[0]
        else:
            assigned[sid] = home.get(sid, sections[0])
    existing = {int(s.get("id")): s for sec in sections for s in sec.iterfind(f"{{{P14}}}sldIdLst/{{{P14}}}sldId")}
    for sec in sections:
        ids = sec.find(f"{{{P14}}}sldIdLst")
        if ids is None:
            ids = etree.SubElement(sec, f"{{{P14}}}sldIdLst")
        for s in list(ids):
            ids.remove(s)
        for sid in order:
            if assigned[sid] is sec:
                el = existing.get(sid)
                if el is None:
                    el = etree.Element(f"{{{P14}}}sldId", id=str(sid))
                ids.append(el)

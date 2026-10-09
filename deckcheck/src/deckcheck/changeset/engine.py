from __future__ import annotations

import hashlib
import json
from abc import ABC, abstractmethod
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import ClassVar, Literal

from lxml import etree
from pptx.oxml.ns import qn
from pydantic import ValidationError

from deckcheck.changeset import chart, slides, text
from deckcheck.changeset.model import (
    OP_KINDS,
    AddSlide,
    Box,
    Change,
    ChangeSet,
    DeleteSlide,
    Edited,
    Executed,
    FillPlaceholder,
    InsertParagraph,
    MoveSlide,
    Op,
    Paragraph,
    ReplaceText,
    Review,
    ReviewChange,
    ReviewShape,
    ReviewSlide,
    SetCell,
    SetChartValue,
)
from deckcheck.model import DeckError, open_presentation, read_deck
from deckcheck.package import Package, PartError

DERIVED = {
    "before": "the engine reads this from the source deck; remove it",
    "after": "the engine derives this from the op; remove it",
    "structural": "the engine derives this from the op kind; remove it",
    "depends_on": "the engine derives this from the op; remove it",
}
DECISION_FORMAT = 'must be "pending", "keep_new", "keep_old" or {"edited": "<text>"}'


@dataclass(frozen=True)
class Problem:
    where: str
    message: str


class Invalid(Exception):
    def __init__(self, problems: Sequence[Problem]) -> None:
        super().__init__(f"{len(problems)} problems")
        self.problems = tuple(problems)


class Undecided(Exception):
    def __init__(self, pending: Sequence[str]) -> None:
        super().__init__(f"{len(pending)} pending")
        self.pending = tuple(pending)


class Miss(Exception):
    def __init__(self, field: str, message: str) -> None:
        super().__init__(message)
        self.field = field
        self.message = message


FIELD_NOTE = "a think-cell label: its field format is rewritten too, so a refresh keeps the new text"


@dataclass(frozen=True)
class Item:
    """One change as checked against the source deck. Everything here is read from the source, not the maker."""

    change: Change
    target: Target
    slide: int | str
    source_index: int | None
    shape: tuple[int, str] | None = None
    before: str | None = None
    after: str | None = None
    quote: tuple[str, str] | None = None
    span: tuple[int, int] | None = None
    depends_on: str | None = None
    notes: tuple[str, ...] = ()


@dataclass(frozen=True)
class Claim:
    """Something only one change may write, keyed for comparison, and how a second claimant's problem reads."""

    key: tuple
    field: str
    taken: str


@dataclass
class Build:
    """The package _build writes into, and what the slide writes collect for the final slide list."""

    pkg: Package
    deck: slides.Deck
    rids: dict[int, str]
    placements: list[tuple[int, int | None]]
    deleted: set[int]
    tails: dict[etree._Element, etree._Element]


class Target(ABC):
    """What one op writes, found on the pristine source deck by locate, the one dispatch on the op. Each kind
    writes itself in its phase, describes itself for check, and names the slide it touches and what it claims,
    so a kind cannot be located and then silently skipped."""

    phase: ClassVar[int] = 1

    @abstractmethod
    def write(self, build: Build) -> None: ...

    @abstractmethod
    def describe(self, change: Change, pkg: Package) -> Item: ...

    def rank(self) -> tuple[int, int]:
        return self.phase, 0

    @property
    def touches(self) -> slides.SourceSlide | None:
        """The source slide this writes on or moves, which no change may delete."""
        return None

    @property
    def follows(self) -> slides.SourceSlide | None:
        """The source slide this places a slide after, which no change may delete."""
        return None

    def claim(self) -> Claim | None:
        return None


@dataclass(frozen=True)
class OnShape(Target):
    slide: slides.SourceSlide
    shape: etree._Element

    @property
    def touches(self) -> slides.SourceSlide:
        return self.slide

    def _item(self, change: Change, **fields: object) -> Item:
        shape = (slides.shape_id(self.shape), slides.shape_name(self.shape))
        return Item(change=change, target=self, slide=self.slide.id, source_index=self.slide.index, shape=shape, **fields)


@dataclass(frozen=True)
class TextAt(OnShape):
    p: etree._Element
    paragraph: int
    at: int
    before: str
    old: str
    new: str
    splice: text.Splice
    phase: ClassVar[int] = 0

    def rank(self) -> tuple[int, int]:
        # Quotes in one paragraph never overlap, so their starts order them. Splice offsets do not: two abutting
        # quotes can both trim to an insertion at their shared edge, and the later quote's must be written first.
        return self.phase, -self.at

    def write(self, build: Build) -> None:
        text.write_splice(self.p, self.splice)

    def describe(self, change: Change, pkg: Package) -> Item:
        end = self.at + len(self.old)
        touched = [x for x in text.atoms(self.p) if x.lo < self.splice.end and x.hi > self.splice.start]
        return self._item(
            change,
            before=self.before,
            after=self.before[: self.at] + self.new + self.before[end:],
            quote=(self.old, self.new),
            span=(self.at, end),
            notes=(FIELD_NOTE,) if any(x.el.tag == text.A_FLD for x in touched) else (),
        )


@dataclass(frozen=True)
class ParagraphAt(OnShape):
    anchor: etree._Element
    new: str

    def write(self, build: Build) -> None:
        p = text.new_paragraph(self.anchor, self.new)
        build.tails.get(self.anchor, self.anchor).addnext(p)
        build.tails[self.anchor] = p

    def describe(self, change: Change, pkg: Package) -> Item:
        return self._item(change, after=self.new)


@dataclass(frozen=True)
class CellAt(OnShape):
    tc: etree._Element
    row: int
    col: int
    new: str

    def write(self, build: Build) -> None:
        text.write_cell(self.tc, self.new)

    def describe(self, change: Change, pkg: Package) -> Item:
        return self._item(change, before=text.cell_text(self.tc), after=self.new)

    def claim(self) -> Claim:
        shape = slides.shape_id(self.shape)
        return Claim(("cell", self.slide.id, shape, self.row, self.col), "op", f"shape {shape} row {self.row} col {self.col} is already changed")


@dataclass(frozen=True)
class PointAt(OnShape):
    point: chart.Point
    series: int
    index: int
    new: float

    def write(self, build: Build) -> None:
        chart.set_point(build.pkg, self.point, self.new)

    def describe(self, change: Change, pkg: Package) -> Item:
        book = self.point.workbook
        notes = (f"workbook cell {self.point.cell.sheet}!{self.point.cell.ref} in {book}",)
        if book.endswith(".xlsb"):
            notes += ("the .xlsb workbook is rewritten as .xlsx, values only",)
        return self._item(change, before=chart.number_text(self.point.value), after=chart.number_text(self.new), notes=notes)

    def claim(self) -> Claim:
        shape = slides.shape_id(self.shape)
        return Claim(
            ("point", self.slide.id, shape, self.series, self.index), "op", f"shape {shape} series {self.series} point {self.index} is already changed"
        )


@dataclass(frozen=True)
class NewSlide(Target):
    layout: slides.Layout
    id: int
    part: str
    after: slides.SourceSlide | None

    @property
    def follows(self) -> slides.SourceSlide | None:
        return self.after

    def write(self, build: Build) -> None:
        build.rids[self.id] = slides.new_slide(build.pkg, build.deck, self.layout, self.part)
        build.placements.append((self.id, self.after.id if self.after else None))

    def describe(self, change: Change, pkg: Package) -> Item:
        where = f"after {self.after.name}" if self.after else "first"
        return Item(change=change, target=self, slide=change.id, source_index=None, after=f"new slide on layout {self.layout.name!r}, {where}")


@dataclass(frozen=True)
class PlaceholderAt(Target):
    """A fill writes on a slide an earlier phase adds, so it writes last."""

    add: NewSlide
    add_id: str
    placeholder: slides.Placeholder
    paragraphs: tuple[Paragraph, ...]
    phase: ClassVar[int] = 2

    def write(self, build: Build) -> None:
        sp = next(s for s in slides.shapes(slides.shape_tree(build.pkg.xml(self.add.part))) if slides.shape_id(s) == self.placeholder.id)
        text.fill(sp.find(qn("p:txBody")), self.paragraphs)

    def describe(self, change: Change, pkg: Package) -> Item:
        return Item(
            change=change,
            target=self,
            slide=self.add_id,
            source_index=None,
            shape=(self.placeholder.id, self.placeholder.name),
            after="\n".join(p.text for p in self.paragraphs),
            depends_on=self.add_id,
        )

    def claim(self) -> Claim:
        ph = self.placeholder.id
        return Claim(("placeholder", self.add_id, ph), "op", f"placeholder {ph} on {self.add_id}'s slide is already changed")


@dataclass(frozen=True)
class SlideAt(Target):
    slide: slides.SourceSlide

    def write(self, build: Build) -> None:
        build.deleted.add(self.slide.id)

    def describe(self, change: Change, pkg: Package) -> Item:
        title = slides.title(pkg.xml(self.slide.part))
        return Item(change=change, target=self, slide=self.slide.id, source_index=self.slide.index, before=title)

    def claim(self) -> Claim:
        return Claim(("delete", self.slide.id), "op.slide", f"{self.slide.name} is already deleted")


@dataclass(frozen=True)
class MoveTo(Target):
    slide: slides.SourceSlide
    after: slides.SourceSlide | None
    previous: slides.SourceSlide | None

    @property
    def touches(self) -> slides.SourceSlide:
        return self.slide

    @property
    def follows(self) -> slides.SourceSlide | None:
        return self.after

    def write(self, build: Build) -> None:
        build.placements.append((self.slide.id, self.after.id if self.after else None))

    def describe(self, change: Change, pkg: Package) -> Item:
        before = f"after {self.previous.name}" if self.previous else "first"
        after = f"after {self.after.name}" if self.after else "first"
        return Item(change=change, target=self, slide=self.slide.id, source_index=self.slide.index, before=before, after=after)

    def claim(self) -> Claim:
        return Claim(("move", self.slide.id), "op.slide", f"{self.slide.name} is already moved")


@dataclass(frozen=True)
class Checked:
    """A ChangeSet proved against its source deck: every id resolves, every quote matches, nothing conflicts,
    and every decision is admissible, so execute cannot fail and apply fails only on a pending decision."""

    path: Path
    changeset: ChangeSet
    source: bytes
    items: tuple[Item, ...]


@dataclass(frozen=True)
class Applied:
    data: bytes
    kept_new: tuple[str, ...]
    edited: tuple[str, ...]
    kept_old: tuple[str, ...]
    dropped: tuple[tuple[str, str], ...]


@dataclass(frozen=True)
class Context:
    """What locating one op needs beyond the deck: every change's op and each add_slide's ordinal among all
    add_slide changes, so an added slide has the same id and part whichever adds a reviewer keeps."""

    ops: Mapping[str, Op]
    ordinals: Mapping[str, int]

    @classmethod
    def of(cls, cs: ChangeSet) -> Context:
        adds = [c.id for c in cs.changes if isinstance(c.op, AddSlide)]
        return cls({c.id: c.op for c in cs.changes}, {cid: i for i, cid in enumerate(adds, start=1)})


def load(path: Path) -> Checked:
    try:
        raw_text = path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeDecodeError) as e:
        raise DeckError(f"cannot read ChangeSet {path}: {e}") from e
    cs = parse(raw_text)
    source = Path(cs.source.path)
    try:
        data = source.read_bytes()
    except OSError as e:
        raise Invalid([Problem("source.path", f"cannot read {source}: {e.strerror}")]) from e
    except ValueError as e:
        raise Invalid([Problem("source.path", f"{cs.source.path!r} is not a path: {e}")]) from e
    sha = hashlib.sha256(data).hexdigest()
    if sha != cs.source.sha256:
        raise Invalid([Problem("source.sha256", f"{source} has sha256 {sha}; the deck changed since the maker read it")])
    return check(cs, data, path)


def parse(raw_text: str) -> ChangeSet:
    try:
        raw = json.loads(raw_text)
    except json.JSONDecodeError as e:
        raise Invalid([Problem("changeset", f"not JSON: {e}")]) from e
    try:
        return ChangeSet.model_validate_json(raw_text)
    except ValidationError as e:
        problems: dict[str, Problem] = {}
        for err in e.errors():
            where = _where(err["loc"], raw)
            problems.setdefault(where, Problem(where, _message(err)))
        raise Invalid(list(problems.values())) from e


def _where(loc: tuple, raw: object) -> str:
    if len(loc) < 2 or loc[0] != "changes" or not isinstance(loc[1], int):
        return ".".join(map(str, loc)) or "changeset"
    change = raw["changes"][loc[1]] if isinstance(raw, dict) and isinstance(raw.get("changes"), list) else None
    cid = change.get("id") if isinstance(change, dict) else None
    name = cid if isinstance(cid, str) and cid else f"changes[{loc[1]}]"
    rest = list(loc[2:])
    if rest[:1] == ["op"] and len(rest) > 1 and rest[1] in OP_KINDS:
        del rest[1]
    if rest[:1] == ["decision"]:
        rest = ["decision"]
    return " ".join([name, ".".join(map(str, rest))]) if rest else name


def _message(err: Mapping) -> str:
    loc = err["loc"]
    field = loc[-1] if loc else ""
    if err["type"] == "extra_forbidden":
        if len(loc) == 3 and loc[0] == "changes" and field in DERIVED:
            return DERIVED[field]
        if len(loc) == 5 and loc[0] == "changes" and loc[2] == "op":
            return f"{loc[3]} has no field {field!r}"
        return "unknown field; remove it"
    if "decision" in err["loc"]:
        return DECISION_FORMAT
    if err["type"] == "missing":
        return "required"
    return err["msg"]


def check(cs: ChangeSet, data: bytes, path: Path = Path("changeset.json")) -> Checked:
    try:
        pkg = Package(data)
        deck = slides.read_deck(pkg)
    except (PartError, StopIteration) as e:
        raise DeckError(f"cannot read deck {cs.source.path}: {e}") from e
    ctx = Context.of(cs)
    problems: list[Problem] = []
    located: list[tuple[Change, Target]] = []
    for change in cs.changes:
        try:
            located.append((change, locate(pkg, deck, ctx, change.id, change.op)))
        except Miss as m:
            problems.append(Problem(f"{change.id} {m.field}", m.message))
            continue
        except PartError as e:
            problems.append(Problem(f"{change.id} op", f"cannot read the deck: {e}"))
            continue
        problems += _decision_problems(pkg, deck, ctx, change)
    problems += _references(cs, deck) + _conflicts(located)
    if problems:
        raise Invalid(problems)
    return Checked(path, cs, data, tuple(t.describe(c, pkg) for c, t in located))


def _decision_problems(pkg: Package, deck: slides.Deck, ctx: Context, change: Change) -> list[Problem]:
    if not isinstance(change.decision, Edited):
        return []
    try:
        op = change.op.edit(change.decision.edited)
        locate(pkg, deck, ctx, change.id, op)
    except ValueError as e:
        return [Problem(f"{change.id} decision", str(e))]
    except Miss as m:
        return [Problem(f"{change.id} decision", f"the edited text cannot be written: {m.message}")]
    return []


def locate(pkg: Package, deck: slides.Deck, ctx: Context, cid: str, op: Op) -> Target:
    """Find what `op` writes on a pristine package, or raise Miss naming the field at fault. check and _build
    both call this, so a ChangeSet that checks cannot fail to build."""
    match op:
        case ReplaceText():
            return _locate_text(pkg, deck, op)
        case InsertParagraph():
            slide, shape, paras = _paragraphs(pkg, deck, op.slide, op.shape)
            if op.after >= len(paras):
                raise Miss("op.after", f"shape {op.shape} {slides.shape_name(shape)!r} has {_count(len(paras))}")
            _no_newline(op.text, "op.text")
            return ParagraphAt(slide, shape, anchor=paras[op.after], new=op.text)
        case SetCell():
            return _locate_cell(pkg, deck, op)
        case SetChartValue():
            slide = _slide(deck, op.slide, "op.slide")
            shape = _shape(pkg, slide, op.shape, "chart")
            try:
                point = chart.locate_point(pkg, slide.part, shape, op.series, op.point)
            except chart.ChartError as e:
                raise Miss("op", f"shape {op.shape} {slides.shape_name(shape)!r}: {e}") from e
            if point.value != op.old:
                raise Miss("op.old", f"series {op.series} point {op.point} caches {chart.number_text(point.value)}, not {chart.number_text(op.old)}")
            return PointAt(slide, shape, point=point, series=op.series, index=op.point, new=op.new)
        case AddSlide():
            return _locate_add(deck, ctx, cid, op)
        case FillPlaceholder():
            return _locate_fill(pkg, deck, ctx, op)
        case DeleteSlide():
            return SlideAt(_slide(deck, op.slide, "op.slide"))
        case MoveSlide():
            slide = _slide(deck, op.slide, "op.slide")
            if op.after == op.slide:
                raise Miss("op.after", "a slide cannot follow itself")
            after = _slide(deck, op.after, "op.after") if op.after is not None else None
            return MoveTo(slide, after, deck.slides[slide.index - 2] if slide.index > 1 else None)
    raise AssertionError(op)


def _locate_text(pkg: Package, deck: slides.Deck, op: ReplaceText) -> TextAt:
    slide, shape, paras = _paragraphs(pkg, deck, op.slide, op.shape)
    name = f"shape {op.shape} {slides.shape_name(shape)!r}"
    texts = [text.paragraph_text(p) for p in paras]
    fed = {i for i, t in enumerate(texts) if "\n" in t}
    if "\n" in op.old and fed:
        raise Miss("op.old", f'{name} holds a line feed ("\\n") inside the text of {_paragraph_list(fed)}, which no op can write; quote the text on one side of it')
    _no_newline(op.old, "op.old")
    _no_newline(op.new, "op.new")
    if op.paragraph is not None and op.paragraph >= len(paras):
        raise Miss("op.paragraph", f"{name} has {_count(len(paras))}")
    hits = [(i, at) for i, t in enumerate(texts) for at in text.occurrences(t, op.old)]
    mine = [(i, at) for i, at in hits if op.paragraph in (None, i)]
    if not mine:
        where = f" paragraph {op.paragraph}" if op.paragraph is not None else ""
        elsewhere = f"; it occurs in {_paragraph_list({i for i, _ in hits})}" if hits else ""
        raise Miss("op.old", f"{op.old!r} is not in {name}{where}{elsewhere}; its text is {text.clip(chr(10).join(texts))!r}")
    if len(mine) > 1:
        raise Miss(
            "op.old",
            f"{op.old!r} occurs {len(mine)} times in {name}, in {_paragraph_list({i for i, _ in mine})}; "
            "quote more of the text or name the paragraph",
        )
    i, at = mine[0]
    try:
        splice = text.plan_splice(paras[i], at, op.old, op.new)
    except text.SpliceError as e:
        raise Miss("op.new", str(e)) from e
    return TextAt(slide, shape, p=paras[i], paragraph=i, at=at, before=texts[i], old=op.old, new=op.new, splice=splice)


def _locate_cell(pkg: Package, deck: slides.Deck, op: SetCell) -> CellAt:
    slide = _slide(deck, op.slide, "op.slide")
    shape = _shape(pkg, slide, op.shape, "table")
    tbl = shape.find(f".//{qn('a:tbl')}") if shape.tag == qn("p:graphicFrame") else None
    if tbl is None:
        raise Miss("op.shape", f"shape {op.shape} {slides.shape_name(shape)!r} is not a table")
    rows = tbl.findall(qn("a:tr"))
    if op.row >= len(rows):
        raise Miss("op.row", f"the table has {len(rows)} rows, numbered 0 to {len(rows) - 1}")
    cells = rows[op.row].findall(qn("a:tc"))
    if op.col >= len(cells):
        raise Miss("op.col", f"row {op.row} has {len(cells)} cells, numbered 0 to {len(cells) - 1}")
    tc = cells[op.col]
    if tc.get("hMerge") in ("1", "true") or tc.get("vMerge") in ("1", "true"):
        raise Miss("op.col", f"row {op.row} col {op.col} is merged into a neighbouring cell; set that cell")
    if any("\n" in text.paragraph_text(p) for p in text.cell_paragraphs(tc)):
        raise Miss("op", f'row {op.row} col {op.col} holds a line feed ("\\n") inside a paragraph, so "\\n" cannot mark where its paragraphs split')
    source = text.cell_text(tc)
    if source != op.old:
        raise Miss("op.old", f"row {op.row} col {op.col} reads {source!r}, not {op.old!r}")
    try:
        text.plan_cell(tc, op.new)
    except text.SpliceError as e:
        raise Miss("op.new", str(e)) from e
    return CellAt(slide, shape, tc=tc, row=op.row, col=op.col, new=op.new)


def _locate_add(deck: slides.Deck, ctx: Context, cid: str, op: AddSlide) -> NewSlide:
    found = [layout for layout in deck.layouts if layout.name == op.layout]
    if not found:
        names = sorted({layout.name for layout in deck.layouts})
        raise Miss("op.layout", f"no layout {op.layout!r} on the masters the deck's slides use; layouts: {_list(map(repr, names))}")
    if len(found) > 1:
        raise Miss("op.layout", f"{len(found)} masters have a layout named {op.layout!r}, so the name does not say which")
    after = _slide(deck, op.after, "op.after") if op.after is not None else None
    ordinal = ctx.ordinals[cid]
    slide_id = max((s.id for s in deck.slides), default=255) + ordinal
    if slide_id > slides.MAX_SLIDE_ID:
        raise Miss("op", "no slide id is free above the deck's highest")
    return NewSlide(found[0], slide_id, f"ppt/slides/slide{deck.last_part + ordinal}.xml", after)


def _locate_fill(pkg: Package, deck: slides.Deck, ctx: Context, op: FillPlaceholder) -> PlaceholderAt:
    add = ctx.ops.get(op.slide)
    if add is None:
        raise Miss("op.slide", f"no change {op.slide!r}; fill_placeholder.slide names an add_slide change")
    if not isinstance(add, AddSlide):
        raise Miss("op.slide", f"{op.slide} is a {add.kind} change, not an add_slide")
    try:
        new = _locate_add(deck, ctx, op.slide, add)
    except Miss as m:
        raise Miss("op.slide", f"{op.slide} cannot add its slide: {m.message}") from m
    ph = next((p for p in new.layout.placeholders if p.id == op.shape), None)
    if ph is None:
        listed = _list(f"{p.id} {p.name!r} ({p.type})" for p in new.layout.placeholders)
        raise Miss("op.shape", f"layout {new.layout.name!r} has no placeholder {op.shape}; its placeholders: {listed}")
    if not ph.holds_text:
        raise Miss("op.shape", f"placeholder {op.shape} {ph.name!r} is a {ph.type} placeholder and holds no text")
    for i, para in enumerate(op.paragraphs):
        _no_newline(para.text, f"op.paragraphs.{i}.text")
    return PlaceholderAt(new, op.slide, ph, op.paragraphs)


def _slide(deck: slides.Deck, slide_id: int, field: str) -> slides.SourceSlide:
    slide = deck.slide(slide_id)
    if slide is None:
        raise Miss(field, _no_slide(deck, slide_id))
    return slide


def _no_slide(deck: slides.Deck, slide_id: int) -> str:
    return f"no slide {slide_id}; slide ids: {_list(f'{s.id} (slide {s.index})' for s in deck.slides)}"


ShapeKind = Literal["text", "table", "chart"]


def _shape(pkg: Package, slide: slides.SourceSlide, shape_id: int, kind: ShapeKind) -> etree._Element:
    found = [s for s in slides.shapes(slides.shape_tree(pkg.xml(slide.part))) if slides.shape_id(s) == shape_id]
    if len(found) == 1:
        return found[0]
    if found:
        raise Miss("op.shape", f"{slide.name} has {len(found)} shapes with id {shape_id}, so the id does not say which")
    listed = [_describe(s, kind) for s in slides.shapes(slides.shape_tree(pkg.xml(slide.part)))]
    label = {"text": "shapes with text", "table": "tables", "chart": "charts"}[kind]
    raise Miss("op.shape", f"{slide.name} has no shape {shape_id}; {label}: {_list(d for d in listed if d) or 'none'}")


def _describe(shape: etree._Element, kind: ShapeKind) -> str | None:
    """How a bad shape id's message lists `shape` as a candidate, or None when an op of `kind` cannot address it."""
    name = f"{slides.shape_id(shape)} {slides.shape_name(shape)!r}"
    match kind:
        case "table":
            return name if shape.find(f".//{qn('a:tbl')}") is not None else None
        case "chart":
            return name if shape.find(f".//{{{chart.C_NS}}}chart") is not None else None
    body = shape.find(qn("p:txBody"))
    words = " ".join(text.paragraph_text(p) for p in body.iterfind(qn("a:p"))).strip() if body is not None else ""
    return f"{name} ({text.clip(words, 30)!r})" if words else None


def _paragraphs(pkg: Package, deck: slides.Deck, slide_id: int, shape_id: int):
    slide = _slide(deck, slide_id, "op.slide")
    shape = _shape(pkg, slide, shape_id, "text")
    body = shape.find(qn("p:txBody"))
    if body is None:
        hint = "; use set_cell" if shape.find(f".//{qn('a:tbl')}") is not None else ""
        hint = hint or ("; use set_chart_value" if shape.find(f".//{{{chart.C_NS}}}chart") is not None else "")
        raise Miss("op.shape", f"shape {shape_id} {slides.shape_name(shape)!r} holds no text{hint}")
    return slide, shape, body.findall(qn("a:p"))


def _no_newline(value: str, field: str) -> None:
    if "\n" in value:
        raise Miss(field, '"\\n" would start a new paragraph; use insert_paragraph, or "\\v" for a line break')


def _paragraph_list(indices: set[int]) -> str:
    return f"{'paragraph' if len(indices) == 1 else 'paragraphs'} {_list(sorted(indices))}"


def _count(n: int) -> str:
    return f"{n} paragraphs, numbered 0 to {n - 1}" if n else "no paragraphs"


def _list(items) -> str:
    return ", ".join(str(i) for i in items)


def _references(cs: ChangeSet, deck: slides.Deck) -> list[Problem]:
    """Ids that repeat, and ids that name nothing: an ask_id, or a flag's or held item's slides."""
    problems: list[Problem] = []
    for kind, items in (("changes", cs.changes), ("asks", cs.asks), ("flags", cs.flags), ("held", cs.held)):
        first: dict[str, int] = {}
        for i, item in enumerate(items):
            if item.id in first:
                problems.append(Problem(f"{item.id} id", f"{kind}[{first[item.id]}] and {kind}[{i}] share this id"))
            first.setdefault(item.id, i)
    for item in (*cs.flags, *cs.held):
        for slide_id in item.slides:
            if deck.slide(slide_id) is None:
                problems.append(Problem(f"{item.id} slides", _no_slide(deck, slide_id)))
    asks = [a.id for a in cs.asks]
    for c in cs.changes:
        if c.ask_id not in asks:
            problems.append(Problem(f"{c.id} ask_id", f"no ask {c.ask_id!r}; asks: {_list(asks) or 'none'}"))
    return problems


def _conflicts(located: Sequence[tuple[Change, Target]]) -> list[Problem]:
    """Judged over every change whatever its decision, so any subset a reviewer keeps is conflict-free."""
    deleted: dict[int, str] = {}
    for c, t in located:
        if isinstance(t, SlideAt):
            deleted.setdefault(t.slide.id, c.id)
    problems: list[Problem] = []
    claims: dict[tuple, str] = {}
    quotes: dict[tuple, list[tuple[int, int, str, str]]] = {}
    for c, t in located:
        claim = t.claim()
        if claim and claim.key in claims:
            problems.append(Problem(f"{c.id} {claim.field}", f"{claim.taken} by {claims[claim.key]}"))
        elif claim:
            claims[claim.key] = c.id
        if t.touches and t.touches.id in deleted:
            problems.append(Problem(f"{c.id} op.slide", f"{t.touches.name} is deleted by {deleted[t.touches.id]}"))
        if t.follows and t.follows.id in deleted:
            problems.append(Problem(f"{c.id} op.after", f"{t.follows.name} is deleted by {deleted[t.follows.id]}; nothing can follow it"))
        if isinstance(t, TextAt):
            spans = quotes.setdefault((t.slide.id, slides.shape_id(t.shape), t.paragraph), [])
            end = t.at + len(t.old)
            other = next((s for s in spans if t.at < s[1] and s[0] < end), None)
            if other:
                problems.append(Problem(f"{c.id} op.old", f"{t.old!r} overlaps {other[2]}'s quote {other[3]!r} in paragraph {t.paragraph}"))
            spans.append((t.at, end, c.id, t.old))
    return problems + _cycles(located)


def _cycles(located: Sequence[tuple[Change, Target]]) -> list[Problem]:
    after = {t.slide.id: (t.after.id if t.after else None, c.id, t.slide) for c, t in located if isinstance(t, MoveTo)}
    problems = []
    for start, (first, cid, slide) in after.items():
        seen, at = {start}, first
        while at in after and at not in seen:
            seen.add(at)
            at = after[at][0]
        if at == start:
            problems.append(Problem(f"{cid} op.after", f"{slide.name} is placed after a slide that is placed after it"))
    return problems


def execute(checked: Checked) -> bytes:
    """The maker-output deck: every change as written. Equal to apply with every decision keep_new."""
    return _build(checked, [(c.id, c.op) for c in checked.changeset.changes])


def apply(checked: Checked) -> Applied:
    """The final deck, from a fresh copy of the source: keep_new and edited changes replayed, keep_old ones not."""
    changes = checked.changeset.changes
    pending = [c.id for c in changes if c.decision == "pending"]
    if pending:
        raise Undecided(pending)
    kept_adds = {c.id for c in changes if isinstance(c.op, AddSlide) and c.decision != "keep_old"}
    ops: list[tuple[str, Op]] = []
    kept_new, edited, kept_old, dropped = [], [], [], []
    for c in changes:
        if c.decision == "keep_old":
            kept_old.append(c.id)
            continue
        if isinstance(c.op, FillPlaceholder) and c.op.slide not in kept_adds:
            dropped.append((c.id, c.op.slide))
            continue
        if isinstance(c.decision, Edited):
            ops.append((c.id, c.op.edit(c.decision.edited)))
            edited.append(c.id)
        else:
            ops.append((c.id, c.op))
            kept_new.append(c.id)
    return Applied(_build(checked, ops), tuple(kept_new), tuple(edited), tuple(kept_old), tuple(dropped))


def _build(checked: Checked, ops: Sequence[tuple[str, Op]]) -> bytes:
    """The one write path: a pure function of the source bytes and the ops. Every target is located on the
    pristine trees before anything is written, so no write shifts another change's address."""
    pkg = Package(checked.source)
    deck = slides.read_deck(pkg)
    ctx = Context.of(checked.changeset)
    targets = [locate(pkg, deck, ctx, cid, op) for cid, op in ops]
    build = Build(pkg, deck, {s.id: s.rid for s in deck.slides}, [], set(), {})
    for t in sorted(targets, key=lambda t: t.rank()):
        t.write(build)
    order = slides.final_order([s.id for s in deck.slides], build.deleted, build.placements)
    slides.write_order(pkg, deck, [(sid, build.rids[sid]) for sid in order], dict(build.placements))
    return pkg.to_bytes()


def review(checked: Checked, executed_path: Path, executed: bytes) -> Review:
    """The review view of `executed`, the deck execute wrote: positions and added slides' titles are read from
    it, everything else from what check resolved on the source."""
    cs = checked.changeset
    source = Package(checked.source)
    built = Package(executed)
    placed = {s.id: s for s in slides.read_deck(built).slides}
    adds = {i.change.id: i.target.id for i in checked.items if isinstance(i.target, NewSlide)}
    executed_index = {sid: s.index for sid, s in placed.items()} | {cid: placed[sid].index for cid, sid in adds.items()}
    boxes = _boxes(checked.source, cs.source.path)
    slide_rows = [
        ReviewSlide(key=s.id, source_index=s.index, executed_index=executed_index.get(s.id), title=slides.title(source.xml(s.part)))
        for s in slides.read_deck(source).slides
    ] + [
        ReviewSlide(key=cid, source_index=None, executed_index=placed[sid].index, title=slides.title(built.xml(placed[sid].part)))
        for cid, sid in adds.items()
    ]
    changes = [
        ReviewChange(
            id=i.change.id,
            ask_id=i.change.ask_id,
            kind=i.change.op.kind,
            structural=i.change.op.structural,
            slide=i.slide,
            source_index=i.source_index,
            executed_index=executed_index.get(i.slide),
            shape=_review_shape(i, boxes),
            before=i.before,
            after=i.after,
            span=i.span,
            depends_on=i.depends_on,
            admits=i.change.op.admits,
            decision=i.change.decision,
            rationale=i.change.rationale,
            refs=i.change.refs,
            notes=i.notes,
        )
        for i in checked.items
    ]
    return Review(
        source=cs.source,
        executed=Executed(path=str(executed_path), sha256=hashlib.sha256(executed).hexdigest()),
        meeting=cs.meeting,
        asks=cs.asks,
        flags=cs.flags,
        held=cs.held,
        slides=tuple(slide_rows),
        changes=tuple(changes),
    )


def _boxes(data: bytes, path: str) -> dict[tuple[int, int], Box]:
    deck = read_deck(open_presentation(data, path), path, "")
    return {
        (slide.index, shape.xml.shape_id): Box(x=shape.left, y=shape.top, w=shape.width, h=shape.height)
        for slide in deck.slides
        for shape in slide.shapes
    }


def _review_shape(item: Item, boxes: Mapping[tuple[int, int], Box]) -> ReviewShape | None:
    if item.shape is None:
        return None
    shape_id, name = item.shape
    box = boxes.get((item.source_index, shape_id)) if item.source_index is not None else None
    return ReviewShape(id=shape_id, name=name, box=box)

from __future__ import annotations

import hashlib
import io
import math
import re
import zipfile
from collections import Counter
from collections.abc import Callable
from dataclasses import dataclass, replace
from typing import Any, Literal

from lxml import etree
from pptx.oxml.text import CT_RegularTextRun, CT_TextField, CT_TextLineBreak

from deckcheck.model import Deck, DeckError, Paragraph, Shape, Slide, Violation, open_presentation, read_deck
from deckcheck.rules import RuleSet, run_rules, slide_fonts, small_runs, visual_box

MAX_PASSES = 3

ABBREVIATIONS = frozenset(
    {"etc", "mgmt", "asst", "reps", "approx", "incl", "excl", "esp", "dept", "govt", "vs", "cf", "inc", "ltd", "corp"}
)
# Their glyphs sit on letter codes, so swapping the typeface turns symbols into letters.
SYMBOL_FONTS = frozenset({"Wingdings", "Wingdings 2", "Wingdings 3", "Webdings", "Symbol", "Marlett"})

Why = Literal["report-only", "declined", "did-not-stick", "pass-limit"]


@dataclass(frozen=True)
class Change:
    what: Literal["text", "size", "font", "box"]
    before: str
    after: str


@dataclass(frozen=True)
class Fixed:
    violation: Violation
    changes: tuple[Change, ...]
    pass_no: int


@dataclass(frozen=True)
class Reported:
    violation: Violation
    why: Why
    detail: str


Outcome = Fixed | Reported


@dataclass(frozen=True)
class FixResult:
    data: bytes
    input_sha256: str
    output_sha256: str
    passes: int
    outcomes: tuple[Outcome, ...]

    @property
    def fixed(self) -> tuple[Fixed, ...]:
        return tuple(o for o in self.outcomes if isinstance(o, Fixed))

    @property
    def remaining(self) -> tuple[Reported, ...]:
        return tuple(o for o in self.outcomes if isinstance(o, Reported))


@dataclass(frozen=True)
class Declined:
    reason: str


FixerFn = Callable[[Any, Deck, dict[str, Any]], tuple[Change, ...] | Declined]


# lxml hands back the same element object while a reference to it lives, and every key holds one, so a
# Slide, Shape, or Paragraph keeps its key across passes although each pass builds new model objects.
def _key(v: Violation) -> tuple[str, object]:
    return (v.rule, v.target.xml)


def fix_deck(data: bytes, rules: RuleSet, path: str) -> FixResult:
    sha = hashlib.sha256(data).hexdigest()
    prs = open_presentation(data, path)
    at_open = [_canonical(s) for s in prs.slides]
    fixed: dict[tuple, Fixed] = {}
    declined: dict[tuple, str] = {}
    for passes in range(1, MAX_PASSES + 1):
        deck = read_deck(prs, path, sha)
        final = run_rules(deck, rules)
        if passes == 1:
            _check_read_is_pure(prs, at_open, path)
        applied = 0
        for v in final:
            if v.rule not in FIXERS or _key(v) in fixed:
                continue
            result = FIXERS[v.rule](v.target, deck, rules.params[v.rule])
            if isinstance(result, Declined):
                declined[_key(v)] = result.reason
            else:
                fixed[_key(v)] = Fixed(v, result, passes)
                applied += 1
        if not applied:
            break
    else:
        final = run_rules(read_deck(prs, path, sha), rules)

    still = {_key(v) for v in final}
    outcomes: list[Outcome] = [f for k, f in fixed.items() if k not in still]
    for v in final:
        k = _key(v)
        if v.rule not in FIXERS:
            outcomes.append(Reported(v, "report-only", "fix by hand"))
        elif k in fixed:
            outcomes.append(Reported(v, "did-not-stick", f"fixed in pass {fixed[k].pass_no} and still fires"))
        elif k in declined:
            outcomes.append(Reported(v, "declined", declined[k]))
        else:
            outcomes.append(Reported(v, "pass-limit", f"first fired after pass {MAX_PASSES}"))
    outcomes.sort(key=lambda o: (o.violation.slide, o.violation.rule, o.violation.shape or ""))

    written = {
        s.part.partname.membername: s.part.blob
        for s, before in zip(prs.slides, at_open, strict=True)
        if _canonical(s) != before
    }
    out = _repack(data, written) if written else data
    return FixResult(out, sha, hashlib.sha256(out).hexdigest(), passes, tuple(outcomes))


def _canonical(slide) -> bytes:
    return etree.tostring(slide._element, method="c14n")


def _check_read_is_pure(prs, at_open: list[bytes], path: str) -> None:
    for index, (s, before) in enumerate(zip(prs.slides, at_open, strict=True), start=1):
        if _canonical(s) != before:
            raise DeckError(
                f"reading {path} changed slide {index}, so fix cannot tell its own edits apart; nothing written"
            )


# python-pptx's save re-serialises every part and drops what it does not model, such as a relationship
# whose target is missing, so the output is the input's zip with only the written slide parts replaced.
def _repack(data: bytes, parts: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(buf, "w") as dst:
        for info in src.infolist():
            entry = zipfile.ZipInfo(info.filename, info.date_time)
            entry.compress_type, entry.external_attr = info.compress_type, info.external_attr
            dst.writestr(entry, parts[info.filename] if info.filename in parts else src.read(info))
    return buf.getvalue()


def plural(n: int, one: str, many: str) -> str:
    return f"{n} {one if n == 1 else many}"


def fonts_to_replace(slide: Slide, max_fonts: int) -> list[str] | Declined:
    present = slide_fonts(slide)
    runs = [(s.kind == "title", r) for s in slide.shapes for p in s.paragraphs for r in p.runs]
    usage = Counter[str]()
    for _, r in runs:
        usage[r.font] += len(r.text)
    kept = {slide.theme.major, slide.theme.minor} | SYMBOL_FONTS
    candidates = sorted(present - kept, key=lambda f: (-usage[f], f))
    for keep in range(len(candidates) - 1, -1, -1):
        drop = candidates[keep:]
        after = (present - set(drop)) | {
            slide.theme.major if is_title else slide.theme.minor for is_title, r in runs if r.font in drop
        }
        if len(after) <= max_fonts:
            return sorted(drop)
    return Declined("theme and symbol fonts alone exceed the maximum")


def fix_fonts(slide: Slide, deck: Deck, params: dict[str, Any]) -> tuple[Change, ...] | Declined:
    drop = fonts_to_replace(slide, params["max"])
    if isinstance(drop, Declined):
        return drop
    refs = {True: ("+mj-lt", slide.theme.major), False: ("+mn-lt", slide.theme.minor)}
    hits = Counter[tuple[str, bool]]()
    for s in slide.shapes:
        for p in s.paragraphs:
            for r in p.runs:
                if r.font in drop:
                    rPr = r.xml.get_or_add_rPr()
                    rPr._remove_latin()
                    rPr.get_or_add_latin().typeface = refs[s.kind == "title"][0]
                    hits[(r.font, s.kind == "title")] += 1
    return tuple(
        Change("font", font, f"{refs[is_title][0]} ({refs[is_title][1]}), {plural(n, 'run', 'runs')}")
        for (font, is_title), n in sorted(hits.items())
    )


DOTTED_WORD = re.compile(r"[^\W\d_]+(?:\.[^\W\d_]+)*$")


def end_punctuation(text: str, chars: str) -> int | Declined:
    text = text.rstrip()
    stem = text.rstrip(chars)
    tail = text[len(stem) :]
    if ".." in tail or stem.endswith("…"):
        return Declined("ellipsis")
    if not stem.strip():
        return Declined("the bullet is only punctuation")
    if stem.rstrip() != stem and stem.rstrip()[-1] in chars:
        return Declined("punctuation separated by spaces")
    word = DOTTED_WORD.search(stem)
    if (
        tail.startswith(".")
        and word
        and (word.group(0).lower() in ABBREVIATIONS or "." in word.group(0) or len(word.group(0)) == 1)
    ):
        return Declined(f"abbreviation {word.group(0)}.")
    return len(tail)


def fix_bullet_end(p: Paragraph, deck: Deck, params: dict[str, Any]) -> tuple[Change, ...] | Declined:
    n = end_punctuation(p.text, params["chars"])
    if isinstance(n, Declined):
        return n
    edits: list[tuple[CT_RegularTextRun, str]] = []
    left = n
    for child in reversed(p.xml.content_children):
        if left == 0:
            break
        body = child.text.rstrip()
        if isinstance(child, CT_TextLineBreak) or not body:
            continue
        if isinstance(child, CT_TextField):
            return Declined("ends in a field")
        k = min(left, len(body))
        edits.append((child, body[: len(body) - k] + child.text[len(body) :]))
        left -= k
    for run, text in edits:
        run.text = text
    before = p.text.strip()
    return (Change("text", before, before[: len(before) - n]),)


def centipoints_at_least(pt: float) -> int:
    return math.ceil(round(pt * 100, 6))


def fix_min_size(s: Shape, deck: Deck, params: dict[str, Any]) -> tuple[Change, ...] | Declined:
    sz = centipoints_at_least(params["min_pt"])
    sizes = Counter[float]()
    for r in small_runs(s, params["min_pt"]):
        r.xml.get_or_add_rPr().sz = sz
        sizes[r.size_pt] += 1
    return tuple(
        Change("size", f"{pt:g}pt", f"{sz / 100:g}pt, {plural(n, 'run', 'runs')}") for pt, n in sorted(sizes.items())
    )


def fitted_box(s: Shape, slide_w: int, slide_h: int) -> tuple[int, int, int, int] | Declined:
    if s.to_slide.rotated:
        return Declined("inside a rotated or flipped group")
    left, top, right, bottom = visual_box(s)
    width, height = s.width, s.height
    if right - left > slide_w or bottom - top > slide_h:
        if s.kind == "table":
            return Declined("table is larger than the slide")
        if s.rotation % 90 == 0:
            odd = round(s.rotation / 90) % 2 == 1
            width, height = min(width, slide_h if odd else slide_w), min(height, slide_w if odd else slide_h)
        else:
            k = min(slide_w / (right - left), slide_h / (bottom - top))
            width, height = round(width * k), round(height * k)
        s = replace(
            s,
            left=round(s.left + (s.width - width) / 2),
            top=round(s.top + (s.height - height) / 2),
            width=width,
            height=height,
        )
        left, top, right, bottom = visual_box(s)
    dx = -left if left < 0 else min(0, slide_w - right)
    dy = -top if top < 0 else min(0, slide_h - bottom)
    return s.left + dx, s.top + dy, width, height


def _box(left: int, top: int, width: int, height: int) -> str:
    return f"({left}, {top}) {width}x{height} EMU"


def fix_bounds(s: Shape, deck: Deck, params: dict[str, Any]) -> tuple[Change, ...] | Declined:
    box = fitted_box(s, deck.slide_width, deck.slide_height)
    if isinstance(box, Declined):
        return box
    left, top, width, height = box
    t = s.to_slide
    s.xml.x = round((left - t.dx) / t.sx)
    s.xml.y = round((top - t.dy) / t.sy)
    s.xml.cx = round(width / t.sx)
    s.xml.cy = round(height / t.sy)
    return (Change("box", _box(s.left, s.top, s.width, s.height), _box(*box)),)


FIXERS: dict[str, FixerFn] = {
    "max-fonts-per-slide": fix_fonts,
    "no-bullet-end-punctuation": fix_bullet_end,
    "min-font-size": fix_min_size,
    "within-slide-bounds": fix_bounds,
}

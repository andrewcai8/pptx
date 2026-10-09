"""Measure the XML python-pptx adds when code only reads a deck, to keep scenario.READ_ARTIFACTS honest.

Usage, from the repo root:
    uv run --project deckcheck python evals/read_artifacts.py

Reads every corpus deck, plus a deck python-pptx builds itself, through each getter below and prints every element the
read added, with whether the scorer ignores it. Exit 1 when a READ_ARTIFACTS entry is never measured.
"""

from __future__ import annotations

import copy
import io
import sys
from collections import Counter
from collections.abc import Callable, Iterator

import yaml
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from deckcheck.model import read_deck
from scenario import READ_ARTIFACTS, corpus, strip_read_artifacts

MARK = "{urn:read-artifacts}seen"


def _shapes(prs) -> Iterator:
    def walk(shapes):
        for s in shapes:
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(s.shapes)
            else:
                yield s

    for slide in prs.slides:
        yield from walk(slide.shapes)


def _paragraphs(prs) -> Iterator:
    for s in _shapes(prs):
        if s.has_text_frame:
            yield from s.text_frame.paragraphs
        if s.has_table:
            for cell in s.table.iter_cells():
                yield from cell.text_frame.paragraphs


def _runs(prs) -> Iterator:
    for p in _paragraphs(prs):
        yield from p.runs


GETTERS: dict[str, Callable] = {
    "deckcheck read_deck": lambda prs: read_deck(prs, "", ""),
    "shape.text_frame.text": lambda prs: [s.text_frame.text for s in _shapes(prs) if s.has_text_frame],
    "cell.text_frame.text": lambda prs: [c.text_frame.text for s in _shapes(prs) if s.has_table for c in s.table.iter_cells()],
    "paragraph.level": lambda prs: [p.level for p in _paragraphs(prs)],
    "paragraph.alignment": lambda prs: [p.alignment for p in _paragraphs(prs)],
    "paragraph.font": lambda prs: [p.font.size for p in _paragraphs(prs)],
    "run.font": lambda prs: [(r.font.size, r.font.name, r.font.bold) for r in _runs(prs)],
    "run.font.color": lambda prs: [r.font.color for r in _runs(prs)],
    "run.hyperlink": lambda prs: [r.hyperlink.address for r in _runs(prs)],
    "shape.fill": lambda prs: [s.fill.type for s in _shapes(prs) if hasattr(s, "fill")],
    "shape.line.fill": lambda prs: [s.line.fill.type for s in _shapes(prs) if hasattr(s, "line")],
    "shape.line.color": lambda prs: [s.line.color for s in _shapes(prs) if hasattr(s, "line")],
    "chart series values": lambda prs: [series.values for s in _shapes(prs) if s.has_chart for plot in s.chart.plots for series in plot.series],
    "chart.chart_title": lambda prs: [s.chart.chart_title for s in _shapes(prs) if s.has_chart],
}


def _python_pptx_deck() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "A title python-pptx wrote"
    slide.placeholders[1].text_frame.text = "A bullet python-pptx wrote"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def added(data: bytes, getter: Callable) -> tuple[Counter, set[tuple[str, str]]]:
    prs = Presentation(io.BytesIO(data))
    parts = [p for p in prs.part.package.iter_parts() if hasattr(p, "_element")]
    for p in parts:
        for el in p._element.iter(etree.Element):
            el.set(MARK, "")
    getter(prs)
    found: Counter = Counter()
    pairs: set[tuple[str, str]] = set()
    for p in parts:
        for el in p._element.iter(etree.Element):
            parent = el.getparent()
            if MARK not in el.attrib and parent is not None and MARK in parent.attrib:
                holder = etree.Element(parent.tag)
                holder.append(copy.deepcopy(el))
                pairs |= {(e.getparent().tag, e.tag) for e in holder.iter(etree.Element) if e is not holder}
                strip_read_artifacts(holder)
                found[(parent.tag, el.tag, not len(holder))] += 1
    return found, pairs


def main() -> int:
    decks = yaml.safe_load(corpus.MANIFEST.read_text())["decks"]
    sources = [corpus.fetch(d).read_bytes() for d in decks] + [_python_pptx_deck()]
    measured: set[tuple[str, str]] = set()
    for name, getter in GETTERS.items():
        total: Counter = Counter()
        for data in sources:
            found, pairs = added(data, getter)
            total += found
            measured |= pairs
        for (parent, tag, ignored), n in sorted(total.items()):
            what = "ignored" if ignored else "counts as a change"
            print(f"{name:22} adds {n:6} {etree.QName(tag).localname} in {etree.QName(parent).localname}, {what}")
        if not total:
            print(f"{name:22} adds nothing")
    if stale := sorted(set(READ_ARTIFACTS) - measured):
        print("READ_ARTIFACTS lists elements no getter adds: " + ", ".join(f"{etree.QName(c).localname} in {etree.QName(p).localname}" for p, c in stale))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())

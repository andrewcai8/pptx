import hashlib
import io
import json
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

import score
from deckcheck.changeset.cli import main as changeset
from scenario import ROOT

PASSES = {
    "fmcg-diagnostic-timeline": "SCENARIO PASS fmcg-diagnostic-timeline (26 checks, 7 intent checks deferred)",
    "insurance-workshop-prep": "SCENARIO PASS insurance-workshop-prep (26 checks, 4 intent checks deferred)",
    "rcc-flexibility-wording": "SCENARIO PASS rcc-flexibility-wording (28 checks, 3 intent checks deferred)",
    "retail-impact-title": "SCENARIO PASS retail-impact-title (36 checks, 5 intent checks deferred)",
    "solar-market-refresh": "SCENARIO PASS solar-market-refresh (41 checks, 6 intent checks deferred)",
}
TURN = re.compile(r"^\[(\d\d:\d\d:\d\d)\] ([^(]+?) \([^)]*\): (.*)$")
SOLAR_MARKET_SLIDE = 2147478638


@pytest.fixture(autouse=True)
def at_repo_root(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(ROOT)


def fixture(name: str) -> Path:
    return ROOT / "evals" / name / "changeset.json"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run(argv: list, capsys: pytest.CaptureFixture[str]) -> tuple[int, str]:
    code = changeset([str(a) for a in argv])
    return code, capsys.readouterr().out


def decided(name: str, decision: str, path: Path) -> Path:
    doc = json.loads(fixture(name).read_text())
    for c in doc["changes"]:
        c["decision"] = decision
    path.write_text(json.dumps(doc))
    return path


@pytest.mark.parametrize("name", sorted(PASSES))
def test_each_fixture_executed_passes_its_scenario(name: str, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = tmp_path / "executed.pptx"
    assert run(["execute", fixture(name), "--out", out], capsys)[0] == 0

    code = score.main([name, str(out), "--out", str(tmp_path)])

    assert (code, capsys.readouterr().out.splitlines()[0]) == (0, PASSES[name])


@pytest.mark.parametrize("name", sorted(PASSES))
def test_decisions_replay_each_fixture_onto_its_source_byte_for_byte(name: str, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    source = ROOT / json.loads(fixture(name).read_text())["source"]["path"]
    source_sha = sha(source)
    executed = tmp_path / "executed.pptx"
    run(["execute", fixture(name), "--out", executed], capsys)
    executed_sha = sha(executed)
    old, new, again = tmp_path / "old.pptx", tmp_path / "new.pptx", tmp_path / "again.pptx"

    run(["apply", decided(name, "keep_old", tmp_path / "old.json"), "--out", old], capsys)
    run(["apply", decided(name, "keep_new", tmp_path / "new.json"), "--out", new], capsys)
    run(["apply", tmp_path / "new.json", "--out", again], capsys)

    # Equal bytes for an edited deck hold for one zlib build: untouched members are recompressed.
    assert old.read_bytes() == source.read_bytes()
    assert new.read_bytes() == executed.read_bytes() == again.read_bytes()
    assert (sha(source), sha(executed)) == (source_sha, executed_sha)


@pytest.mark.parametrize("name", sorted(PASSES))
def test_every_ref_quotes_its_transcript_turn_verbatim(name: str) -> None:
    turns = {}
    for line in (ROOT / "evals" / name / "transcript.md").read_text().splitlines():
        m = TURN.match(line)
        if m:
            turns[m.group(1)] = (m.group(2), m.group(3))
    doc = json.loads(fixture(name).read_text())
    refs = [r for key in ("asks", "changes", "flags", "held") for item in doc[key] for r in item["refs"]]

    assert [r for r in refs if r["t"] not in turns or turns[r["t"]][0] != r["speaker"] or r["quote"] not in turns[r["t"]][1]] == []


def market_slide(path: Path):
    return next(s for s in Presentation(str(path)).slides if s.slide_id == SOLAR_MARKET_SLIDE)


def test_the_solar_chart_edit_moves_the_cache_and_the_workbook_to_410(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = tmp_path / "executed.pptx"
    run(["execute", fixture("solar-market-refresh"), "--out", out], capsys)

    chart = next(s for s in market_slide(out).shapes if s.shape_id == 8).chart
    book = chart.part.chart_workbook.xlsx_part
    sheet = openpyxl.load_workbook(io.BytesIO(book.blob))["Sheet1"]
    assert (
        chart.plots[0].series[0].values,
        [c.value for c in sheet[1]],
        str(book.partname),
        book.content_type,
    ) == (
        (410.0, 2000.0, 11000.0),
        [410, 2000, 11000],
        "/ppt/embeddings/Microsoft_Excel_Binary_Worksheet.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def test_the_solar_think_cell_label_reads_37_in_its_text_and_its_field_format(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = tmp_path / "executed.pptx"
    run(["execute", fixture("solar-market-refresh"), "--out", out], capsys)

    label = next(s for s in market_slide(out).shapes if s.shape_id == 69)
    field = label.text_frame.paragraphs[0]._p.find(qn("a:fld"))
    assert (field.findtext(qn("a:t")), field.get("type")[8:].replace("'", "")) == ("+37%", "+37%")


def test_a_solar_chart_point_whose_xlsb_cell_is_blank_is_refused(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    doc = json.loads(fixture("solar-market-refresh").read_text())
    deck = tmp_path / "blank-cell.pptx"
    with zipfile.ZipFile(ROOT / doc["source"]["path"]) as src, zipfile.ZipFile(deck, "w") as dst:
        for info in src.infolist():
            body = src.read(info)
            if info.filename == "ppt/charts/chart1.xml":
                body = body.replace(b"Sheet1!$A$1:$C$1", b"Sheet1!$D$1:$F$1")
            dst.writestr(info, body)
    doc["source"] = {"path": str(deck), "sha256": sha(deck)}
    doc["changes"] = [c for c in doc["changes"] if c["id"] == "chart-2022-bar"]
    cs = tmp_path / "changeset.json"
    cs.write_text(json.dumps(doc))

    assert run(["validate", cs], capsys) == (
        1,
        f"INVALID {cs}: 1 problem\n  chart-2022-bar op: shape 8 'Chart 7': workbook cell Sheet1!D1 is blank, not a number\n",
    )

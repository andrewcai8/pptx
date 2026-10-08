from pptx.oxml.xmlchemy import OxmlElement

from deckedit import DeckEdit, variant

OLD = "Flexibility to teams"
NEW = "Teams choose their own office days within the site cap"


@variant()
def good(d: DeckEdit) -> None:
    """Slide 4 flexibility bullet reworded to the wording Daniel settled on."""
    d.replace(4, OLD, NEW)


@variant(fails={("scope", 6), ("missing", 4)})
def wrong_slide(d: DeckEdit) -> None:
    """The new wording overwrites the flexible working model bullet on slide 6 and slide 4 keeps the old bullet."""
    d.replace(
        6,
        "Working model will also need to be flexible during re-opening, with ability to moderate up/down, based on local risk factors (e.g., virus resurgence)",
        NEW,
    )


@variant(base=good, fails={("style", 4)})
def trailing_period(d: DeckEdit) -> None:
    """The reworded bullet ends with a full stop, a new no-bullet-end-punctuation breach."""
    d.replace(4, NEW, NEW + ".")


@variant(base=good, fails={("non-change", 7)})
def edited_timing_slide(d: DeckEdit) -> None:
    """Slide 7 gets a provincial-announcement caveat after the team agreed to leave it."""
    d.replace(
        7,
        "Government regulations and company guiding principles should determine timing for re-opening offices",
        "Government regulations and company guiding principles should determine timing for re-opening offices (pending provincial announcement)",
    )


@variant(fails={("missing", 4), ("forbidden", 4)})
def kept_first_answer(d: DeckEdit) -> None:
    """Slide 4 uses Daniel's first wording, which the team talked him out of."""
    d.replace(4, OLD, "Let teams pick their own schedule")


@variant(base=good, fails={("non-change", 12)})
def called_out_calgary(d: DeckEdit) -> None:
    """The site page names Calgary's later opening although Daniel parked it until next week."""
    d.replace(12, "Determine office site and maintenance modifications required for safe return", "Determine office site and maintenance modifications required for safe return, with Calgary to follow")


@variant(base=good, fails={("non-change", 4)})
def cap_number_on_slide(d: DeckEdit) -> None:
    """The reworded bullet quotes Mississauga's 380 a day although Daniel said to keep the number off the slide."""
    d.replace(4, NEW, NEW + " (e.g., Mississauga: 380 a day)")


@variant(base=good, fails={("scope", 4)})
def dropped_profitability_bullet(d: DeckEdit) -> None:
    """The flexibility bullet is reworded, and the business profitability bullet next to it is emptied."""
    d.replace(4, "Business profitability, market share, etc.", "")


@variant()
def reworded_daily_cap(d: DeckEdit) -> None:
    """Slide 4 says teams choose their office days within each site's daily cap, a faithful rewording."""
    d.replace(4, OLD, "Teams choose their own office days within each site's daily cap")


@variant()
def reworded_facilities_cap(d: DeckEdit) -> None:
    """Slide 4 says teams pick their office days within the cap Facilities sets for each site, a faithful rewording."""
    d.replace(4, OLD, "Teams pick their office days, within the cap Facilities sets for each site")


@variant()
def reworded_daily_limit(d: DeckEdit) -> None:
    """Slide 4 says teams choose their days in the office within the site's daily limit, a faithful rewording."""
    d.replace(4, OLD, "Teams choose their days in the office, within the site's daily limit")


@variant(base=good, fails={("scope", 6)})
def bullets_off_working_model(d: DeckEdit) -> None:
    """An empty buNone on slide 6's eight factor paragraphs hides their bullets on a slide nobody asked about."""
    factors = next(s for s in d.slide(6).shapes if s.name == "ee4pContent1")
    for p in factors.text_frame.paragraphs:
        p._p.get_or_add_pPr().insert_element_before(OxmlElement("a:buNone"), "a:buAutoNum", "a:buChar", "a:buBlip", "a:tabLst", "a:defRPr", "a:extLst")


def read_every_notes_page(d: DeckEdit) -> None:
    for s in d.slides():
        s.notes_slide.notes_text_frame.text


@variant(fails={("missing", 4)})
def read_notes_only(d: DeckEdit) -> None:
    """Every notes page is read and nothing is edited, so it scores exactly like an untouched copy of the source."""
    read_every_notes_page(d)


@variant(base=good)
def read_notes_after_edit(d: DeckEdit) -> None:
    """The reworded deck with every notes page read, which adds empty notes pages a reader never sees."""
    read_every_notes_page(d)


@variant(base=good, fails={("non-change", 7)})
def reminder_note_on_timing_slide(d: DeckEdit) -> None:
    """Slide 7 gets a speaker note to revisit it after the provincial guidance, though the team agreed to leave slide 7."""
    d.slide(7).notes_slide.notes_text_frame.text = "Revisit once the provincial workplace guidance is out"

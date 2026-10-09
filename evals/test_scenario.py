import hashlib
import json
import re
import shutil
import uuid
import zipfile
from decimal import Decimal
from pathlib import Path

import pytest
import yaml
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

import score
from deckedit import DeckEdit, load_variants
from scenario import ROOT, BadScenario, Charts, open_scenario, snapshot


def deck(path: Path, change=None) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Market outlook"
    slide.placeholders[1].text_frame.text = "Demand grows 4% a year"
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(5), Inches(3), Inches(1)).text_frame.text = "Callout"
    if change:
        change(slide)
    prs.save(path)
    return path


def seen(path: Path) -> tuple:
    snap = snapshot(path)
    return snap.looks, snap.shared


def paragraphs(slide):
    return [p for s in slide.shapes for p in s.text_frame.paragraphs]


def bullets_off(slide):
    for p in slide.placeholders[1].text_frame.paragraphs:
        p._p.get_or_add_pPr().insert_element_before(OxmlElement("a:buNone"), "a:tabLst", "a:defRPr", "a:extLst")


def box(slide):
    return slide.shapes[2]


READS = {
    "paragraph.level": lambda s: [p.level for p in paragraphs(s)],
    "paragraph.font": lambda s: [p.font.size for p in paragraphs(s)],
    "run.font": lambda s: [r.font.size for p in paragraphs(s) for r in p.runs],
    "shape.line.fill": lambda s: box(s).line.fill.type,
    "slide.notes_slide": lambda s: s.notes_slide.notes_text_frame.text,
}

EDITS = {
    "empty buNone": bullets_off,
    "empty noFill on a shape filled by its style": lambda s: box(s).fill.background(),
    "empty effectLst": lambda s: box(s)._element.spPr.append(OxmlElement("a:effectLst")),
    "normAutofit": lambda s: setattr(s.shapes.title.text_frame, "auto_size", MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE),
    "run.font.color": lambda s: [r.font.color for p in paragraphs(s) for r in p.runs],
    "shape.line.color": lambda s: box(s).line.color,
    "a notes page with text": lambda s: setattr(s.notes_slide.notes_text_frame, "text", "Say the number slowly"),
}


@pytest.mark.parametrize("read", READS.values(), ids=READS.keys())
def test_a_python_pptx_read_leaves_the_slide_unchanged(tmp_path, read):
    source = seen(deck(tmp_path / "source.pptx"))
    assert seen(deck(tmp_path / "read.pptx", read)) == source
    assert seen(deck(tmp_path / "edited.pptx", bullets_off)) != source


@pytest.mark.parametrize("edit", EDITS.values(), ids=EDITS.keys())
def test_an_edit_outside_the_read_artifacts_changes_the_slide(tmp_path, edit):
    assert seen(deck(tmp_path / "edited.pptx", edit)) != seen(deck(tmp_path / "source.pptx"))


def scenario_at(d: Path, changes: list | None = None, make=deck, non_changes: list | None = None) -> Path:
    d.mkdir(parents=True)
    source = make(d / "input.pptx")
    deck_ref = {"file": "input.pptx", "sha256": hashlib.sha256(source.read_bytes()).hexdigest()}
    (d / "expected.yaml").write_text(yaml.safe_dump({"deck": deck_ref, "changes": changes or [], "non_changes": non_changes or []}))
    (d / "transcript.md").write_text("[00:00:05] Ana Ruiz (Principal, Kestrel Advisory): Nothing to change today.\n")
    return d


@pytest.fixture
def private_dir():
    top = ROOT / "private" / f"pytest-{uuid.uuid4().hex[:8]}"
    yield top
    shutil.rmtree(top, ignore_errors=True)


@pytest.mark.parametrize(("folder", "opens"), [("private", True), ("artifacts", False), ("evals", False)])
def test_a_private_deck_scenario_opens_only_under_the_repos_private_folder(folder, opens):
    top = ROOT / folder / f"pytest-{uuid.uuid4().hex[:8]}"
    try:
        d = scenario_at(top / "client-scenario")
        if opens:
            assert open_scenario(d).ref == "file:input.pptx"
        else:
            with pytest.raises(BadScenario, match="never enters git"):
                open_scenario(d)
    finally:
        shutil.rmtree(top, ignore_errors=True)


EDIT_SLIDE_1 = [{"id": "c1", "kind": "edit-text", "intent": "Reword the callout.", "said": ["00:00:05"], "slides": {1: {}}}]


def chart_deck(path: Path, plot: str = "barChart") -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    data = CategoryChartData()
    data.categories = ["2022", "2027"]
    data.add_series("Market", (410, 2000))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), data).chart
    chart._chartSpace.plotArea.find(qn("c:barChart")).tag = qn(f"c:{plot}")
    prs.save(path)
    return path


@pytest.mark.parametrize("plot", ["bar3DChart", "line3DChart", "pie3DChart", "stockChart", "surfaceChart", "surface3DChart", "ofPieChart"])
def test_a_chart_python_pptx_cannot_read_is_reported_by_type(tmp_path, plot):
    assert snapshot(chart_deck(tmp_path / "unreadable.pptx", plot)).charts == (Charts((), (plot,)),)
    assert snapshot(chart_deck(tmp_path / "bar.pptx")).charts == (Charts((Decimal("410.0"), Decimal("2000.0")), ()),)


def built(tmp_path: Path, scenario: str, variant: str) -> Path:
    sc = open_scenario(scenario)
    d = DeckEdit(sc)
    next(v for v in load_variants(sc.dir / "build.py") if v.name == variant).apply(d)
    d.save(tmp_path / "output.pptx")
    return tmp_path / "output.pptx"


@pytest.mark.parametrize(
    ("variant", "failures"),
    [
        ("chart_drawn_in_3d", [("scope", 10, "c1 does not ask to change the type of slide 10's chart, but barChart became bar3DChart")]),
        ("untouched_chart_drawn_in_3d", [("scope", 12, "slide 12 changed a chart, image, media, or notes part, but no change asks for it")]),
    ],
)
def test_a_3d_chart_on_a_copy_of_the_solar_deck_fails_scope(tmp_path, variant, failures):
    assert score.main(["solar-market-refresh", str(built(tmp_path, "solar-market-refresh", variant))]) == score.FAIL
    report = json.loads((tmp_path / "score.json").read_text())
    assert report["verdict"] == "FAIL"
    assert [(f["code"], f["slide"], f["message"]) for f in report["failures"]] == failures


@pytest.mark.parametrize(
    ("variant", "message"),
    [
        ("chart_2027_bar_overwritten", "c1 does not ask to change these chart values on slide 10, but they are gone: 2000"),
        ("chart_rebuilt_without_potential", "c1 does not ask to change these chart values on slide 10, but they are gone: 11000"),
    ],
)
def test_a_solar_chart_value_the_edit_does_not_target_must_survive(tmp_path, variant, message):
    assert score.main(["solar-market-refresh", str(built(tmp_path, "solar-market-refresh", variant))]) == score.FAIL
    report = json.loads((tmp_path / "score.json").read_text())
    assert [(f["code"], f["slide"], f["message"]) for f in report["failures"]] == [("lost", 10, message)]


def test_a_chart_the_scorer_cannot_read_leaves_the_deck_unreadable(private_dir, tmp_path, capsys):
    gone = [{"id": "n1", "kind": "not-a-change", "said": ["00:00:05"], "why": "The 2022 bar stays.", "absent": [{"chart": 380}]}]
    d = scenario_at(private_dir / "bar3d", make=lambda path: chart_deck(path, "bar3DChart"), non_changes=gone)
    shutil.copy(d / "input.pptx", tmp_path / "output.pptx")
    assert score.main([str(d), str(tmp_path / "output.pptx")]) == score.BAD
    assert capsys.readouterr().out.splitlines()[0] == "SCENARIO UNREADABLE: [unreadable] slide 1 chart: bar3DChart cannot be read; check by hand"
    assert json.loads((tmp_path / "score.json").read_text())["verdict"] == "UNREADABLE"


def test_a_chart_relationship_to_a_missing_part_is_a_bad_output(tmp_path, capsys):
    source = open_scenario("solar-market-refresh").source_path
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(tmp_path / "output.pptx", "w") as zout:
        for name in zin.namelist():
            data = zin.read(name)
            if name == "ppt/slides/_rels/slide10.xml.rels":
                data = re.sub(rb'Target="\.\./charts/chart\d+\.xml"', b'Target="../charts/chartMISSING.xml"', data, count=1)
            zout.writestr(name, data)
    assert score.main(["solar-market-refresh", str(tmp_path / "output.pptx")]) == score.BAD
    assert "error: cannot read deck" in capsys.readouterr().err
    assert not (tmp_path / "score.json").exists()


def test_a_chart_fact_cannot_target_a_source_chart_that_cannot_be_read(private_dir):
    change = [{"id": "c1", "kind": "update-number", "intent": "Redraw the bar.", "said": ["00:00:05"], "slides": {1: {"forbid": [{"chart": 410}]}}}]
    assert open_scenario(scenario_at(private_dir / "bar", change, chart_deck)).source.charts[0].values == (Decimal("410.0"), Decimal("2000.0"))
    with pytest.raises(BadScenario, match="c1 slide 1: a chart fact cannot target a slide whose bar3DChart chart cannot be read"):
        open_scenario(scenario_at(private_dir / "bar3d", change, lambda path: chart_deck(path, "bar3DChart")))


UNCLEAR_CALLOUT = {"id": "n1", "kind": "ambiguous", "said": ["00:00:05"], "why": "Nobody said which callout.", "slides": [1], "flag": "Which callout should change?"}


@pytest.mark.parametrize(
    ("flags", "raised", "missing", "unmatched", "printed"),
    [
        (None, [], ["n1"], [], "flags (reported, not scored): raised none; missing n1; 0 unmatched; no flags.json"),
        ([{"question": "Which callout?", "said": ["00:00:05"]}], ["n1"], [], [], "flags (reported, not scored): raised n1; missing none; 0 unmatched"),
        ([{"question": "Which callout?", "slides": [1]}], ["n1"], [], [], "flags (reported, not scored): raised n1; missing none; 0 unmatched"),
        ([{"question": "Is the title final?", "slides": [2]}], [], ["n1"], ["Is the title final?"], "flags (reported, not scored): raised none; missing n1; 1 unmatched"),
    ],
)
def test_flags_json_next_to_the_deck_is_reported_beside_the_verdict(private_dir, tmp_path, capsys, flags, raised, missing, unmatched, printed):
    d = scenario_at(private_dir / "flags", non_changes=[UNCLEAR_CALLOUT])
    shutil.copy(d / "input.pptx", tmp_path / "output.pptx")
    if flags is not None:
        (tmp_path / "flags.json").write_text(json.dumps(flags))
    assert score.main([str(d), str(tmp_path / "output.pptx")]) == score.OK
    assert capsys.readouterr().out.splitlines()[1] == printed
    report = json.loads((tmp_path / "score.json").read_text())
    assert (report["verdict"], report["flags"]["raised"], report["flags"]["missing"], report["flags"]["unmatched"]) == ("PASS", raised, missing, unmatched)


@pytest.mark.parametrize(
    ("text", "printed"),
    [
        ('\ufeff[{"question": "Which callout?", "slides": [1]}]', "flags (reported, not scored): raised n1; missing none; 0 unmatched"),
        ('[{"question": "Which callout?", "said": ["0:05"], "id": "f1", "reason": "unclear"}]', "flags (reported, not scored): raised n1; missing none; 0 unmatched"),
        ('[{"question": "Which callout?", "slides": ["1"]}]', "flags (reported, not scored): raised n1; missing none; 0 unmatched"),
        ('[{"question": "Which callout?", "said": "00:00:05"}]', "flags (reported, not scored): raised n1; missing none; 0 unmatched"),
        ("not json", "flags (reported, not scored): raised none; missing n1; 0 unmatched; unreadable (flags.json: Expecting value: line 1 column 1 (char 0))"),
        ('{"question": "Which callout?"}', "flags (reported, not scored): raised none; missing n1; 0 unmatched; unreadable (flags.json: expected a list of flags)"),
        ('[{"question": "Which callout?"}]', "flags (reported, not scored): raised none; missing n1; 0 unmatched; unreadable (flags.json[0]: name the transcript turns in said or the source slides in slides)"),
        (
            '[{"question": "Which callout?", "said": ["5 minutes in"]}, {"question": "Which callout?", "slides": [1]}]',
            "flags (reported, not scored): raised n1; missing none; 0 unmatched; unreadable (flags.json[0].said: '5 minutes in' is not a transcript timestamp such as 00:08:05)",
        ),
        ('[{"question": "Which callout?", "said": ["1:05"]}]', "flags (reported, not scored): raised none; missing n1; 1 unmatched"),
        ('[{"question": "Which callout?", "slides": ["slide 1"]}]', "flags (reported, not scored): raised none; missing n1; 0 unmatched; unreadable (flags.json[0].slides: 'slide 1' is not a source slide number)"),
        (
            '[{"question": "Which callout?", "slides": ["' + "9" * 5000 + '"]}]',
            "flags (reported, not scored): raised none; missing n1; 0 unmatched; unreadable (flags.json[0].slides: '" + "9" * 27 + "..." + "9" * 28 + "' is not a source slide number)",
        ),
        (
            '[{"question": "Which callout?", "slides": [' + "9" * 5000 + "]}]",
            "flags (reported, not scored): raised none; missing n1; 0 unmatched; unreadable (flags.json: Exceeds the limit (4300 digits) for integer string conversion: value has 5000 digits; use sys.set_int_max_str_digits() to increase the limit)",
        ),
        ("[" * 100_000 + "]" * 100_000, "flags (reported, not scored): raised none; missing n1; 0 unmatched; unreadable (flags.json: nested too deeply to read)"),
    ],
)
def test_a_flags_json_in_any_shape_is_reported_and_the_deck_still_scored(private_dir, tmp_path, capsys, text, printed):
    d = scenario_at(private_dir / "flags", non_changes=[UNCLEAR_CALLOUT])
    shutil.copy(d / "input.pptx", tmp_path / "output.pptx")
    (tmp_path / "flags.json").write_text(text, encoding="utf-8")
    assert score.main([str(d), str(tmp_path / "output.pptx")]) == score.OK
    assert capsys.readouterr().out.splitlines()[1] == printed
    assert json.loads((tmp_path / "score.json").read_text())["verdict"] == "PASS"


def test_a_flags_file_never_hides_a_failing_deck(tmp_path, capsys):
    output = built(tmp_path, "insurance-workshop-prep", "kept_first_answer")
    (tmp_path / "flags.json").write_text(json.dumps([{"question": "Which regulator slide?", "slides": [18], "id": "f1"}]))
    assert score.main(["insurance-workshop-prep", str(output)]) == score.FAIL
    report = json.loads((tmp_path / "score.json").read_text())
    assert (report["verdict"], report["flags"]["raised"], report["flags"]["unreadable"]) == ("FAIL", ["n2"], [])
    assert capsys.readouterr().out.splitlines()[1] == "flags (reported, not scored): raised n2; missing none; 0 unmatched"


def test_a_garbage_flags_file_still_lets_the_deck_fail(tmp_path, capsys):
    output = built(tmp_path, "insurance-workshop-prep", "kept_first_answer")
    (tmp_path / "flags.json").write_bytes(b"\xff\xfe garbage")
    assert score.main(["insurance-workshop-prep", str(output)]) == score.FAIL
    assert json.loads((tmp_path / "score.json").read_text())["verdict"] == "FAIL"
    assert capsys.readouterr().out.splitlines()[1].startswith("flags (reported, not scored): raised none; missing n2; 0 unmatched; unreadable (flags.json: ")


def test_a_solar_flag_on_the_slide_that_holds_the_tariff_raises_the_tariff_question(tmp_path):
    output = built(tmp_path, "solar-market-refresh", "good")
    (tmp_path / "flags.json").write_text(json.dumps([{"question": "Which of Priya's tariffs replaces the $0.08?", "slides": [10]}]))
    assert score.main(["solar-market-refresh", str(output)]) == score.OK
    assert json.loads((tmp_path / "score.json").read_text())["flags"]["raised"] == ["n1"]


def test_flag_slides_names_only_a_slide_a_change_edits(private_dir):
    ask = {**UNCLEAR_CALLOUT, "slides": [], "flag_slides": [1]}
    with pytest.raises(BadScenario, match="n1.flag_slides: no change edits or deletes slide 1, so list it under slides"):
        open_scenario(scenario_at(private_dir / "unedited", non_changes=[ask]))
    sc = open_scenario(scenario_at(private_dir / "edited", EDIT_SLIDE_1, non_changes=[ask]))
    assert sc.non_changes[0].cited_by == {1}


def two_slides(path: Path) -> Path:
    prs = Presentation()
    for text in ("Demand grows 4% a year", "Supply grows 4% a year"):
        prs.slides.add_slide(prs.slide_layouts[1]).placeholders[1].text_frame.text = text
    prs.save(path)
    return path


def test_a_deck_wide_forbid_must_be_absent_from_every_slide_no_change_edits(private_dir):
    def change(forbid: dict) -> list[dict]:
        return [{"id": "c1", "kind": "update-number", "intent": "Move the growth rate.", "said": ["00:00:05"], "slides": {1: {"forbid": [forbid]}}}]

    assert open_scenario(scenario_at(private_dir / "slide", change({"percent": "4%"}), two_slides)).changes[0].slides[1].forbid[0].where == "slide"
    with pytest.raises(BadScenario, match="c1 slide 1 forbid '4%': a deck-wide forbid must be absent from every slide no change edits, but source slide 2 has it"):
        open_scenario(scenario_at(private_dir / "deck", change({"percent": "4%", "where": "deck"}), two_slides))


def test_a_deck_wide_forbid_fails_every_slide_that_states_it(tmp_path):
    output = built(tmp_path, "retail-impact-title", "tripled_on_another_slide")
    assert score.main(["retail-impact-title", str(output)]) == score.FAIL
    report = json.loads((tmp_path / "score.json").read_text())
    assert [(f["code"], f["slide"], f["message"]) for f in report["failures"] if f["code"] == "forbidden"] == [
        ("forbidden", 5, "c1: 'tripled' is on slide 5; it was abandoned after 00:06:54")
    ]

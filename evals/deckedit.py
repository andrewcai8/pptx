"""Build golden outputs: python-pptx edits addressed by SOURCE slide number, declared as @variant functions."""

from __future__ import annotations

import importlib.util
import io
import sys
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.slide import Slide as PptxSlide

from deckcheck.cli import write_atomic
from deckcheck.model import open_presentation
from scenario import ROOT, Scenario
from score import Code

BUILD_ROOT = ROOT / "artifacts/evals/build"


class DeckEdit:
    def __init__(self, sc: Scenario) -> None:
        self._prs = open_presentation(sc.source_path.read_bytes(), sc.source_path)
        self._list = self._prs.slides._sldIdLst
        self._ids = dict(enumerate(self._list, start=1))
        self._slides = dict(enumerate(self._prs.slides, start=1))
        self._new: dict[str, PptxSlide] = {}
        self._tail: dict[int, object] = {}
        self._deleted: set[int] = set()

    def slide(self, slide: int | str) -> PptxSlide:
        if isinstance(slide, str):
            return self._new[slide]
        if slide in self._deleted:
            raise ValueError(f"slide {slide} was deleted")
        return self._slides[slide]

    def slides(self) -> list[PptxSlide]:
        return [s for k, s in self._slides.items() if k not in self._deleted] + list(self._new.values())

    def replace(self, slide: int, old: str, new: str, *, count: int = 1) -> None:
        if not old:
            raise ValueError("replace needs a non-empty string to find")
        paragraphs = [(p, list(p.r_lst)) for p in self.slide(slide)._element.iter(qn("a:p"))]
        hits = [(runs, at) for _, runs in paragraphs for at in _find_all("".join(r.text for r in runs), old)]
        if len(hits) != count:
            raise ValueError(f"slide {slide}: {old!r} occurs {len(hits)} times in single paragraphs, expected {count}")
        for runs, at in reversed(hits):
            _splice(runs, at, at + len(old), new)

    def add_slide(self, change_id: str, *, after: int, layout: str, title: str, bullets: Sequence[str] = ()) -> None:
        source_layout = next((s.slide_layout for s in self._slides.values() if s.slide_layout.name == layout), None)
        if source_layout is None:
            raise ValueError(f"no source slide uses layout {layout!r}")
        slide = self._prs.slides.add_slide(source_layout)
        if slide.shapes.title is None:
            raise ValueError(f"layout {layout!r} has no title placeholder")
        slide.shapes.title.text = title
        if bullets:
            body = next((ph for ph in slide.placeholders if ph.placeholder_format.type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)), None)
            if body is None:
                raise ValueError(f"layout {layout!r} has no body placeholder for bullets")
            body.text_frame.text = bullets[0]
            for b in bullets[1:]:
                body.text_frame.add_paragraph().text = b
        self._new[change_id] = slide
        self._place(self._list[-1], after)

    def delete_slide(self, slide: int) -> None:
        self.slide(slide)
        sld_id = self._ids[slide]
        self._list.remove(sld_id)
        self._prs.part.drop_rel(sld_id.rId)
        self._deleted.add(slide)

    def move_slide(self, slide: int, *, after: int) -> None:
        self.slide(slide)
        if after == slide:
            raise ValueError(f"slide {slide} cannot follow itself")
        self._place(self._ids[slide], after)

    def set_size(self, slide: int, find: str, pt: float) -> None:
        paragraphs = [p for p in self.slide(slide)._element.iter(qn("a:p")) if find in "".join(r.text for r in p.r_lst)]
        if not paragraphs:
            raise ValueError(f"slide {slide}: no paragraph contains {find!r}")
        for p in paragraphs:
            for r in p.r_lst:
                r.get_or_add_rPr().set("sz", str(round(pt * 100)))

    def save(self, path: Path) -> None:
        path = Path(path).resolve()
        if ROOT in path.parents and ROOT / "artifacts" not in path.parents:
            raise ValueError(f"{path}: decks inside the repo go under artifacts/ only")
        buf = io.BytesIO()
        self._prs.save(buf)
        write_atomic(path, buf.getvalue())

    def _place(self, sld_id, after: int) -> None:
        if after in self._deleted:
            raise ValueError(f"slide {after} was deleted; nothing can follow it")
        anchor = self._tail.get(after, None if after == 0 else self._ids[after])
        if anchor is None:
            self._list.insert(0, sld_id)
        else:
            anchor.addnext(sld_id)
        self._tail[after] = sld_id


def _find_all(text: str, old: str) -> list[int]:
    found, at = [], text.find(old)
    while at >= 0:
        found.append(at)
        at = text.find(old, at + len(old))
    return found


def _splice(runs: list, start: int, end: int, new: str) -> None:
    offset = 0
    first = True
    for r in runs:
        text = r.text
        lo, hi = offset, offset + len(text)
        offset = hi
        if hi <= start or lo >= end:
            continue
        head = text[: max(start - lo, 0)] if first else ""
        tail = text[end - lo :] if end <= hi else ""
        r.text = head + (new if first else "") + tail
        first = False


@dataclass(frozen=True)
class Variant:
    name: str
    build: Callable[[DeckEdit], None]
    base: Variant | None
    fails: frozenset[tuple[Code, int | str]]
    doc: str

    def apply(self, d: DeckEdit) -> None:
        if self.base:
            self.base.apply(d)
        self.build(d)


def variant(*, base: Variant | None = None, fails: Iterable[tuple[str, int | str]] = ()) -> Callable[[Callable[[DeckEdit], None]], Variant]:
    parsed = frozenset((Code(code), slide) for code, slide in fails)

    def wrap(fn: Callable[[DeckEdit], None]) -> Variant:
        doc = (fn.__doc__ or "").strip()
        if not doc or "\n" in doc:
            raise ValueError(f"variant {fn.__name__} needs a one-line docstring naming what it gets right or wrong")
        return Variant(fn.__name__, fn, base, parsed, doc)

    return wrap


def load_variants(build_py: Path) -> list[Variant]:
    spec = importlib.util.spec_from_file_location(f"golden_build_{build_py.parent.name.replace('-', '_')}", build_py)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    seen: dict[int, Variant] = {}
    for v in vars(module).values():
        if isinstance(v, Variant):
            seen.setdefault(id(v), v)
    return list(seen.values())


def build_all(sc: Scenario, out_root: Path = BUILD_ROOT) -> list[tuple[Variant, Path]]:
    built = []
    for v in load_variants(sc.dir / "build.py"):
        d = DeckEdit(sc)
        v.apply(d)
        path = out_root / sc.name / v.name / "output.pptx"
        d.save(path)
        print(f"built {sc.name}/{v.name} -> {path}", file=sys.stderr)
        built.append((v, path))
    return built

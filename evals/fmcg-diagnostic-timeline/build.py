from pptx.util import Emu, Pt

from deckedit import DeckEdit, variant

TITLE = "STEP 1: Diagnose in 6 weeks, ready for the 2 December steering committee"
BULLETS = (
    "6 weeks from kick-off on 19 October to readout",
    "All 10 levers assessed against peers, across both pages of the imperatives house",
    "Deliverables: peer benchmark and a heatmap of where you are ahead, behind or at risk on each lever",
    "Readout at the steering committee on 2 December",
)


# The layout's level 0 renders as unbulleted 10pt text, unlike the 16 to 18pt bullets on the deck's own text slides.
def add_diagnostic(d: DeckEdit, *, after: int = 13, title: str = TITLE, bullets=BULLETS) -> None:
    d.add_slide("c1", after=after, layout="D. Title and Text", title=title, bullets=bullets)
    body = next(ph for ph in d.slide("c1").placeholders if ph.placeholder_format.idx != 0)
    for p in body.text_frame.paragraphs:
        p.level = 1
        for r in p.runs:
            r.font.size = Pt(16)


@variant()
def good(d: DeckEdit) -> None:
    """Diagnostic slide on the deck's own text layout, straight after Next steps, with 6 weeks and 10 levers."""
    add_diagnostic(d)


@variant(fails={("structure", "c1")})
def appended_at_end(d: DeckEdit) -> None:
    """Diagnostic slide appended after the last slide instead of straight after Next steps."""
    add_diagnostic(d, after=16)


@variant(fails={("structure", "c1")})
def placed_after_contacts(d: DeckEdit) -> None:
    """Diagnostic slide placed after the contacts page, the placement the client took back."""
    add_diagnostic(d, after=14)


@variant(fails={("layout", "c1")})
def title_only_with_textbox(d: DeckEdit) -> None:
    """Diagnostic slide built on D. Title Only with the bullets in a hand-drawn text box."""
    d.add_slide("c1", after=13, layout="D. Title Only", title=TITLE)
    slide = d.slide("c1")
    title = slide.shapes.title
    box = slide.shapes.add_textbox(title.left, Emu(title.top + title.height + Pt(18)), title.width, Pt(260))
    tf = box.text_frame
    tf.text = BULLETS[0]
    for b in BULLETS[1:]:
        tf.add_paragraph().text = b


@variant(fails={("style", "c1"), ("non-change", "c1")})
def tbd_team_bullet(d: DeckEdit) -> None:
    """Diagnostic slide carries a "Team size TBD" placeholder bullet the client asked to leave off."""
    add_diagnostic(d, bullets=BULLETS + ("Team size TBD",))


@variant(fails={("missing", "c1"), ("forbidden", "c1")})
def kept_eight_weeks(d: DeckEdit) -> None:
    """Diagnostic slide keeps the superseded 8 weeks instead of the 6 weeks the client settled on."""
    add_diagnostic(d, title=TITLE.replace("6 weeks", "8 weeks"), bullets=tuple(b.replace("6 weeks", "8 weeks") for b in BULLETS))


@variant(base=good, fails={("non-change", 14)})
def edited_contacts(d: DeckEdit) -> None:
    """Contacts page edited to swap an original author for the consultant, which marketing ruled out."""
    d.replace(14, "Matt Gamber", "Nadia Brennan")


@variant(fails={("structure", "c1"), ("non-change", 14)})
def after_contacts_and_edited_contacts(d: DeckEdit) -> None:
    """Diagnostic slide placed after the contacts page, and the contacts page edited to name the consultant."""
    add_diagnostic(d, after=14)
    d.replace(14, "Matt Gamber", "Nadia Brennan")

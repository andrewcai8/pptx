"""What a golden scenario is: expected.yaml and transcript.md parsed into types, linted against the real source deck."""

from __future__ import annotations

import contextlib
import copy
import csv
import hashlib
import importlib.util
import os
import re
import sys
import unicodedata
from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass
from functools import cached_property
from pathlib import Path
from typing import Any, Literal

import yaml
from lxml import etree
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part, XmlPart

from deckcheck.model import Deck, DeckError, Slide, open_presentation, read_bytes, read_deck

ROOT = Path(__file__).resolve().parents[1]
EVALS = ROOT / "evals"
PRIVATE_ENV = "GOLDEN_PRIVATE_DIR"


def _load_corpus():
    spec = importlib.util.spec_from_file_location("verify_pptx_corpus", ROOT / ".claude/skills/verify-pptx/scripts/corpus.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


corpus = _load_corpus()

Where = Literal["title", "slide"]
EditKind = Literal["edit-text", "update-number", "restyle"]
EDIT_KINDS: tuple[EditKind, ...] = ("edit-text", "update-number", "restyle")


class BadScenario(Exception):
    pass


class DeckUnreachable(Exception):
    pass


class SourceTampered(Exception):
    def __init__(self, name: str, ref: str, path: Path, got: str, pinned: str) -> None:
        super().__init__(f"source deck {path} hashes to {got}, the scenario pins {pinned}; something wrote into the source")
        self.name, self.ref, self.path, self.got, self.pinned = name, ref, path, got, pinned


@dataclass(frozen=True)
class Turn:
    at: str
    speaker: str
    text: str


@dataclass(frozen=True)
class FromSaid:
    at: str


@dataclass(frozen=True)
class FromData:
    data: str
    row: str
    column: str


@dataclass(frozen=True)
class FromSlide:
    slide: int


Provenance = FromSaid | FromData | FromSlide

Phrase = tuple[str, ...]


@dataclass(frozen=True)
class Require:
    text: Phrase
    where: Where
    source: Provenance | None


@dataclass(frozen=True)
class Forbid:
    text: Phrase
    where: Where
    superseded: str | None


@dataclass(frozen=True)
class Facts:
    require: tuple[Require, ...]
    forbid: tuple[Forbid, ...]


@dataclass(frozen=True)
class Edit:
    id: str
    kind: EditKind
    intent: str
    said: tuple[str, ...]
    slides: dict[int, Facts]
    intent_checks: tuple[str, ...]


@dataclass(frozen=True)
class AddSlide:
    id: str
    after: int
    layout: str
    intent: str
    said: tuple[str, ...]
    facts: Facts
    intent_checks: tuple[str, ...]


@dataclass(frozen=True)
class DeleteSlide:
    id: str
    slide: int
    intent: str
    said: tuple[str, ...]
    intent_checks: tuple[str, ...]


@dataclass(frozen=True)
class MoveSlide:
    id: str
    slide: int
    after: int
    intent: str
    said: tuple[str, ...]
    intent_checks: tuple[str, ...]


Change = Edit | AddSlide | DeleteSlide | MoveSlide


@dataclass(frozen=True)
class NotAChange:
    id: str
    said: tuple[str, ...]
    why: str
    slides: tuple[int, ...]
    absent: tuple[Phrase, ...]


@dataclass(frozen=True)
class Ambiguous:
    id: str
    said: tuple[str, ...]
    why: str
    slides: tuple[int, ...]
    absent: tuple[Phrase, ...]
    flag: str


NonChange = NotAChange | Ambiguous


@dataclass(frozen=True)
class CorpusDeck:
    id: str
    sha256: str


@dataclass(frozen=True)
class PrivateDeck:
    file: str
    sha256: str


DeckRef = CorpusDeck | PrivateDeck


@dataclass(frozen=True)
class SourceSlot:
    slide: int


@dataclass(frozen=True)
class NewSlot:
    change: str


Slot = SourceSlot | NewSlot


@dataclass(frozen=True)
class Look:
    xml: str
    parts: str


@dataclass(frozen=True)
class Snapshot:
    deck: Deck
    ids: tuple[int, ...]
    looks: tuple[Look, ...]
    shared: str


@dataclass(frozen=True)
class Scenario:
    name: str
    dir: Path
    deck: DeckRef
    source_path: Path
    source: Snapshot
    transcript: tuple[Turn, ...]
    changes: tuple[Change, ...]
    non_changes: tuple[NonChange, ...]
    skeleton: tuple[Slot, ...]

    @property
    def ref(self) -> str:
        return f"corpus:{self.deck.id}" if isinstance(self.deck, CorpusDeck) else f"file:{self.deck.file}"

    @cached_property
    def position(self) -> dict[Slot, int]:
        return {slot: i for i, slot in enumerate(self.skeleton, start=1)}

    @cached_property
    def edits(self) -> dict[int, tuple[Edit, ...]]:
        targets: dict[int, tuple[Edit, ...]] = {}
        for c in self.changes:
            if isinstance(c, Edit):
                for k in c.slides:
                    targets[k] = (*targets.get(k, ()), c)
        return targets

    @cached_property
    def deleted(self) -> dict[int, DeleteSlide]:
        return {c.slide: c for c in self.changes if isinstance(c, DeleteSlide)}

    @cached_property
    def moved(self) -> dict[int, MoveSlide]:
        return {c.slide: c for c in self.changes if isinstance(c, MoveSlide)}

    @cached_property
    def frozen(self) -> tuple[int, ...]:
        return tuple(s.slide for s in self.skeleton if isinstance(s, SourceSlot) and s.slide not in self.edits)

    @cached_property
    def named(self) -> dict[int, NonChange]:
        return {k: nc for nc in self.non_changes for k in nc.slides}

    @cached_property
    def fact_targets(self) -> tuple[tuple[Edit | AddSlide, Slot, Facts], ...]:
        out: list[tuple[Edit | AddSlide, Slot, Facts]] = []
        for c in self.changes:
            if isinstance(c, Edit):
                out += [(c, SourceSlot(k), facts) for k, facts in c.slides.items()]
            elif isinstance(c, AddSlide):
                out.append((c, NewSlot(c.id), c.facts))
        return tuple(out)

    @cached_property
    def deferred(self) -> tuple[tuple[str, str], ...]:
        checks = [(c.id, line) for c in self.changes for line in c.intent_checks]
        flags = [(nc.id, nc.flag) for nc in self.non_changes if isinstance(nc, Ambiguous)]
        return tuple(checks + flags)

    @property
    def words(self) -> int:
        return sum(len(t.text.split()) for t in self.transcript)


def skeleton(n_source: int, changes: Sequence[Change]) -> tuple[Slot, ...]:
    gone = {c.slide for c in changes if isinstance(c, DeleteSlide | MoveSlide)}
    inserts: dict[int, list[Slot]] = {}
    for c in changes:
        if isinstance(c, MoveSlide):
            inserts.setdefault(c.after, []).append(SourceSlot(c.slide))
        elif isinstance(c, AddSlide):
            inserts.setdefault(c.after, []).append(NewSlot(c.id))
    slots: list[Slot] = list(inserts.get(0, []))
    for k in range(1, n_source + 1):
        if k not in gone:
            slots.append(SourceSlot(k))
        slots += inserts.get(k, [])
    return tuple(slots)


WHITESPACE = re.compile(r"\s+")
QUOTES = str.maketrans({"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"'})
NUMBER = re.compile(r"\d+(?:[.,]\d+)*")


def normalize(text: str) -> str:
    return WHITESPACE.sub(" ", unicodedata.normalize("NFC", text).translate(QUOTES)).strip().casefold()


def _joined(a: str, b: str, c: str) -> bool:
    return a.isalnum() and b.isalnum() or a.isdigit() and b in ".," and c.isdigit()


def contains(fact: str, text: str) -> bool:
    fact, text = normalize(fact), f"  {normalize(text)}  "
    start = text.find(fact)
    while start >= 0:
        end = start + len(fact)
        if not _joined(fact[0], text[start - 1], text[start - 2]) and not _joined(fact[-1], text[end], text[end + 1]):
            return True
        start = text.find(fact, start + 1)
    return False


def holds(phrase: Phrase, text: str) -> str | None:
    return next((alt for alt in phrase if contains(alt, text)), None)


def says_number(number: str, text: str) -> bool:
    return re.search(rf"(?<![\d.,]){re.escape(number)}(?![\d]|[.,]\d)", normalize(text)) is not None


def texts(slide: Slide, where: Where) -> list[str]:
    if where == "title":
        return [slide.title]
    return [p.text for shape in slide.shapes for p in shape.paragraphs]


def find(phrase: Phrase, slide: Slide, where: Where = "slide") -> str | None:
    return next((alt for t in texts(slide, where) if (alt := holds(phrase, t))), None)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


SHARED = frozenset({CT.PML_SLIDE_LAYOUT, CT.PML_SLIDE_MASTER, CT.OFC_THEME})
NOT_SLIDE_CONTENT = frozenset({RT.SLIDE, RT.SLIDE_LAYOUT, RT.SLIDE_MASTER, RT.NOTES_MASTER})


def snapshot(path: Path) -> Snapshot:
    data = read_bytes(path)
    prs = open_presentation(data, path)
    deck = read_deck(prs, str(path), hashlib.sha256(data).hexdigest())
    slides = list(prs.slides)
    shared = sorted(_content(p) for p in prs.part.package.iter_parts() if p.content_type in SHARED)
    return Snapshot(deck, tuple(s.slide_id for s in slides), tuple(Look(_content(s.part), _related(s.part)) for s in slides), _digest(shared))


def _digest(items: Iterable[str]) -> str:
    return hashlib.sha256("\n".join(items).encode()).hexdigest()


# python-pptx adds empty elements (a:rPr, a:pPr) when code merely reads run.font or paragraph.level,
# so an untouched slide must compare equal with those dropped.
def _content(part: Part) -> str:
    if not isinstance(part, XmlPart):
        return hashlib.sha256(part.blob).hexdigest()
    root = copy.deepcopy(part._element)
    for el in reversed(list(root.iter(etree.Element))):
        if el is not root and not len(el) and not el.attrib and not el.text:
            el.getparent().remove(el)
    return hashlib.sha256(etree.tostring(root, method="c14n")).hexdigest()


def _related(part: Part) -> str:
    return _digest(
        sorted(
            f"{rel.rId} {rel.reltype} " + (rel.target_ref if rel.is_external else _content(rel.target_part) + _related(rel.target_part))
            for rel in part.rels.values()
            if rel.reltype not in NOT_SLIDE_CONTENT
        )
    )


def resolve(arg: str | Path) -> Path:
    path = Path(arg)
    if path.is_dir():
        if not (path / "expected.yaml").is_file():
            raise BadScenario(f"{path} has no expected.yaml")
        return path.resolve()
    name = str(arg)
    if "/" in name or name in ("", ".", ".."):
        raise BadScenario(f"no scenario directory {name}")
    private_root = os.environ.get(PRIVATE_ENV)
    candidates = [EVALS / name] + ([Path(private_root) / name] if private_root else [])
    found = [p.resolve() for p in candidates if (p / "expected.yaml").is_file()]
    if len(found) > 1:
        raise BadScenario(f"scenario {name} exists in both {found[0]} and {found[1]}; rename one")
    if not found:
        raise BadScenario(f"no scenario {name} in {' or '.join(str(p) for p in candidates)}")
    return found[0]


def public(names: Sequence[str] = ()) -> list[Path]:
    if not names:
        return [d for d in sorted(EVALS.iterdir()) if (d / "expected.yaml").is_file()]
    if missing := [n for n in names if not (EVALS / n / "expected.yaml").is_file()]:
        raise BadScenario(f"no public scenario {', '.join(missing)} in {EVALS}")
    return [EVALS / n for n in names]


def open_scenario(arg: str | Path) -> Scenario:
    d = resolve(arg)
    private = ROOT not in d.parents
    try:
        doc = yaml.safe_load((d / "expected.yaml").read_text())
    except (OSError, yaml.YAMLError) as e:
        raise BadScenario(f"{d / 'expected.yaml'}: {e}") from e
    doc = _keys(doc, "expected.yaml", required=("deck", "changes"), optional=("non_changes",))
    deck_ref = _deck_ref(doc["deck"], private)
    transcript = parse_transcript(d / "transcript.md")
    changes = tuple(_change(raw, f"changes[{i}]") for i, raw in enumerate(_list(doc["changes"], "changes")))
    non_changes = tuple(_non_change(raw, f"non_changes[{i}]") for i, raw in enumerate(_list(doc.get("non_changes") or [], "non_changes")))
    path = _source_path(d.name, deck_ref, d)
    try:
        source = snapshot(path)
    except DeckError as e:
        raise BadScenario(str(e)) from e
    sc = Scenario(d.name, d, deck_ref, path, source, transcript, changes, non_changes, skeleton(len(source.ids), changes))
    lint(sc)
    return sc


TURN = re.compile(r"^\[(\d\d:\d\d:\d\d)\] ([^:]+): (.+)$")
TIMESTAMP = re.compile(r"^\d\d:\d\d:\d\d$")


def parse_transcript(path: Path) -> tuple[Turn, ...]:
    try:
        lines = path.read_text().splitlines()
    except OSError as e:
        raise BadScenario(f"cannot read {path}: {e}") from e
    turns: list[Turn] = []
    for n, line in enumerate(lines, start=1):
        if m := TURN.match(line):
            turns.append(Turn(m[1], m[2].strip(), m[3].strip()))
        elif line.strip() and (turns or line.startswith("[")):
            raise BadScenario(f"{path.name} line {n} is not a turn of the form '[hh:mm:ss] Name (Role, Org): text'")
    if not turns:
        raise BadScenario(f"{path.name} has no turns")
    stamps = [t.at for t in turns]
    if stamps != sorted(set(stamps)):
        raise BadScenario(f"{path.name} timestamps must be unique and ascending")
    return tuple(turns)


def _keys(raw: Any, where: str, required: Sequence[str], optional: Sequence[str] = ()) -> dict[str, Any]:
    if not isinstance(raw, dict):
        raise BadScenario(f"{where}: expected a mapping")
    if missing := [k for k in required if k not in raw]:
        raise BadScenario(f"{where}: missing {', '.join(missing)}")
    if extra := sorted(str(k) for k in set(raw) - set(required) - set(optional)):
        raise BadScenario(f"{where}: unknown field {', '.join(extra)}")
    return raw


def _list(raw: Any, where: str) -> list[Any]:
    if not isinstance(raw, list):
        raise BadScenario(f"{where}: expected a list")
    return raw


def _str(raw: Any, where: str) -> str:
    if not isinstance(raw, str) or not raw.strip():
        raise BadScenario(f"{where}: expected a non-empty string")
    return raw


def _int(raw: Any, where: str) -> int:
    if type(raw) is not int:
        raise BadScenario(f"{where}: expected an integer slide number")
    return raw


def _stamp(raw: Any, where: str) -> str:
    if not isinstance(raw, str) or not TIMESTAMP.match(raw):
        raise BadScenario(f"{where}: expected a quoted hh:mm:ss timestamp")
    return raw


def _stamps(raw: Any, where: str) -> tuple[str, ...]:
    stamps = tuple(_stamp(s, f"{where}[{i}]") for i, s in enumerate(_list(raw, where)))
    if not stamps:
        raise BadScenario(f"{where}: name at least one transcript turn")
    return stamps


def _strs(raw: Any, where: str) -> tuple[str, ...]:
    return tuple(_str(s, f"{where}[{i}]") for i, s in enumerate(_list(raw or [], where)))


def _where(raw: Any, where: str) -> Where:
    if raw not in ("title", "slide"):
        raise BadScenario(f"{where}: where must be title or slide")
    return raw


def _deck_ref(raw: Any, private: bool) -> DeckRef:
    if isinstance(raw, dict) and "file" in raw:
        raw = _keys(raw, "deck", required=("file", "sha256"))
        if not private:
            raise BadScenario("deck: a deck file never enters the repo; put the scenario in $GOLDEN_PRIVATE_DIR")
        return PrivateDeck(_str(raw["file"], "deck.file"), _str(raw["sha256"], "deck.sha256"))
    raw = _keys(raw, "deck", required=("corpus", "sha256"))
    return CorpusDeck(_str(raw["corpus"], "deck.corpus"), _str(raw["sha256"], "deck.sha256"))


def _provenance(raw: Any, where: str) -> Provenance:
    if not isinstance(raw, dict) or len(set(raw) & {"said", "data", "slide"}) != 1:
        raise BadScenario(f"{where}: from must be one of {{said}}, {{data, row, column}}, {{slide}}")
    if "said" in raw:
        return FromSaid(_stamp(_keys(raw, where, ("said",))["said"], f"{where}.said"))
    if "slide" in raw:
        return FromSlide(_int(_keys(raw, where, ("slide",))["slide"], f"{where}.slide"))
    raw = _keys(raw, where, ("data", "row", "column"))
    return FromData(_str(raw["data"], f"{where}.data"), str(raw["row"]), _str(raw["column"], f"{where}.column"))


def _phrase(raw: Any, where: str) -> Phrase:
    if isinstance(raw, str):
        return (_str(raw, where),)
    alternatives = _strs(raw, where)
    if not alternatives:
        raise BadScenario(f"{where}: expected a string or a non-empty list of alternatives")
    return alternatives


def _facts(raw: dict[str, Any], where: str) -> Facts:
    require = []
    for i, r in enumerate(_list(raw.get("require") or [], f"{where}.require")):
        w = f"{where}.require[{i}]"
        r = _keys(r, w, ("text",), ("where", "from"))
        source = _provenance(r["from"], f"{w}.from") if "from" in r else None
        require.append(Require(_phrase(r["text"], f"{w}.text"), _where(r.get("where", "slide"), w), source))
    forbid = []
    for i, f in enumerate(_list(raw.get("forbid") or [], f"{where}.forbid")):
        w = f"{where}.forbid[{i}]"
        f = _keys(f, w, ("text",), ("where", "superseded"))
        superseded = _stamp(f["superseded"], f"{w}.superseded") if "superseded" in f else None
        forbid.append(Forbid(_phrase(f["text"], f"{w}.text"), _where(f.get("where", "slide"), w), superseded))
    return Facts(tuple(require), tuple(forbid))


COMMON = ("id", "kind", "intent", "said")


def _change(raw: Any, where: str) -> Change:
    if not isinstance(raw, dict):
        raise BadScenario(f"{where}: expected a mapping")
    kind = raw.get("kind")
    cid = _str(raw.get("id"), f"{where}.id")
    where = f"{where} ({cid})"
    if kind in EDIT_KINDS:
        raw = _keys(raw, where, (*COMMON, "slides"), ("intent_checks",))
        if not isinstance(raw["slides"], dict) or not raw["slides"]:
            raise BadScenario(f"{where}.slides: expected a mapping of source slide number to require/forbid")
        slides = {
            _int(k, f"{where}.slides key"): _facts(_keys(v or {}, f"{where}.slides.{k}", (), ("require", "forbid")), f"{where}.slides.{k}")
            for k, v in raw["slides"].items()
        }
        return Edit(cid, kind, _str(raw["intent"], f"{where}.intent"), _stamps(raw["said"], f"{where}.said"), slides, _strs(raw.get("intent_checks"), f"{where}.intent_checks"))
    if kind == "add-slide":
        raw = _keys(raw, where, (*COMMON, "after", "layout"), ("require", "forbid", "intent_checks"))
        return AddSlide(
            cid,
            _int(raw["after"], f"{where}.after"),
            _str(raw["layout"], f"{where}.layout"),
            _str(raw["intent"], f"{where}.intent"),
            _stamps(raw["said"], f"{where}.said"),
            _facts(raw, where),
            _strs(raw.get("intent_checks"), f"{where}.intent_checks"),
        )
    if kind == "delete-slide":
        raw = _keys(raw, where, (*COMMON, "slide"), ("intent_checks",))
        return DeleteSlide(cid, _int(raw["slide"], f"{where}.slide"), _str(raw["intent"], f"{where}.intent"), _stamps(raw["said"], f"{where}.said"), _strs(raw.get("intent_checks"), f"{where}.intent_checks"))
    if kind == "move-slide":
        raw = _keys(raw, where, (*COMMON, "slide", "after"), ("intent_checks",))
        return MoveSlide(
            cid,
            _int(raw["slide"], f"{where}.slide"),
            _int(raw["after"], f"{where}.after"),
            _str(raw["intent"], f"{where}.intent"),
            _stamps(raw["said"], f"{where}.said"),
            _strs(raw.get("intent_checks"), f"{where}.intent_checks"),
        )
    raise BadScenario(f"{where}.kind: expected one of {', '.join((*EDIT_KINDS, 'add-slide', 'delete-slide', 'move-slide'))}")


def _non_change(raw: Any, where: str) -> NonChange:
    if not isinstance(raw, dict):
        raise BadScenario(f"{where}: expected a mapping")
    kind = raw.get("kind")
    nid = _str(raw.get("id"), f"{where}.id")
    where = f"{where} ({nid})"
    if kind not in ("not-a-change", "ambiguous"):
        raise BadScenario(f"{where}.kind: expected not-a-change or ambiguous")
    raw = _keys(raw, where, ("id", "kind", "said", "why"), ("slides", "absent", "flag") if kind == "ambiguous" else ("slides", "absent"))
    said = _stamps(raw["said"], f"{where}.said")
    why = _str(raw["why"], f"{where}.why")
    slides = tuple(_int(s, f"{where}.slides[{i}]") for i, s in enumerate(_list(raw.get("slides") or [], f"{where}.slides")))
    absent = tuple(
        _phrase(_keys(a, f"{where}.absent[{i}]", ("text",))["text"], f"{where}.absent[{i}].text")
        for i, a in enumerate(_list(raw.get("absent") or [], f"{where}.absent"))
    )
    if kind == "not-a-change":
        return NotAChange(nid, said, why, slides, absent)
    if "flag" not in raw:
        raise BadScenario(f"{where}: an ambiguous non-change needs the flag the maker should raise")
    return Ambiguous(nid, said, why, slides, absent, _str(raw["flag"], f"{where}.flag"))


def _source_path(name: str, ref: DeckRef, d: Path) -> Path:
    if isinstance(ref, PrivateDeck):
        path = (d / ref.file).resolve()
        if not path.is_file():
            raise BadScenario(f"deck.file: no file {path}")
        if (got := sha256(path)) != ref.sha256:
            raise SourceTampered(name, f"file:{ref.file}", path, got, ref.sha256)
        return path
    decks = yaml.safe_load(corpus.MANIFEST.read_text())["decks"]
    entry = next((e for e in decks if e["id"] == ref.id), None)
    if entry is None:
        raise BadScenario(f"deck.corpus: no deck {ref.id} in {corpus.MANIFEST}")
    if entry["sha256"] != ref.sha256:
        raise BadScenario(f"deck.sha256: the manifest pins {ref.id} at {entry['sha256']}, the scenario at {ref.sha256}")
    cached = corpus.CACHE / f"{ref.sha256}.pptx"
    if cached.is_file() and (got := sha256(cached)) != ref.sha256:
        raise SourceTampered(name, f"corpus:{ref.id}", cached, got, ref.sha256)
    try:
        with contextlib.redirect_stdout(sys.stderr):
            return corpus.fetch(entry)
    except corpus.Unreachable as e:
        raise DeckUnreachable(f"{ref.id} unreachable: {e}") from e
    except corpus.BadDeck as e:
        raise BadScenario(f"{ref.id}: {e}") from e


def lint(sc: Scenario) -> None:
    for problem in _problems(sc):
        raise BadScenario(problem)


def _problems(sc: Scenario) -> Iterator[str]:
    slides = sc.source.deck.slides
    n = len(slides)
    turns = {t.at: t for t in sc.transcript}
    ids = [c.id for c in sc.changes] + [nc.id for nc in sc.non_changes]
    if dup := sorted({i for i in ids if ids.count(i) > 1}):
        yield f"ids {', '.join(dup)} are used twice"
    for item in (*sc.changes, *sc.non_changes):
        for at in item.said:
            if at not in turns:
                yield f"{item.id}.said: no transcript turn at {at}"

    def in_range(k: int, where: str, low: int = 1) -> Iterator[str]:
        if not low <= k <= n:
            yield f"{where}: slide {k} is outside the deck's {low}..{n}"

    structural: list[int] = []
    for c in sc.changes:
        match c:
            case Edit():
                for k in c.slides:
                    yield from in_range(k, f"{c.id}.slides")
            case AddSlide():
                yield from in_range(c.after, f"{c.id}.after", low=0)
                if c.layout not in {s.layout_name for s in slides}:
                    yield f"{c.id}.layout: no source slide uses layout {c.layout!r}"
            case DeleteSlide():
                yield from in_range(c.slide, f"{c.id}.slide")
                structural.append(c.slide)
            case MoveSlide():
                yield from in_range(c.slide, f"{c.id}.slide")
                yield from in_range(c.after, f"{c.id}.after", low=0)
                structural.append(c.slide)
                if c.after in (c.slide, c.slide - 1):
                    yield f"{c.id}: moving slide {c.slide} after {c.after} leaves it where it is"
    if dup := sorted({k for k in structural if structural.count(k) > 1}):
        yield f"slides {', '.join(map(str, dup))} are deleted or moved twice"
    for c in sc.changes:
        if isinstance(c, AddSlide | MoveSlide) and c.after in structural:
            yield f"{c.id}.after: slide {c.after} is itself deleted or moved; anchor on a slide that stays"
    for k, edits in sc.edits.items():
        if k in sc.deleted:
            yield f"slide {k} is both edited ({edits[0].id}) and deleted ({sc.deleted[k].id})"
        if any(e.kind == "restyle" for e in edits) and any(e.kind != "restyle" for e in edits):
            yield f"slide {k} is both restyled and text-edited; a restyle must keep the slide's text"
    yield from _fact_problems(sc, turns)

    targets = set(sc.edits) | set(sc.deleted)
    named = [k for nc in sc.non_changes for k in nc.slides]
    for nc in sc.non_changes:
        for k in nc.slides:
            yield from in_range(k, f"{nc.id}.slides")
            if k in targets:
                yield f"{nc.id}.slides: slide {k} is edited or deleted, so a non-change cannot name it"
        for phrase in nc.absent:
            if found := next((alt for s in slides if (alt := find(phrase, s))), None):
                yield f"{nc.id}.absent: {found!r} is already in the source deck"
    if dup := sorted({k for k in named if named.count(k) > 1}):
        yield f"slides {', '.join(map(str, dup))} are named by two non-changes"


def _said(alt: str, turn: Turn) -> bool:
    numbers = NUMBER.findall(alt)
    return contains(alt, turn.text) or bool(numbers) and all(says_number(num, turn.text) for num in numbers)


def _fact_problems(sc: Scenario, turns: dict[str, Turn]) -> Iterator[str]:
    slides = sc.source.deck.slides
    for change, slot, facts in sc.fact_targets:
        src = slides[slot.slide - 1] if isinstance(slot, SourceSlot) and 1 <= slot.slide <= len(slides) else None
        label = f"{change.id} slide {slot.slide}" if isinstance(slot, SourceSlot) else f"{change.id}"
        for r in facts.require:
            for alt in (a for a in r.text if any(ch.isdigit() for ch in a)):
                if r.source is None:
                    yield f"{label} require {alt!r}: a fact with a number needs `from`"
                else:
                    yield from (f"{label} require {alt!r}: {p}" for p in _provenance_problems(sc, alt, r.source, turns))
        for f in facts.forbid:
            if f.superseded is not None:
                if f.superseded not in turns:
                    yield f"{label} forbid {f.text!r}: no transcript turn at {f.superseded}"
                elif not any(_said(alt, turns[f.superseded]) for alt in f.text):
                    yield f"{label} forbid {f.text!r}: the turn at {f.superseded} must say the abandoned value"
                if src is not None and (found := find(f.text, src, f.where)):
                    yield f"{label} forbid {found!r}: a superseded value must be absent from the source slide"
            elif src is not None and not find(f.text, src, f.where):
                yield f"{label} forbid {f.text!r}: an old value must be on the source slide ({f.where})"


def _provenance_problems(sc: Scenario, text: str, source: Provenance, turns: dict[str, Turn]) -> Iterator[str]:
    match source:
        case FromSaid(at):
            if at not in turns:
                yield f"no transcript turn at {at}"
            elif missing := [num for num in NUMBER.findall(text) if not says_number(num, turns[at].text)]:
                yield f"the turn at {at} does not say {', '.join(missing)}"
        case FromSlide(k):
            if not 1 <= k <= len(sc.source.deck.slides):
                yield f"from.slide {k} is outside the deck"
            elif not find((text,), sc.source.deck.slides[k - 1]):
                yield f"the text is not on source slide {k}"
        case FromData(data, row, column):
            path = (sc.dir / data).resolve()
            if sc.dir not in path.parents or not path.is_file():
                yield f"from.data: no file {data} in the scenario directory"
                return
            with path.open(newline="") as f:
                rows = list(csv.reader(f))
            header = rows[0] if rows else []
            hits = [r for r in rows[1:] if r and r[0] == row]
            if column not in header:
                yield f"from.data: {data} has no column {column!r}"
            elif len(hits) != 1:
                yield f"from.data: {data} has {len(hits)} rows keyed {row!r}, expected 1"
            elif not says_number(cell := hits[0][header.index(column)].strip(), text):
                yield f"from.data: {data} {row}/{column} is {cell!r}, which the fact text does not contain"

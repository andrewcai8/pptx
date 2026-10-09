import copy

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from deckedit import DeckEdit, variant

SURVEY_NOTE = "3. Insurers' Association member survey 2026, 2027 projection"
CENSUS_NOTE = "1. U.S. Census Bureau, Population Division.  2014 estimate of population; Generations as defined by Pew Research Center, 2014"


def footnote_paragraphs(d: DeckEdit) -> list:
    return [p._p for p in next(s for s in d.slide(5).shapes if s.name == "Footnote").text_frame.paragraphs]


def cite_survey(d: DeckEdit, *, text: str = SURVEY_NOTE, first: bool = False) -> None:
    notes = footnote_paragraphs(d)
    note = copy.deepcopy(notes[-1])
    run, *rest = note.r_lst
    run.text = text
    for r in rest:
        note.remove(r)
    if first:
        notes[0].addprevious(note)
    else:
        notes[-1].addnext(note)


def rewrite_footnote(d: DeckEdit, line: int, text: str) -> None:
    note = footnote_paragraphs(d)[line]
    run, *rest = note.r_lst
    run.text = text
    for r in rest:
        note.remove(r)


def shape(d: DeckEdit, name: str):
    return next(s for s in d.slide(5).shapes if s.name == name)


def note_box(d: DeckEdit, text: str) -> None:
    """A new text box under the Footnote, set in the Footnote's own font and size."""
    footnote = shape(d, "Footnote")
    run = footnote.text_frame.paragraphs[0].runs[0]
    box = d.slide(5).shapes.add_textbox(footnote.left, footnote.top + footnote.height, footnote.width, Pt(12))
    box.text_frame.text = text
    font = box.text_frame.paragraphs[0].runs[0].font
    font.size, font.name = run.font.size or Pt(8), run.font.name


def refresh(d: DeckEdit) -> None:
    d.replace(5, "ca. 50%", "ca. 48%")
    d.delete_slide(2)
    d.move_slide(14, after=11)


@variant()
def good(d: DeckEdit) -> None:
    """Survey figure in the headline with a new footnote line citing the survey, credentials page gone, innovation leads the three roles."""
    refresh(d)
    cite_survey(d)


@variant()
def survey_added_to_footnote_1(d: DeckEdit) -> None:
    """Footnote 1 is extended to cite the survey for the headline, the literal reading of "I'll point the footnote at your survey file"."""
    refresh(d)
    d.replace(5, CENSUS_NOTE, CENSUS_NOTE + "; headline share: Insurers' Association member survey (2027 projection)")


@variant()
def survey_footnote_listed_first(d: DeckEdit) -> None:
    """The survey citation is a new footnote line above the census one, the other reading of pointing the footnote at the survey file."""
    refresh(d)
    cite_survey(d, text="Headline share: Insurers' Association member survey 2026 (2027 projection)", first=True)


@variant(base=good, fails={("lost", 5)})
def fixed_callout_spelling(d: DeckEdit) -> None:
    """Beyond the headline and footnote, slide 5's callout gets "Millenials" corrected, a line no change asks about."""
    d.replace(5, "...and Millenials have", "...and Millennials have")


@variant(base=good, fails={("guessed", 18)})
def flag_in_speaker_note(d: DeckEdit) -> None:
    """The open question about the regulator slide is written as a speaker note on slide 18 instead of in flags.json, so it is an edit."""
    d.slide(18).notes_slide.notes_text_frame.text = "Marko: which regulator slide should be punchier? Check before editing."


@variant(base=good, fails={("guessed", 18)})
def guessed_regulator_slide(d: DeckEdit) -> None:
    """Picks slide 18 as "the regulator slide" and rewrites it instead of asking which one."""
    d.replace(18, "have to closely cooperate", "must co-own the digital agenda")


@variant(fails={("missing", 5), ("forbidden", 5)})
def kept_first_answer(d: DeckEdit) -> None:
    """Uses the 45% Marko remembered and ignores the survey file he read out later."""
    d.replace(5, "ca. 50%", "ca. 45%")
    d.delete_slide(2)
    d.move_slide(14, after=11)


@variant(fails={("forbidden", 5)})
def hedged_range_45_to_48(d: DeckEdit) -> None:
    """The headline hedges with "ca. 45-48%", keeping the 45% Marko took back once he had the survey file."""
    refresh(d)
    d.replace(5, "ca. 48%", "ca. 45-48%")
    cite_survey(d)


@variant(fails={("structure", 2)})
def kept_credentials(d: DeckEdit) -> None:
    """Applies the number and the move but leaves the credentials page in."""
    d.replace(5, "ca. 50%", "ca. 48%")
    d.move_slide(14, after=11)


@variant(base=good, fails={("scope", 1), ("non-change", 1)})
def acted_on_logistics(d: DeckEdit) -> None:
    """Puts the Oversight Office on the cover although the invitation is the client's action and unconfirmed."""
    d.replace(1, "Ljubljana, September 4th, 2017", "Joint workshop with the Oversight Office, Ljubljana")


@variant(base=good, fails={("non-change", 12), ("non-change", 13), ("non-change", 14)})
def renumbered_roles(d: DeckEdit) -> None:
    """Renumbers the role pages after the move although the client said to keep the numbers."""
    d.replace(14, "3", "1")
    d.replace(12, "1", "2")
    d.replace(13, "2", "3")


@variant(base=good, fails={("scope", 9)})
def restyled_frozen_slide(d: DeckEdit) -> None:
    """Enlarges the title on slide 9, which nobody asked about; the text is the same, the XML is not."""
    d.set_size(9, "There are typical challenges", 32)


@variant(base=good, fails={("style", 5)})
def shrunk_footnote(d: DeckEdit) -> None:
    """Shrinks the slide 5 survey footnote to 6pt, below the house minimum."""
    d.set_size(5, "BCG digital satisfaction survey", 6)


@variant(base=good, fails={("structure", 18), ("guessed", 18)})
def dropped_regulator_slide(d: DeckEdit) -> None:
    """Deletes the closing regulator slide as a way of making "the regulator slide" punchier."""
    d.delete_slide(18)


@variant(fails={("structure", 14)})
def moved_after_wrong_slide(d: DeckEdit) -> None:
    """Puts genuine innovation after Digitize the core instead of straight after the roles overview."""
    d.replace(5, "ca. 50%", "ca. 48%")
    d.delete_slide(2)
    d.move_slide(14, after=12)


@variant(base=good, fails={("structure", 5)})
def moved_edited_slide(d: DeckEdit) -> None:
    """The generations slide gets the survey figure and is also moved after slide 6, which nobody asked for."""
    d.move_slide(5, after=6)


@variant(base=good, fails={("guessed", 18)})
def cleared_regulator_panel(d: DeckEdit) -> None:
    """The white panel behind slide 18's cartoon is made transparent, a guess at making "the regulator slide" land harder."""
    next(s for s in d.slide(18).shapes if s.name == "Rectangle 5").fill.background()


@variant(base=good, fails={("non-change", 17)})
def autofit_sandbox_title(d: DeckEdit) -> None:
    """Slide 17's title is set to shrink on overflow, which renders it smaller on the sandbox page Marko called fine."""
    d.slide(17).shapes.title.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


@variant(fails={("forbidden", 5)})
def hedged_range_45_to_48_in_words(d: DeckEdit) -> None:
    """The headline hedges with "ca. 45 to 48%", the same range as 45-48% written with "to"."""
    refresh(d)
    d.replace(5, "ca. 48%", "ca. 45 to 48%")
    cite_survey(d)


@variant(fails={("lost", 5)})
def survey_replaced_footnote_2(d: DeckEdit) -> None:
    """Footnote 2, the BCG survey behind the callout, is overwritten with the member survey, the likely misreading of "point the footnote at your survey file"."""
    refresh(d)
    rewrite_footnote(d, 1, "2. Insurers' Association member survey 2026, 2027 projection")


@variant(fails={("lost", 5)})
def survey_replaced_census_line(d: DeckEdit) -> None:
    """Footnote 1 is overwritten with the member survey, so the curve loses its US census source."""
    refresh(d)
    rewrite_footnote(d, 0, "1. Insurers' Association member survey 2026, 2027 projection")


@variant(fails={("lost", 5)})
def footnote_rewritten(d: DeckEdit) -> None:
    """The whole footnote is rewritten into one survey line and a short source line."""
    refresh(d)
    rewrite_footnote(d, 0, "Headline share: Insurers' Association member survey 2026 (2027 projection)")
    rewrite_footnote(d, 1, "Curve: U.S. Census Bureau 2014")


@variant(base=good, intent=("c1", "adds nothing unrelated to the survey citation"))
def lunch_added_to_footnote(d: DeckEdit) -> None:
    """Beside the survey line, the footnote gains a sentence about lunch, which no change asks for."""
    cite_survey(d, text="Lunch at the workshop is provided by the hotel")


@variant(base=good, intent=("c1", "adds nothing unrelated to the survey citation"))
def lunch_added_in_a_text_box(d: DeckEdit) -> None:
    """Slide 5 gains a new text box about lunch at the workshop, which no change asks for."""
    title = d.slide(5).shapes.title
    d.slide(5).shapes.add_textbox(title.left, title.top + title.height, title.width, title.height).text_frame.text = "Lunch at the workshop is provided by the hotel"


@variant(fails={("forbidden", 5)})
def old_share_in_footnote(d: DeckEdit) -> None:
    """The footnote cites the survey but keeps the old ca. 50% beside it, a value the meeting replaced."""
    refresh(d)
    cite_survey(d, text="3. Insurers' Association member survey 2026, 2027 projection; ca. 50% in the 2017 material")


@variant()
def title_split_into_two_paragraphs(d: DeckEdit) -> None:
    """The headline is split at its dash into two paragraphs, as pressing Enter in PowerPoint gives, with the survey figure in the second."""
    refresh(d)
    cite_survey(d)
    d.slide(5).shapes.title.text_frame.text = "The emergence of new Generations further accelerates the digitalization\nBy 2027 in Slovenia Millennials will account for ca. 48% of the client base"


@variant()
def survey_footnote_in_a_new_text_box(d: DeckEdit) -> None:
    """The survey citation is footnote 3 in a new text box under the Footnote, not inside it."""
    refresh(d)
    note_box(d, "3. Insurers' Association member survey (2026), 2027 projection")


@variant()
def source_line_in_a_new_box(d: DeckEdit) -> None:
    """A new box under the footnotes says "Source: Insurers' Association member survey 2026" for the headline."""
    refresh(d)
    note_box(d, "Source: Insurers' Association member survey 2026")


@variant()
def footnotes_renumbered_survey_first(d: DeckEdit) -> None:
    """The survey becomes footnote 1, and the census and BCG footnotes are renumbered 2 and 3."""
    refresh(d)
    cite_survey(d, text="1. Insurers' Association member survey 2026, 2027 projection", first=True)
    d.replace(5, CENSUS_NOTE, "2" + CENSUS_NOTE[1:])
    d.replace(5, "2. BCG digital satisfaction survey", "3. BCG digital satisfaction survey")


@variant(fails={("lost", 5)})
def curve_graphic_deleted(d: DeckEdit) -> None:
    """The generations curve, an embedded graph, is deleted from slide 5 while the headline is updated."""
    refresh(d)
    cite_survey(d)
    curve = shape(d, "Object 12")
    curve._element.getparent().remove(curve._element)


@variant(fails={("lost", 5)})
def curve_graphic_moved_off_slide(d: DeckEdit) -> None:
    """The generations curve is dragged past the slide's right edge, so the slide no longer shows it."""
    refresh(d)
    cite_survey(d)
    shape(d, "Object 12").left = Inches(30)


@variant(base=good, intent=("c1", "an extended footnote gains only words about the survey"))
def footnote_extended_with_unrelated_words(d: DeckEdit) -> None:
    """Footnote 1 is extended with the survey and with a remark about the workshop venue, which no change asks for."""
    d.replace(5, CENSUS_NOTE, CENSUS_NOTE + "; headline share: member survey; workshop venue to be confirmed")


@variant()
def headline_reworded_whole(d: DeckEdit) -> None:
    """The whole headline, one sentence across a line break, is reworded around the survey figure; the footnote cites the survey."""
    refresh(d)
    cite_survey(d)
    d.slide(5).shapes.title.text_frame.text = "New generations speed up digitalization: by 2027 millennials will be ca. 48% of Slovenian insurers' clients"


def survey_over_footnote_2(d: DeckEdit, *, left=None, top=None, white: bool = False) -> None:
    """Footnote 2 is overwritten with the member survey, and its old text goes into a new 8pt box at left, top."""
    old = footnote_paragraphs(d)[1].xpath("string(.)")
    refresh(d)
    rewrite_footnote(d, 1, "2. Insurers' Association member survey 2026, 2027 projection")
    footnote = shape(d, "Footnote")
    box = d.slide(5).shapes.add_textbox(footnote.left if left is None else left, footnote.top if top is None else top, footnote.width, Pt(12))
    box.text_frame.text = old
    font = box.text_frame.paragraphs[0].runs[0].font
    font.size = Pt(8)
    if white:
        font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        tree = d.slide(5).shapes._spTree
        tree.insert(2, box._element)


@variant(fails={("lost", 5)})
def footnote_2_kept_off_slide(d: DeckEdit) -> None:
    """Footnote 2 is overwritten with the member survey, and its old text is parked in a text box past the slide's right edge."""
    survey_over_footnote_2(d, left=Inches(30))


@variant(base=good, fails={("lost", 5)})
def footnote_box_pushed_off_slide(d: DeckEdit) -> None:
    """The whole Footnote box, survey line included, is dragged below the slide, so no footnote shows."""
    shape(d, "Footnote").top = Inches(30)


@variant(intent=("c1", "hidden in white"))
def footnote_2_kept_in_white_behind_curve(d: DeckEdit) -> None:
    """Footnote 2 is overwritten with the member survey, and its old text sits at 8pt white behind the curve, where no reader sees it."""
    curve = shape(d, "Object 12")
    survey_over_footnote_2(d, left=curve.left, top=curve.top, white=True)


@variant(fails={("lost", 5)})
def footnote_2_kept_in_hidden_box(d: DeckEdit) -> None:
    """Footnote 2 is overwritten with the member survey, and its old text sits in a box the selection pane hides."""
    survey_over_footnote_2(d)
    d.slide(5).shapes[-1]._element.nvSpPr.cNvPr.set("hidden", "1")


@variant(base=good, fails={("lost", 5)})
def curve_graphic_hidden(d: DeckEdit) -> None:
    """The census curve graphic stays in place but is hidden, so the slide shows no curve."""
    shape(d, "Object 12")._element[0][0].set("hidden", "1")
